#!/usr/bin/env python3
from __future__ import annotations

import argparse
import csv
import datetime
import hashlib
import http.cookiejar
import io
import json
import os
import re
import socket
import shutil
import ssl
import sys
import time
import urllib.error
import urllib.parse
import urllib.request
import uuid
import unicodedata
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from typing import Any, Iterable
import xml.etree.ElementTree as ET

try:
    import fitz  # type: ignore[import-not-found]
except ImportError:
    fitz = None

try:
    from openpyxl import load_workbook  # type: ignore[import-not-found]
except ImportError:
    load_workbook = None

try:
    from PIL import Image  # type: ignore[import-not-found]
except ImportError:
    Image = None

try:
    from pypdf import PdfReader  # type: ignore[import-not-found]
except ImportError:
    PdfReader = None

try:
    import pytesseract  # type: ignore[import-not-found]
except ImportError:
    pytesseract = None

try:
    import xlrd  # type: ignore[import-not-found]
except ImportError:
    xlrd = None

try:
    import certifi  # type: ignore[import-not-found]
except ImportError:
    certifi = None

DEFAULT_BASE_URL = "https://api.sos-contador.com/api-comunidad"
DEFAULT_SOFT_BASE_URL = "https://soft.sos-contador.com"
DEFAULT_OFINUBE_ORQUESTAR_URL = "https://ci025qoff5.execute-api.sa-east-1.amazonaws.com/testing/agente/orquestar"
SKILL_ROOT = Path(__file__).resolve().parents[1]


def resolve_runtime_home() -> Path:
    configured = os.environ.get("SOS_CONTADOR_HOME", "").strip()
    candidate = Path(configured).expanduser() if configured else Path.home() / ".sos-contador"
    return candidate.resolve()


RUNTIME_HOME = resolve_runtime_home()
LOCAL_ENV_PATH = RUNTIME_HOME / ".env.local"
LEGACY_LOCAL_ENV_PATH = SKILL_ROOT / ".env.local"
CUIT_CATALOG_CACHE_PATH = RUNTIME_HOME / ".cuit_catalog_cache.json"
CUIT_ALIASES_PATH = RUNTIME_HOME / ".cuit_aliases.json"
DOCUMENT_CATALOG_CACHE_PATH = RUNTIME_HOME / ".document_catalog_cache.json"
DRAFT_CACHE_DIR = RUNTIME_HOME / ".draft_cache"
LOCAL_DATA_DIR = RUNTIME_HOME / "local"
DOCUMENT_PROFILE_ROOT = LOCAL_DATA_DIR / "document_profiles"
PUBLIC_API_CATALOG_PATH = SKILL_ROOT / "references" / "public-api-operations.json"
CUIT_CATALOG_MAX_AGE_SECONDS = 12 * 60 * 60
DOCUMENT_CATALOG_MAX_AGE_SECONDS = 12 * 60 * 60
DRAFT_TTL_SECONDS = 2 * 60 * 60
OCR_ENV_VARS = ("SOS_CONTADOR_TESSERACT_CMD", "TESSERACT_CMD")
CUIT_CHECK_WEIGHTS = (5, 4, 3, 2, 7, 6, 5, 4, 3, 2)
CUIT_ALLOWED_PREFIXES = {"20", "23", "24", "27", "30", "33", "34"}
PROFILE_SCOPE_LINE = "line"
PROFILE_SCOPE_TEXT = "text"
PROFILE_DOCUMENT_KIND_COBRO = "cobro_recibo"
PROFILE_DOCUMENT_DATE_TOKEN = "__DOCUMENT_DATE__"
AFIP_DRAFT_KIND = "afip_mis_comprobantes"
COMPRA_DOCUMENT_DRAFT_KIND = "compra_documento"
AFIP_PURCHASE_PROVIDER_PROVINCE_ID = 19
AFIP_PURCHASE_PROVIDER_COND_IVA_A = 1
AFIP_PURCHASE_PROVIDER_COND_IVA_C = 3
LEGAL_ENTITY_SUFFIXES = {
    "sa",
    "s a",
    "srl",
    "s r l",
    "sas",
    "s a s",
    "srlu",
    "s r l u",
    "srlu",
    "sociedad anonima",
    "sociedad de responsabilidad limitada",
    "sociedad por acciones simplificada",
}
MUTATING_METHODS = {"POST", "PUT", "DELETE", "PATCH"}
COMPROBANTE_MUTATION_PATH_HINTS = (
    "compra",
    "venta",
    "cobro",
    "pago",
    "afip/import",
    "back/comprobante_altamodi.asp",
)
COMPROBANTE_CANCELLATION_KEY_HINTS = (
    "cancelado",
    "anulado",
    "anular",
    "fechabaja",
    "fecha_baja",
    "baja",
    "eliminadodefinitivo",
)
COMPROBANTE_CANCELLATION_VALUE_HINTS = {
    "anulado",
    "anular",
    "cancelado",
    "cancelar",
    "dar de baja",
    "baja",
    "eliminado",
}
TOKEN_KEYS_JWT = ("jwt", "token", "access_token", "accessToken", "bearer")
TOKEN_KEYS_JWTC = ("jwtc", "jwt", "token", "access_token", "accessToken", "bearer")
SENSITIVE_PREVIEW_KEYS = {
    "authorization",
    "bearer",
    "clave",
    "jwt",
    "jwtc",
    "password",
    "secret",
    "token",
}
HTTP_RETRY_ATTEMPTS = 3
HTTP_RETRY_SLEEP_SECONDS = 0.5
DEPRECATED_WORK_CUIT_ENV_KEYS = (
    "SOS_CONTADOR_CUIT_ID",
    "SOS_CUIT_ID",
    "SOS_CONTADOR_CUIT",
    "SOS_CUIT",
)
WORK_CUIT_POSITIVE_HINTS = (
    "datos del comprador",
    "comprador",
    "adquirente",
    "receptor",
    "titular",
    "datos del receptor",
    "razon social comprador",
)
WORK_CUIT_NEGATIVE_HINTS = (
    "datos del vendedor",
    "vendedor",
    "emisor",
    "proveedor",
    "cuit emisor",
    "datos del proveedor",
)


class CLIError(RuntimeError):
    pass


def build_api_ssl_context() -> ssl.SSLContext:
    """Build a verified TLS context without falling back to insecure HTTPS."""
    configured_bundle = first_env("SOS_CONTADOR_CA_BUNDLE", "SSL_CERT_FILE")
    if configured_bundle:
        bundle_path = Path(configured_bundle).expanduser()
        if not bundle_path.is_file():
            raise CLIError(f"El bundle CA configurado no existe o no es un archivo: {bundle_path}")
        try:
            return ssl.create_default_context(cafile=str(bundle_path))
        except (OSError, ssl.SSLError) as exc:
            raise CLIError(f"No se pudo cargar el bundle CA configurado: {bundle_path}") from exc

    if certifi is not None:
        return ssl.create_default_context(cafile=certifi.where())
    return ssl.create_default_context()


def load_local_env_file(path: Path) -> bool:
    if not path.exists():
        return False
    for raw_line in path.read_text(encoding="utf-8").splitlines():
        line = raw_line.strip()
        if not line or line.startswith("#"):
            continue
        if line.startswith("export "):
            line = line[7:].strip()
        if "=" not in line:
            continue
        key, value = line.split("=", 1)
        key = key.strip()
        value = value.strip()
        if not key or key in os.environ:
            continue
        if len(value) >= 2 and value[0] == value[-1] and value[0] in {"'", '"'}:
            value = value[1:-1]
        os.environ[key] = value
    return True


LOADED_LOCAL_ENV_PATH: Path | None = None
if load_local_env_file(LOCAL_ENV_PATH):
    LOADED_LOCAL_ENV_PATH = LOCAL_ENV_PATH
elif LEGACY_LOCAL_ENV_PATH != LOCAL_ENV_PATH and load_local_env_file(LEGACY_LOCAL_ENV_PATH):
    LOADED_LOCAL_ENV_PATH = LEGACY_LOCAL_ENV_PATH
LOCAL_ENV_LOADED = LOADED_LOCAL_ENV_PATH is not None


def load_json_file(path: Path, default: Any) -> Any:
    if not path.exists():
        return default
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return default


def load_public_api_catalog() -> dict[str, Any]:
    try:
        payload = json.loads(PUBLIC_API_CATALOG_PATH.read_text(encoding="utf-8"))
    except FileNotFoundError as exc:
        raise CLIError(f"No existe el catalogo de API publica: {PUBLIC_API_CATALOG_PATH}") from exc
    except json.JSONDecodeError as exc:
        raise CLIError(f"El catalogo de API publica no es JSON valido: {exc}") from exc

    operations = payload.get("operations") if isinstance(payload, dict) else None
    if not isinstance(operations, list) or not operations:
        raise CLIError("El catalogo de API publica no contiene operaciones.")

    seen_ids: set[str] = set()
    for operation in operations:
        if not isinstance(operation, dict):
            raise CLIError("El catalogo de API publica contiene una operacion invalida.")
        operation_id = str(operation.get("id") or "").strip()
        method = str(operation.get("method") or "").upper()
        raw_path = str(operation.get("path") or "").strip()
        path = raw_path.strip("/")
        auth_mode = str(operation.get("auth") or "").lower()
        body_mode = str(operation.get("body") or "none").lower()
        query_keys = operation.get("query", [])
        side_effect = operation.get("side_effect")
        if not operation_id or operation_id in seen_ids:
            raise CLIError(f"ID de operacion ausente o duplicado en el catalogo: {operation_id or '(vacio)'}")
        if method not in {"GET", "POST", "PUT", "DELETE", "PATCH"}:
            raise CLIError(f"Metodo invalido para {operation_id}: {method}")
        if not path:
            raise CLIError(f"Path ausente para {operation_id}.")
        if "://" in raw_path or raw_path.startswith(("/", "\\")) or "\\" in raw_path:
            raise CLIError(f"Path no relativo o inseguro para {operation_id}: {raw_path}")
        if auth_mode not in {"none", "jwt", "jwtc"}:
            raise CLIError(f"Modo de autenticacion invalido para {operation_id}: {auth_mode}")
        if body_mode not in {"none", "optional", "required"}:
            raise CLIError(f"Modo de body invalido para {operation_id}: {body_mode}")
        if not isinstance(query_keys, list) or any(not isinstance(key, str) or not key for key in query_keys):
            raise CLIError(f"Lista de query invalida para {operation_id}.")
        if len(query_keys) != len(set(query_keys)):
            raise CLIError(f"La operacion {operation_id} contiene queries duplicadas.")
        if not isinstance(side_effect, bool):
            raise CLIError(f"side_effect debe ser booleano para {operation_id}.")
        if method in {"PUT", "DELETE", "PATCH"} and not side_effect:
            raise CLIError(f"La operacion mutante {operation_id} no puede declarar side_effect=false.")
        if method == "GET" and side_effect:
            raise CLIError(f"La operacion GET {operation_id} no puede declarar side_effect=true.")
        if "invokable" in operation and not isinstance(operation["invokable"], bool):
            raise CLIError(f"invokable debe ser booleano para {operation_id}.")
        seen_ids.add(operation_id)
    return payload


def public_api_operations() -> list[dict[str, Any]]:
    return list(load_public_api_catalog()["operations"])


def public_api_operation(operation_id: str) -> dict[str, Any]:
    normalized = str(operation_id or "").strip().lower()
    for operation in public_api_operations():
        if str(operation.get("id") or "").lower() == normalized:
            return operation
    raise CLIError(
        f"Operacion documentada desconocida: {operation_id}. "
        "Use 'api catalog' para consultar los IDs disponibles."
    )


def save_json_file(path: Path, payload: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2, sort_keys=True), encoding="utf-8")


def file_exists(path: str | None) -> bool:
    if not path:
        return False
    try:
        return Path(path).exists()
    except OSError:
        return False


def get_tesseract_command() -> str | None:
    explicit = first_env(*OCR_ENV_VARS)
    if explicit:
        candidate = str(explicit).strip()
        if candidate:
            if pytesseract is not None:
                pytesseract.pytesseract.tesseract_cmd = candidate
            return candidate
    discovered = shutil.which("tesseract")
    if discovered and pytesseract is not None:
        pytesseract.pytesseract.tesseract_cmd = discovered
    return discovered


def ocr_is_available() -> bool:
    if pytesseract is None or Image is None:
        return False
    command = get_tesseract_command()
    return bool(command and (file_exists(command) or shutil.which(command)))


def ensure_bound_client(client: SOSContadorClient, *, cuit: str | None = None, cuit_id: str | None = None) -> SOSContadorClient:
    normalized_cuit = digits_only(cuit or "") or None
    normalized_id = str(cuit_id).strip() if cuit_id else None
    if normalized_cuit and getattr(client, "explicit_cuit", None) == normalized_cuit:
        return client
    if normalized_id and getattr(client, "explicit_cuit_id", None) == normalized_id:
        return client
    if not normalized_cuit and not normalized_id:
        return client
    return SOSContadorClient(
        base_url=client.base_url,
        explicit_cuit_id=normalized_id,
        explicit_cuit=normalized_cuit,
    )


def load_document_catalog_cache() -> dict[str, Any]:
    payload = load_json_file(DOCUMENT_CATALOG_CACHE_PATH, {})
    if not isinstance(payload, dict):
        return {}
    return payload


def save_document_catalog_cache(payload: dict[str, Any]) -> None:
    save_json_file(DOCUMENT_CATALOG_CACHE_PATH, payload)


def catalog_cache_key(kind: str, cuit: str | None) -> str:
    return f"{kind}:{digits_only(cuit or '') or 'global'}"


def get_cached_document_catalog(kind: str, *, cuit: str | None = None) -> list[dict[str, Any]] | None:
    payload = load_document_catalog_cache()
    entry = payload.get(catalog_cache_key(kind, cuit))
    if not isinstance(entry, dict):
        return None
    fetched_at = parse_iso_datetime(entry.get("fetched_at"))
    if fetched_at is None:
        return None
    age = datetime.datetime.now(datetime.timezone.utc) - fetched_at
    if age.total_seconds() > DOCUMENT_CATALOG_MAX_AGE_SECONDS:
        return None
    items = entry.get("items")
    if not isinstance(items, list):
        return None
    return [item for item in items if isinstance(item, dict)]


def set_cached_document_catalog(kind: str, items: list[dict[str, Any]], *, cuit: str | None = None) -> None:
    payload = load_document_catalog_cache()
    payload[catalog_cache_key(kind, cuit)] = {
        "fetched_at": utc_now_iso(),
        "items": items,
    }
    save_document_catalog_cache(payload)


def cleanup_expired_drafts() -> None:
    if not DRAFT_CACHE_DIR.exists():
        return
    threshold = time.time() - DRAFT_TTL_SECONDS
    for path in DRAFT_CACHE_DIR.glob("*.json"):
        try:
            if path.stat().st_mtime < threshold:
                path.unlink()
        except OSError:
            continue


def draft_file_path(draft_id: str) -> Path:
    return DRAFT_CACHE_DIR / f"{draft_id}.json"


def save_draft_payload(payload: dict[str, Any]) -> Path:
    cleanup_expired_drafts()
    draft_id = str(payload.get("draft_id") or "").strip() or uuid.uuid4().hex[:12]
    payload["draft_id"] = draft_id
    payload.setdefault("created_at", utc_now_iso())
    destination = draft_file_path(draft_id)
    save_json_file(destination, payload)
    return destination


def load_draft_payload(*, draft_id: str | None = None, draft_file: str | None = None) -> dict[str, Any]:
    cleanup_expired_drafts()
    if draft_id and draft_file:
        raise CLIError("Use solo una de estas opciones: --draft-id o --draft-file.")
    if draft_file:
        payload = load_json_file(Path(draft_file), None)
    elif draft_id:
        payload = load_json_file(draft_file_path(draft_id), None)
    else:
        payload = None
    if not isinstance(payload, dict):
        raise CLIError("No se pudo cargar el borrador indicado.")
    created_at = parse_iso_datetime(payload.get("created_at"))
    if created_at is not None:
        age = datetime.datetime.now(datetime.timezone.utc) - created_at
        if age.total_seconds() > DRAFT_TTL_SECONDS:
            raise CLIError("El borrador indicado ya vencio. Genere uno nuevo.")
    return payload


def slugify(value: str | None) -> str:
    normalized = normalize_search_text(value)
    if not normalized:
        return "perfil"
    return normalized.replace(" ", "-")


def profile_counterparty_key(cuit_value: str | None) -> str:
    return digits_only(cuit_value or "") or "_unknown"


def profile_directory(
    *,
    work_cuit: str,
    counterparty_cuit: str | None,
    document_kind: str,
    profile_id: str,
) -> Path:
    return DOCUMENT_PROFILE_ROOT / digits_only(work_cuit) / profile_counterparty_key(counterparty_cuit) / document_kind / profile_id


def normalize_lines_corpus(lines: Iterable[str]) -> str:
    return "\n".join(str(line) for line in lines if str(line).strip())


def load_profile_json(path: Path) -> dict[str, Any]:
    payload = load_json_file(path, {})
    if not isinstance(payload, dict):
        raise CLIError(f"Perfil local invalido: {path}")
    return payload


def make_status(value: Any, state: str, *, source_text: str | None = None, warning: str | None = None) -> dict[str, Any]:
    payload = {"value": value, "estado": state}
    if source_text:
        payload["source_text"] = source_text
    if warning:
        payload["warning"] = warning
    return payload


def normalize_search_text(value: str | None) -> str:
    text = strip_accents(str(value or "")).lower()
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return " ".join(text.split())


def smart_decimal_from_text(value: Any) -> Decimal:
    text = str(value or "").strip()
    if not text:
        raise CLIError("Monto vacio.")
    cleaned = re.sub(r"[^\d,.\-]", "", text)
    if not cleaned:
        raise CLIError(f"Monto invalido '{value}'.")
    if cleaned.count(",") and cleaned.count("."):
        if cleaned.rfind(",") > cleaned.rfind("."):
            cleaned = cleaned.replace(".", "").replace(",", ".")
        else:
            cleaned = cleaned.replace(",", "")
    elif cleaned.count(",") == 1 and cleaned.count(".") == 0:
        cleaned = cleaned.replace(",", ".")
    elif cleaned.count(".") > 1 and cleaned.count(",") == 0:
        cleaned = cleaned.replace(".", "")
    try:
        return Decimal(cleaned)
    except (InvalidOperation, ValueError) as exc:
        raise CLIError(f"Monto invalido '{value}'.") from exc


def parse_optional_decimal_text(value: Any) -> str:
    if value in (None, ""):
        return ""
    return format_decimal_string(smart_decimal_from_text(value))


def guess_date_from_text(value: str | None) -> str | None:
    text = str(value or "").strip()
    if not text:
        return None
    for fmt in ("%d/%m/%Y", "%d/%m/%y", "%Y-%m-%d", "%d-%m-%Y", "%d-%m-%y"):
        try:
            return datetime.datetime.strptime(text, fmt).strftime("%Y-%m-%d")
        except ValueError:
            continue
    return None


def normalize_comprobante_reference(value: str | None) -> str | None:
    text = str(value or "").strip().upper()
    if not text:
        return None
    match = re.search(r"\b([A-Z]{1,2})\s*[- ]?\s*(\d{1,5})\s*[- ]\s*(\d{1,8})\b", text)
    if not match:
        match = re.search(r"\b(?:FC|FACTURA|NC|ND|RECIBO|COMPROBANTE)?\s*([A-Z]{1,2})\s+(\d{1,5})[-/ ](\d{1,8})\b", text)
    if not match:
        return None
    letra, sucursal, numero = match.groups()
    return f"{letra}-{int(digits_only(sucursal) or '0'):04d}-{int(digits_only(numero) or '0'):08d}"


def is_sale_comprobante(reference: str | None) -> bool:
    if not reference:
        return False
    prefix = str(reference).split("-", 1)[0].upper()
    return prefix in {"A", "B", "C", "E", "M", "NC", "ND"}


def utc_now_iso() -> str:
    return datetime.datetime.now(datetime.timezone.utc).isoformat()


def parse_iso_datetime(value: Any) -> datetime.datetime | None:
    if not isinstance(value, str) or not value.strip():
        return None
    text = value.strip()
    try:
        return datetime.datetime.fromisoformat(text)
    except ValueError:
        return None


def strip_accents(value: str) -> str:
    normalized = unicodedata.normalize("NFKD", value)
    return "".join(ch for ch in normalized if not unicodedata.combining(ch))


def normalize_entity_name(value: str) -> str:
    text = strip_accents(value).lower()
    text = re.sub(r"[^a-z0-9\s]", " ", text)
    text = " ".join(text.split())
    if not text:
        return ""

    parts = text.split()
    while parts:
        candidate = " ".join(parts[-3:])
        if candidate in LEGAL_ENTITY_SUFFIXES:
            del parts[-3:]
            continue
        candidate = " ".join(parts[-2:])
        if candidate in LEGAL_ENTITY_SUFFIXES:
            del parts[-2:]
            continue
        if parts[-1] in LEGAL_ENTITY_SUFFIXES:
            parts.pop()
            continue
        break
    return " ".join(parts)


def catalog_entry_tokens(entry: dict[str, Any]) -> set[str]:
    normalized = str(entry.get("normalized_name") or "")
    return {token for token in normalized.split() if token}


def load_cuit_aliases() -> dict[str, str]:
    payload = load_json_file(CUIT_ALIASES_PATH, {})
    if not isinstance(payload, dict):
        return {}
    aliases: dict[str, str] = {}
    for key, value in payload.items():
        if not isinstance(key, str) or not isinstance(value, str):
            continue
        normalized = normalize_entity_name(key)
        target = digits_only(value)
        if normalized and target:
            aliases[normalized] = target
    return aliases


def load_cuit_catalog_cache() -> list[dict[str, Any]]:
    payload = load_json_file(CUIT_CATALOG_CACHE_PATH, [])
    if not isinstance(payload, list):
        return []
    items: list[dict[str, Any]] = []
    for item in payload:
        if not isinstance(item, dict):
            continue
        cuit = digits_only(str(item.get("cuit") or ""))
        cuit_id = str(item.get("id") or "").strip()
        razon_social = str(item.get("razon_social") or "").strip()
        normalized_name = normalize_entity_name(str(item.get("normalized_name") or razon_social))
        if not cuit or not cuit_id:
            continue
        items.append(
            {
                "id": cuit_id,
                "cuit": cuit,
                "razon_social": razon_social or None,
                "normalized_name": normalized_name,
                "fetched_at": item.get("fetched_at"),
            }
        )
    return items


def is_cuit_catalog_stale(entries: list[dict[str, Any]]) -> bool:
    if not entries:
        return True
    fetched_values = [parse_iso_datetime(entry.get("fetched_at")) for entry in entries]
    fetched_values = [value for value in fetched_values if value is not None]
    if not fetched_values:
        return True
    freshest = max(fetched_values)
    age = datetime.datetime.now(datetime.timezone.utc) - freshest
    return age.total_seconds() > CUIT_CATALOG_MAX_AGE_SECONDS


def build_cuit_catalog_entries(payload: Any) -> list[dict[str, Any]]:
    fetched_at = utc_now_iso()
    entries: list[dict[str, Any]] = []
    for item in collect_cuit_entries(payload):
        cuit = digits_only(item.get("cuit") or "")
        cuit_id = str(item.get("id") or "").strip()
        razon_social = str(item.get("razon_social") or "").strip()
        if not cuit or not cuit_id:
            continue
        entries.append(
            {
                "id": cuit_id,
                "cuit": cuit,
                "razon_social": razon_social or None,
                "normalized_name": normalize_entity_name(razon_social),
                "fetched_at": fetched_at,
            }
        )
    return entries


def refresh_cuit_catalog(client: "SOSContadorClient") -> list[dict[str, Any]]:
    session = client.get_login_session()
    entries = build_cuit_catalog_entries(session.login_payload)
    save_json_file(CUIT_CATALOG_CACHE_PATH, entries)
    return entries


def format_cuit_catalog_item(item: dict[str, Any]) -> dict[str, Any]:
    result = {
        "id": item.get("id"),
        "cuit": item.get("cuit"),
        "razon_social": item.get("razon_social"),
    }
    if item.get("normalized_name"):
        result["normalized_name"] = item["normalized_name"]
    return result


def resolve_cuit_name(
    client: "SOSContadorClient",
    *,
    name: str,
    refresh: bool = False,
) -> dict[str, Any]:
    normalized_query = normalize_entity_name(name)
    if not normalized_query:
        raise CLIError("El nombre o alias de CUIT no puede estar vacio.")

    aliases = load_cuit_aliases()
    catalog = refresh_cuit_catalog(client) if refresh else load_cuit_catalog_cache()
    refreshed = refresh

    def ensure_catalog() -> list[dict[str, Any]]:
        nonlocal catalog, refreshed
        if not catalog:
            catalog = refresh_cuit_catalog(client)
            refreshed = True
        return catalog

    def find_exact_matches(entries: list[dict[str, Any]]) -> list[dict[str, Any]]:
        if digits_only(normalized_query) and len(digits_only(normalized_query)) >= 11:
            return [entry for entry in entries if digits_only(str(entry.get("cuit") or "")) == digits_only(normalized_query)]
        return [entry for entry in entries if entry.get("normalized_name") == normalized_query]

    def find_token_matches(entries: list[dict[str, Any]]) -> list[dict[str, Any]]:
        query_tokens = {token for token in normalized_query.split() if token}
        if not query_tokens:
            return []
        matches: list[dict[str, Any]] = []
        for entry in entries:
            if query_tokens.issubset(catalog_entry_tokens(entry)):
                matches.append(entry)
        return matches

    if normalized_query in aliases:
        target_cuit = aliases[normalized_query]
        entries = ensure_catalog()
        alias_matches = [entry for entry in entries if digits_only(str(entry.get("cuit") or "")) == target_cuit]
        if len(alias_matches) == 1:
            result = format_cuit_catalog_item(alias_matches[0])
            result["matched_by"] = "alias"
            result["query"] = name
            return result

    entries = refresh_cuit_catalog(client) if is_cuit_catalog_stale(catalog) and not refreshed else ensure_catalog()
    if entries is not catalog:
        catalog = entries
        refreshed = True
    exact_matches = find_exact_matches(entries)
    if len(exact_matches) == 1:
        result = format_cuit_catalog_item(exact_matches[0])
        result["matched_by"] = "catalog_exact"
        result["query"] = name
        return result
    if len(exact_matches) > 1:
        raise CLIError(
            "La referencia de CUIT es ambigua. Especifique mejor el contribuyente.\n"
            f"{format_output([format_cuit_catalog_item(item) for item in exact_matches])}"
        )

    token_matches = find_token_matches(entries)
    if len(token_matches) == 1:
        result = format_cuit_catalog_item(token_matches[0])
        result["matched_by"] = "catalog_tokens"
        result["query"] = name
        return result
    if len(token_matches) > 1:
        raise CLIError(
            "La referencia de CUIT es ambigua. Especifique mejor el contribuyente.\n"
            f"{format_output([format_cuit_catalog_item(item) for item in token_matches])}"
        )

    if not refreshed:
        return resolve_cuit_name(client, name=name, refresh=True)

    raise CLIError(
        "No se pudo resolver la CUIT a partir del nombre o alias indicado. "
        "Revise los aliases locales o la lista de CUITs accesibles."
    )


def deprecated_work_cuit_env_values() -> dict[str, str]:
    values: dict[str, str] = {}
    for name in DEPRECATED_WORK_CUIT_ENV_KEYS:
        value = os.getenv(name)
        if value:
            values[name] = value
    return values


def missing_work_cuit_error() -> str:
    detail = (
        "Debe indicar el CUIT de trabajo con --cuit-trabajo, --cuit-trabajo-id o --cuit-trabajo-nombre. "
        "Las variables SOS_CONTADOR_CUIT*, SOS_CUIT* y sus aliases locales fueron deprecadas y ya no se usan "
        "para seleccionar el contribuyente operativo."
    )
    deprecated = deprecated_work_cuit_env_values()
    if deprecated:
        detail += f" Variables detectadas e ignoradas: {', '.join(sorted(deprecated.keys()))}."
    return detail


def ensure_cuit_catalog(client: "SOSContadorClient", *, refresh: bool = False) -> list[dict[str, Any]]:
    if refresh:
        return refresh_cuit_catalog(client)
    cached = load_cuit_catalog_cache()
    if not cached or is_cuit_catalog_stale(cached):
        return refresh_cuit_catalog(client)
    return cached


def resolve_work_cuit_target(
    client: "SOSContadorClient",
    *,
    explicit_id: str | None = None,
    explicit_cuit: str | None = None,
    explicit_name: str | None = None,
    required: bool = True,
) -> dict[str, Any] | None:
    selectors = {
        "id": str(explicit_id or "").strip(),
        "cuit": digits_only(str(explicit_cuit or "")),
        "nombre": str(explicit_name or "").strip(),
    }
    used = [key for key, value in selectors.items() if value]
    if len(used) > 1:
        raise CLIError(
            "Indique el CUIT de trabajo de una sola forma: --cuit-trabajo, --cuit-trabajo-id o --cuit-trabajo-nombre."
        )
    if not used:
        if required:
            raise CLIError(missing_work_cuit_error())
        return None

    if selectors["nombre"]:
        match = resolve_cuit_name(client, name=selectors["nombre"])
        return {
            "cuit_id": str(match.get("id") or ""),
            "cuit": digits_only(str(match.get("cuit") or "")),
            "nombre": match.get("razon_social"),
            "matched_by": match.get("matched_by"),
        }

    catalog = ensure_cuit_catalog(client)
    if selectors["id"]:
        matches = [item for item in catalog if str(item.get("id") or "") == selectors["id"]]
        if not matches:
            catalog = ensure_cuit_catalog(client, refresh=True)
            matches = [item for item in catalog if str(item.get("id") or "") == selectors["id"]]
        if len(matches) != 1:
            raise CLIError(f"No se pudo resolver la CUIT de trabajo con id {selectors['id']}.")
        item = matches[0]
        return {
            "cuit_id": str(item.get("id") or ""),
            "cuit": digits_only(str(item.get("cuit") or "")),
            "nombre": item.get("razon_social"),
            "matched_by": "id",
        }

    matches = [item for item in catalog if digits_only(str(item.get("cuit") or "")) == selectors["cuit"]]
    if not matches:
        catalog = ensure_cuit_catalog(client, refresh=True)
        matches = [item for item in catalog if digits_only(str(item.get("cuit") or "")) == selectors["cuit"]]
    unique_ids = {str(item.get("id") or "") for item in matches if str(item.get("id") or "").strip()}
    if len(unique_ids) != 1:
        raise CLIError(f"No se pudo resolver la CUIT de trabajo {selectors['cuit']}.")
    item = next(item for item in matches if str(item.get("id") or "").strip() in unique_ids)
    return {
        "cuit_id": str(item.get("id") or ""),
        "cuit": digits_only(str(item.get("cuit") or "")),
        "nombre": item.get("razon_social"),
        "matched_by": "cuit",
    }


def resolve_work_cuit_target_from_args(
    client: "SOSContadorClient",
    args: argparse.Namespace,
    *,
    required: bool = True,
) -> dict[str, Any] | None:
    return resolve_work_cuit_target(
        client,
        explicit_id=getattr(args, "cuit_trabajo_id", None),
        explicit_cuit=getattr(args, "cuit_trabajo", None),
        explicit_name=getattr(args, "cuit_trabajo_nombre", None),
        required=required,
    )


def bind_business_client(
    client: "SOSContadorClient",
    args: argparse.Namespace,
    *,
    required: bool = True,
) -> tuple["SOSContadorClient", dict[str, Any] | None]:
    target = resolve_work_cuit_target_from_args(client, args, required=required)
    if target is None:
        return client, None
    bound = ensure_bound_client(
        client,
        cuit=target.get("cuit"),
        cuit_id=target.get("cuit_id"),
    )
    return bound, target


@dataclass
class Session:
    jwt: str
    jwtc: str
    cuit_id: str
    cuit: str | None
    login_payload: Any
    credentials_payload: Any


@dataclass
class LoginSession:
    jwt: str
    login_payload: Any


@dataclass
class WebSession:
    cuit_id: str
    cuit: str | None
    nombrecuit: str | None
    login_message: str | None
    default_html: str


class APIError(RuntimeError):
    def __init__(self, status: int, reason: str, payload: Any):
        self.status = status
        self.reason = reason
        self.payload = payload
        detail = format_output(payload) if payload is not None else "(no response body)"
        super().__init__(f"HTTP {status} {reason}\n{detail}")


class SOSContadorClient:
    def __init__(
        self,
        base_url: str | None = None,
        *,
        explicit_cuit_id: str | None = None,
        explicit_cuit: str | None = None,
    ):
        self.base_url = (
            base_url
            or first_env("SOS_CONTADOR_BASE_URL", "SOS_API_BASE_URL")
            or DEFAULT_BASE_URL
        ).rstrip("/")
        self._login_session: LoginSession | None = None
        self._session: Session | None = None
        self.ssl_context = build_api_ssl_context()
        self.explicit_cuit_id = str(explicit_cuit_id).strip() if explicit_cuit_id else None
        self.explicit_cuit = digits_only(explicit_cuit) if explicit_cuit else None

    def get_login_session(self) -> LoginSession:
        if self._login_session is not None:
            return self._login_session
        usuario = require_env("SOS_CONTADOR_USUARIO", "SOS_USUARIO")
        password = require_env("SOS_CONTADOR_PASSWORD", "SOS_PASSWORD")

        login_payload, _ = self._send_request(
            method="POST",
            path="login",
            body={"usuario": usuario, "password": password},
        )
        jwt = extract_token(login_payload, TOKEN_KEYS_JWT)
        if not jwt:
            raise CLIError("No se pudo extraer el JWT del login. Revise la respuesta del endpoint /login.")
        self._login_session = LoginSession(jwt=jwt, login_payload=login_payload)
        return self._login_session

    def get_session(self) -> Session:
        if self._session is not None:
            return self._session

        login_session = self.get_login_session()
        explicit_id = self.explicit_cuit_id
        explicit_cuit = self.explicit_cuit
        if explicit_id is None and explicit_cuit is None:
            raise CLIError(missing_work_cuit_error())

        cuit_id, cuit_value = resolve_cuit_from_login(
            login_session.login_payload,
            explicit_id=explicit_id,
            explicit_cuit=explicit_cuit,
            require_explicit=True,
        )

        credentials_payload, _ = self._send_request(
            method="GET",
            path=f"cuit/credentials/{cuit_id}",
            bearer_token=login_session.jwt,
        )
        jwtc = extract_token(credentials_payload, TOKEN_KEYS_JWTC)
        if not jwtc:
            raise CLIError(
                "No se pudo extraer el JWTC de la credencial de CUIT. Revise la respuesta del endpoint /cuit/credentials/:idcuit."
            )

        self._session = Session(
            jwt=login_session.jwt,
            jwtc=jwtc,
            cuit_id=str(cuit_id),
            cuit=cuit_value,
            login_payload=login_session.login_payload,
            credentials_payload=credentials_payload,
        )
        return self._session

    def request(
        self,
        method: str,
        path: str,
        *,
        query: list[tuple[str, str]] | None = None,
        body: Any = None,
        auth_mode: str = "jwtc",
        out_path: str | None = None,
    ) -> Any:
        bearer_token = None
        if auth_mode == "jwt":
            bearer_token = self.get_login_session().jwt
        elif auth_mode == "jwtc":
            bearer_token = self.get_session().jwtc
        elif auth_mode != "none":
            raise CLIError(f"Modo de autenticacion no soportado: {auth_mode}")

        payload, headers = self._send_request(
            method=method,
            path=path,
            query=query,
            body=body,
            bearer_token=bearer_token,
            out_path=out_path,
        )
        if out_path:
            return {"written_to": out_path, "content_type": headers.get("Content-Type", "")}
        return payload

    def request_raw(
        self,
        method: str,
        path: str,
        *,
        query: list[tuple[str, str]] | None = None,
        body: Any = None,
        auth_mode: str = "jwtc",
    ) -> tuple[bytes, dict[str, str]]:
        bearer_token = None
        if auth_mode == "jwt":
            bearer_token = self.get_login_session().jwt
        elif auth_mode == "jwtc":
            bearer_token = self.get_session().jwtc
        elif auth_mode != "none":
            raise CLIError(f"Modo de autenticacion no soportado: {auth_mode}")

        return self._send_request_raw(
            method=method,
            path=path,
            query=query,
            body=body,
            bearer_token=bearer_token,
        )

    def _send_request(
        self,
        *,
        method: str,
        path: str,
        query: list[tuple[str, str]] | None = None,
        body: Any = None,
        bearer_token: str | None = None,
        out_path: str | None = None,
    ) -> tuple[Any, dict[str, str]]:
        raw, response_headers = self._send_request_raw(
            method=method,
            path=path,
            query=query,
            body=body,
            bearer_token=bearer_token,
        )
        if out_path:
            write_binary_output(out_path, raw)
            return None, response_headers

        return parse_response_body(raw, response_headers), response_headers

    def _send_request_raw(
        self,
        *,
        method: str,
        path: str,
        query: list[tuple[str, str]] | None = None,
        body: Any = None,
        bearer_token: str | None = None,
    ) -> tuple[bytes, dict[str, str]]:
        url = build_url(self.base_url, path, query)
        headers = {"Accept": "*/*"}
        data = None
        if body is not None:
            headers["Content-Type"] = "application/json"
            data = json.dumps(body, ensure_ascii=True).encode("utf-8")
        if bearer_token:
            headers["Authorization"] = f"Bearer {bearer_token}"

        request = urllib.request.Request(url=url, data=data, headers=headers, method=method.upper())
        last_error: BaseException | None = None
        for attempt in range(HTTP_RETRY_ATTEMPTS):
            try:
                with urllib.request.urlopen(request, timeout=60, context=self.ssl_context) as response:
                    raw = response.read()
                    response_headers = dict(response.headers.items())
                return raw, response_headers
            except urllib.error.HTTPError as exc:
                raw = exc.read()
                payload = parse_response_body(raw, dict(exc.headers.items()))
                raise APIError(exc.code, exc.reason, payload) from exc
            except (urllib.error.URLError, TimeoutError, socket.timeout) as exc:
                last_error = exc
                reason = getattr(exc, "reason", exc)
                if isinstance(reason, ssl.SSLCertVerificationError):
                    raise CLIError(
                        "No se pudo validar el certificado TLS de SOS Contador. "
                        "Actualice las dependencias del skill o configure SOS_CONTADOR_CA_BUNDLE "
                        "con un archivo PEM confiable; no desactive la verificacion TLS."
                    ) from exc
                if attempt + 1 >= HTTP_RETRY_ATTEMPTS or not is_timeout_exception(exc):
                    raise CLIError(f"No se pudo conectar con SOS Contador API: {reason}") from exc
                time.sleep(HTTP_RETRY_SLEEP_SECONDS * (attempt + 1))
        raise CLIError(f"No se pudo conectar con SOS Contador API: {last_error}")


class SOSContadorWebClient:
    def __init__(
        self,
        api_client: SOSContadorClient | None = None,
        base_url: str | None = None,
        *,
        explicit_cuit_id: str | None = None,
        explicit_cuit: str | None = None,
    ):
        self.api_client = api_client
        self.base_url = (
            base_url
            or first_env("SOS_CONTADOR_SOFT_BASE_URL", "SOS_SOFT_BASE_URL")
            or DEFAULT_SOFT_BASE_URL
        ).rstrip("/")
        self.cookie_jar = http.cookiejar.CookieJar()
        self.opener = urllib.request.build_opener(urllib.request.HTTPCookieProcessor(self.cookie_jar))
        self._session: WebSession | None = None
        self.explicit_cuit_id = str(explicit_cuit_id).strip() if explicit_cuit_id else None
        self.explicit_cuit = digits_only(explicit_cuit) if explicit_cuit else None

    def _reset_http_session(self) -> None:
        self.cookie_jar = http.cookiejar.CookieJar()
        self.opener = urllib.request.build_opener(urllib.request.HTTPCookieProcessor(self.cookie_jar))

    def get_session(self) -> WebSession:
        if self._session is not None:
            return self._session

        usuario = require_env("SOS_CONTADOR_USUARIO", "SOS_USUARIO")
        password = require_env("SOS_CONTADOR_PASSWORD", "SOS_PASSWORD")
        desired_cuit_id = self.explicit_cuit_id
        desired_cuit = self.explicit_cuit
        if self.api_client is not None and (desired_cuit_id is None or desired_cuit is None):
            api_session = self.api_client.get_session()
            desired_cuit_id = desired_cuit_id or api_session.cuit_id
            desired_cuit = desired_cuit or api_session.cuit
        if desired_cuit_id is None and desired_cuit is None:
            raise CLIError(missing_work_cuit_error())

        login_form = {
            "usuario": usuario,
            "clave": password,
            "idcuit": desired_cuit_id or "",
            "fuente": "",
            "referal": "",
            "telefono": "",
        }
        last_error: CLIError | None = None
        login_payload = None
        for attempt in range(3):
            self._reset_http_session()
            login_text, _ = self._request(
                "POST",
                "back/login.asp",
                form=login_form,
            )
            login_payload = parse_items_xml(login_text)
            code = str(login_payload.get("codigo", ""))
            if code != "0":
                last_error = CLIError(
                    "No se pudo iniciar la web-session de SOS Contador. "
                    f"Respuesta de login.asp: {(login_payload or {}).get('mensaje') or login_text}"
                )
                time.sleep(0.25)
                continue

            for _ in range(3):
                default_html = self.get_default_page()
                info = extract_web_session_from_html(default_html)
                if desired_cuit_id and info["idcuit"] != str(desired_cuit_id):
                    self._request("GET", "back/cambiar_cuit.asp", query=[("nuevocuit", str(desired_cuit_id))])
                    default_html = self.get_default_page()
                    info = extract_web_session_from_html(default_html)

                active_matches_desired_cuit = bool(
                    desired_cuit and info["cuit"] and digits_only(info["cuit"]) == digits_only(desired_cuit)
                )
                id_matches = not desired_cuit_id or info["idcuit"] == str(desired_cuit_id)
                cuit_matches = not desired_cuit or not info["cuit"] or digits_only(info["cuit"]) == digits_only(desired_cuit)

                if id_matches and cuit_matches:
                    self._session = WebSession(
                        cuit_id=info["idcuit"] or str(desired_cuit_id or ""),
                        cuit=info["cuit"] or desired_cuit,
                        nombrecuit=info["nombrecuit"],
                        login_message=str(login_payload.get("mensaje", "")) or None,
                        default_html=default_html,
                    )
                    return self._session

                if desired_cuit_id and info["idcuit"] != str(desired_cuit_id) and not active_matches_desired_cuit:
                    last_error = CLIError(
                        "La web-session no quedo posicionada en la CUIT esperada. "
                        f"Esperada: {desired_cuit_id}. Activa: {info['idcuit'] or '(sin idcuit)'}."
                    )
                elif desired_cuit and info["cuit"] and digits_only(info["cuit"]) != digits_only(desired_cuit):
                    last_error = CLIError(
                        "La web-session quedo en una CUIT distinta a la requerida. "
                        f"Esperada: {desired_cuit}. Activa: {info['cuit']}."
                    )
                else:
                    last_error = CLIError("La web-session no devolvio una CUIT activa utilizable.")
                time.sleep(0.25)

            time.sleep(0.25)

        raise last_error or CLIError("No se pudo iniciar la web-session de SOS Contador.")

    def get_default_page(self) -> str:
        text, _ = self._request("GET", "web/default.asp")
        return text

    def list_comprobantes_range(
        self,
        *,
        idtipo_operacion: int,
        desde: str,
        hasta: str,
        start: int = 0,
        length: int = 200,
    ) -> dict[str, Any]:
        self.get_session()
        form = {
            "object": "comprobante_listado",
            "start": str(start),
            "length": str(length),
            "search": "",
            "idtipo_operacion": str(idtipo_operacion),
            "eliminadodefinitivo": "0",
            "order": [{"column": 1, "dir": "asc", "name": "fecha"}],
            "filters": [
                {
                    "name": "fecha",
                    "logic": "and",
                    "0": {"field": "C.fechaiva", "type": "date", "op": "gte", "data": desde},
                    "1": {"field": "C.fechaiva", "type": "date", "op": "lte", "data": hasta},
                }
            ],
        }
        text, _ = self._request("POST", "back/xml.asp", form=form)
        payload = parse_items_xml(text)
        if str(payload.get("codigo", "")) == "-1":
            raise CLIError(f"Error en web-session xml.asp: {payload.get('mensaje') or text}")
        return payload

    def export_comprobantes_csv(self, *, idtipo_operacion: int, desde: str, hasta: str) -> str:
        self.get_session()
        query = [
            ("csv", "1"),
            ("idtipo_operacion", str(idtipo_operacion)),
            ("fechadesde", desde),
            ("fechahasta", hasta),
            ("sucursal", ""),
            ("idclipro", ""),
            ("idtipo_condicioniva", ""),
            ("letra", ""),
            ("idprovinciaiibb", ""),
            ("conadjunto", ""),
            ("Antiguedad", ""),
        ]
        text, _ = self._request("GET", "back/comprobante_listado_xls.asp", query=query)
        return text

    def get_comprobante_historia(self, *, idcomprobante: str, idtipo_operacion: int) -> dict[str, Any]:
        self.get_session()
        form = {
            "object": "comprobante_historia",
            "idcomprobante": str(idcomprobante),
            "idtipo_operacion": str(idtipo_operacion),
        }
        text, _ = self._request("POST", "back/xml.asp", form=form)
        payload = parse_items_xml(text)
        items = payload.get("items", [])
        if not items or not isinstance(items[0], dict):
            raise CLIError("La web-session no devolvio un detalle utilizable para el comprobante solicitado.")
        return items[0]

    def xml_request(self, object_name: str, *, form: dict[str, Any] | None = None) -> dict[str, Any]:
        self.get_session()
        request_form = {"object": object_name}
        if form:
            request_form.update(form)
        text, _ = self._request("POST", "back/xml.asp", form=request_form)
        payload = parse_items_xml(text)
        if str(payload.get("codigo", "")) == "-1":
            raise CLIError(f"Error en web-session xml.asp: {payload.get('mensaje') or text}")
        return payload

    def get_centros_costo(self) -> list[dict[str, Any]]:
        payload = self.xml_request("centrocosto_predice")
        items = payload.get("items", [])
        if isinstance(items, list):
            return [item for item in items if isinstance(item, dict)]
        if isinstance(items, dict):
            return [items]
        return []

    def get_ultimo_comprobante_numero(self, *, idtipo_operacion: int) -> int:
        payload = self.xml_request("comprobante_ultimo_numero", form={"idtipo_operacion": str(idtipo_operacion)})
        numero = find_value(payload, ("numero", "ultimo_numero", "ultimo", "maximo"))
        try:
            return int(digits_only(str(numero or "0")) or "0")
        except (TypeError, ValueError):
            return 0

    def save_comprobante(self, body: dict[str, Any]) -> dict[str, Any]:
        self.get_session()
        text, _ = self._request("POST", "back/comprobante_altamodi.asp", form=body)
        payload = parse_items_xml(text)
        code = str(payload.get("codigo", ""))
        if code and code != "0":
            raise CLIError(payload.get("mensaje") or f"La web-session devolvio codigo {code} al grabar el comprobante.")
        return payload

    def get_cobropago_vs_compraventa(
        self,
        *,
        idclipro: str,
        idtipo_operacion: int,
        items: list[str],
        numero: str | None = None,
        sucursal: str | None = None,
        tipooperacionbuscar: str | None = None,
    ) -> dict[str, Any]:
        form: dict[str, Any] = {
            "idclipro": str(idclipro),
            "idtipo_operacion": str(idtipo_operacion),
            "items": [str(item) for item in items],
        }
        if numero:
            form["numero"] = str(numero)
        if sucursal:
            form["sucursal"] = str(sucursal)
        if tipooperacionbuscar:
            form["tipooperacionbuscar"] = str(tipooperacionbuscar)
        return self.xml_request("cobropago_compraventa", form=form)

    def save_asociaciones(self, body: dict[str, Any]) -> dict[str, Any]:
        self.get_session()
        text, _ = self._request("POST", "back/asociaciones_altamodi.asp", form=body)
        payload = parse_items_xml(text)
        code = str(payload.get("codigo", ""))
        if code and code != "0":
            raise CLIError(payload.get("mensaje") or f"La web-session devolvio codigo {code} al grabar las asociaciones.")
        return payload

    def download_binary(
        self,
        path: str,
        *,
        query: list[tuple[str, str]] | None = None,
        form: dict[str, Any] | None = None,
    ) -> tuple[bytes, dict[str, str]]:
        self.get_session()
        return self._request_raw("GET" if form is None else "POST", path, query=query, form=form)

    def download_venta_pdf(
        self,
        *,
        venta_id: str,
        out_path: str,
    ) -> dict[str, Any]:
        body = self.build_ofinube_print_payload(venta_id=venta_id)
        agent_payload = dict(body)
        agent_payload["endpoint_original"] = DEFAULT_OFINUBE_ORQUESTAR_URL
        response = self.call_ofinube_agent(agent_payload)
        lote_id = extract_ofinube_lote_id(response)
        lote = self.wait_for_ofinube_lote(lote_id)
        remote_file = str(lote.get("url_documento") or "").strip()
        if not remote_file:
            raise CLIError("La web-session no devolvio la URL final del PDF de la venta solicitada.")
        raw, headers = self._download_absolute_binary(remote_file)
        if not response_looks_like_pdf(raw, headers):
            raise CLIError("La web-session no devolvio un PDF valido para la venta solicitada.")
        write_binary_output(out_path, raw)
        return {
            "written_to": out_path,
            "content_type": headers.get("Content-Type", ""),
            "remote_file": remote_file,
        }

    def build_ofinube_print_payload(self, *, venta_id: str, descargar: bool = False) -> dict[str, Any]:
        self.get_session()
        payload = self._request_form_json(
            "POST",
            "back/sos_agente_armar_json_orquestar.asp",
            form={
                "idscomprobantes": str(venta_id),
                "descargar": bool_to_str(descargar),
            },
        )
        if not isinstance(payload, dict):
            raise CLIError("La web-session no devolvio el payload de impresion esperado.")
        return payload

    def call_ofinube_agent(self, body: dict[str, Any]) -> dict[str, Any]:
        self.get_session()
        payload = self._request_json_json("POST", "back/llamar_agente_ofinube.asp", body=body)
        if not isinstance(payload, dict):
            raise CLIError("La web-session no devolvio una respuesta JSON valida del agente de impresion.")
        status_code = str(payload.get("statusCode") or "")
        if status_code and status_code != "200":
            raise CLIError(payload.get("mensaje") or f"El agente de impresion devolvio status {status_code}.")
        return payload

    def wait_for_ofinube_lote(self, lote_id: str, *, attempts: int = 30, sleep_seconds: float = 1.0) -> dict[str, Any]:
        self.get_session()
        last_lote: dict[str, Any] | None = None
        for _ in range(attempts):
            payload = self._request_form_json(
                "POST",
                "back/sos_agente_obtener_lote.asp",
                form={"id_lote": str(lote_id)},
            )
            if not isinstance(payload, dict):
                raise CLIError("La web-session devolvio un lote de impresion invalido.")
            code = str(payload.get("codigo") or "")
            if code and code != "0":
                raise CLIError(payload.get("mensaje") or f"No se pudo consultar el lote {lote_id}.")
            datos = payload.get("datos")
            if not isinstance(datos, list) or not datos or not isinstance(datos[0], dict):
                raise CLIError(f"La web-session no devolvio datos utilizables para el lote {lote_id}.")
            lote = datos[0]
            last_lote = lote
            estado = str(lote.get("estado") or "").strip().upper()
            if estado == "FINALIZADO":
                return lote
            if estado == "ERROR":
                raise CLIError(f"El lote de impresion {lote_id} termino en ERROR.")
            time.sleep(sleep_seconds)
        estado = str((last_lote or {}).get("estado") or "").strip() or "desconocido"
        raise CLIError(f"El lote de impresion {lote_id} no finalizo a tiempo (estado {estado}).")

    def _download_absolute_binary(self, url: str) -> tuple[bytes, dict[str, str]]:
        request = urllib.request.Request(url=url, headers={"Accept": "*/*"}, method="GET")
        last_error: BaseException | None = None
        for attempt in range(HTTP_RETRY_ATTEMPTS):
            try:
                with urllib.request.urlopen(request, timeout=60) as response:
                    raw = response.read()
                    response_headers = dict(response.headers.items())
                return raw, response_headers
            except urllib.error.HTTPError as exc:
                detail = exc.read().decode("utf-8", errors="replace")
                raise CLIError(f"HTTP {exc.code} {exc.reason}\n{detail}") from exc
            except (urllib.error.URLError, TimeoutError, socket.timeout) as exc:
                last_error = exc
                if attempt + 1 >= HTTP_RETRY_ATTEMPTS or not is_timeout_exception(exc):
                    raise CLIError(
                        f"No se pudo descargar el PDF generado por la web-session: {getattr(exc, 'reason', exc)}"
                    ) from exc
                time.sleep(HTTP_RETRY_SLEEP_SECONDS * (attempt + 1))
        raise CLIError(f"No se pudo descargar el PDF generado por la web-session: {last_error}")

    def _request(
        self,
        method: str,
        path: str,
        *,
        query: list[tuple[str, str]] | None = None,
        form: dict[str, Any] | None = None,
    ) -> tuple[str, dict[str, str]]:
        raw, response_headers = self._request_raw(method, path, query=query, form=form)
        return raw.decode("utf-8", errors="replace"), response_headers

    def _request_raw(
        self,
        method: str,
        path: str,
        *,
        query: list[tuple[str, str]] | None = None,
        form: dict[str, Any] | None = None,
    ) -> tuple[bytes, dict[str, str]]:
        url = build_url(self.base_url, path, query)
        headers = {"Accept": "*/*"}
        data = None
        if form is not None:
            headers["Content-Type"] = "application/x-www-form-urlencoded; charset=UTF-8"
            data = encode_form_data(form)
        request = urllib.request.Request(url=url, data=data, headers=headers, method=method.upper())
        last_error: BaseException | None = None
        for attempt in range(HTTP_RETRY_ATTEMPTS):
            try:
                with self.opener.open(request, timeout=60) as response:
                    raw = response.read()
                    response_headers = dict(response.headers.items())
                break
            except urllib.error.HTTPError as exc:
                raw = exc.read()
                detail = raw.decode("utf-8", errors="replace")
                raise CLIError(f"HTTP {exc.code} {exc.reason}\n{detail}") from exc
            except (urllib.error.URLError, TimeoutError, socket.timeout) as exc:
                last_error = exc
                if attempt + 1 >= HTTP_RETRY_ATTEMPTS or not is_timeout_exception(exc):
                    raise CLIError(
                        f"No se pudo conectar con la web-session de SOS Contador: {getattr(exc, 'reason', exc)}"
                    ) from exc
                time.sleep(HTTP_RETRY_SLEEP_SECONDS * (attempt + 1))
        else:
            raise CLIError(f"No se pudo conectar con la web-session de SOS Contador: {last_error}")
        return raw, response_headers

    def _request_form_json(
        self,
        method: str,
        path: str,
        *,
        query: list[tuple[str, str]] | None = None,
        form: dict[str, Any] | None = None,
    ) -> Any:
        raw, headers = self._request_raw(method, path, query=query, form=form)
        return parse_response_body(raw, headers)

    def _request_json_json(
        self,
        method: str,
        path: str,
        *,
        query: list[tuple[str, str]] | None = None,
        body: Any = None,
    ) -> Any:
        raw, headers = self._request_json_raw(method, path, query=query, body=body)
        return parse_response_body(raw, headers)

    def _request_json_raw(
        self,
        method: str,
        path: str,
        *,
        query: list[tuple[str, str]] | None = None,
        body: Any = None,
    ) -> tuple[bytes, dict[str, str]]:
        url = build_url(self.base_url, path, query)
        headers = {
            "Accept": "*/*",
            "Content-Type": "application/json",
        }
        data = json.dumps(body or {}, ensure_ascii=True).encode("utf-8")
        request = urllib.request.Request(url=url, data=data, headers=headers, method=method.upper())
        last_error: BaseException | None = None
        for attempt in range(HTTP_RETRY_ATTEMPTS):
            try:
                with self.opener.open(request, timeout=60) as response:
                    raw = response.read()
                    response_headers = dict(response.headers.items())
                return raw, response_headers
            except urllib.error.HTTPError as exc:
                raw = exc.read()
                detail = raw.decode("utf-8", errors="replace")
                raise CLIError(f"HTTP {exc.code} {exc.reason}\n{detail}") from exc
            except (urllib.error.URLError, TimeoutError, socket.timeout) as exc:
                last_error = exc
                if attempt + 1 >= HTTP_RETRY_ATTEMPTS or not is_timeout_exception(exc):
                    raise CLIError(
                        f"No se pudo conectar con la web-session de SOS Contador: {getattr(exc, 'reason', exc)}"
                    ) from exc
                time.sleep(HTTP_RETRY_SLEEP_SECONDS * (attempt + 1))
        raise CLIError(f"No se pudo conectar con la web-session de SOS Contador: {last_error}")


def first_env(*names: str) -> str | None:
    for name in names:
        value = os.getenv(name)
        if value:
            return value
    return None


def require_env(name: str, *aliases: str) -> str:
    value = first_env(name, *aliases)
    if value:
        return value
    alias_text = f" (aliases: {', '.join(aliases)})" if aliases else ""
    raise CLIError(
        f"Falta la variable de entorno requerida: {name}{alias_text}. "
        f"Definala en el entorno del proceso o en {LOCAL_ENV_PATH}."
    )


def is_timeout_exception(exc: BaseException) -> bool:
    if isinstance(exc, TimeoutError | socket.timeout):
        return True
    if isinstance(exc, urllib.error.URLError):
        reason = exc.reason
        if isinstance(reason, TimeoutError | socket.timeout):
            return True
        return "timed out" in str(reason).lower()
    return "timed out" in str(exc).lower()


def build_url(base_url: str, path: str, query: list[tuple[str, str]] | None = None) -> str:
    if path.startswith("http://") or path.startswith("https://"):
        url = path
    else:
        url = f"{base_url}/{path.lstrip('/')}"
    if query:
        encoded = urllib.parse.urlencode(query, doseq=True)
        separator = "&" if "?" in url else "?"
        url = f"{url}{separator}{encoded}"
    return url


def write_binary_output(path: str, raw: bytes) -> str:
    destination = Path(path)
    destination.parent.mkdir(parents=True, exist_ok=True)
    destination.write_bytes(raw)
    return str(destination)


def encode_form_data(form: dict[str, Any]) -> bytes:
    pairs: list[tuple[str, str]] = []
    for key, value in form.items():
        pairs.extend(flatten_form_pairs(str(key), value))
    return urllib.parse.urlencode(pairs, doseq=True).encode("utf-8")


def flatten_form_pairs(prefix: str, value: Any) -> list[tuple[str, str]]:
    if isinstance(value, dict):
        pairs: list[tuple[str, str]] = []
        for key, inner in value.items():
            pairs.extend(flatten_form_pairs(f"{prefix}[{key}]", inner))
        return pairs
    if isinstance(value, list):
        pairs: list[tuple[str, str]] = []
        for index, inner in enumerate(value):
            pairs.extend(flatten_form_pairs(f"{prefix}[{index}]", inner))
        return pairs
    if value is None:
        return [(prefix, "")]
    return [(prefix, str(value))]


def parse_items_xml(text: str) -> dict[str, Any]:
    try:
        root = ET.fromstring(text)
    except ET.ParseError as exc:
        raise CLIError(f"No se pudo interpretar la respuesta XML de SOS Contador: {exc}") from exc
    result: dict[str, Any] = {}
    items: list[dict[str, Any]] = []
    for child in root:
        if child.tag == "item":
            items.append(xml_element_to_value(child))
        else:
            value = xml_element_to_value(child)
            if child.tag in result:
                existing = result[child.tag]
                if isinstance(existing, list):
                    existing.append(value)
                else:
                    result[child.tag] = [existing, value]
            else:
                result[child.tag] = value
    if items:
        result["items"] = items
    return result


def xml_element_to_value(element: ET.Element) -> Any:
    children = list(element)
    if not children:
        return (element.text or "").strip()

    result: dict[str, Any] = {}
    repeated: dict[str, list[Any]] = {}
    for child in children:
        value = xml_element_to_value(child)
        if child.tag in result:
            repeated.setdefault(child.tag, [result[child.tag]])
            repeated[child.tag].append(value)
            result[child.tag] = repeated[child.tag]
        else:
            result[child.tag] = value
    return result


def extract_web_session_from_html(html: str) -> dict[str, str | None]:
    match = re.search(r"window\.session\s*=\s*\{(?P<body>.*?)\}", html, re.S)
    if not match:
        raise CLIError("No se pudo extraer la sesion activa desde /web/default.asp.")
    body = match.group("body")

    def find(name: str) -> str | None:
        entry = re.search(rf"{re.escape(name)}\s*:\s*\"([^\"]*)\"", body)
        if entry:
            return entry.group(1)
        return None

    return {
        "idcuit": find("idcuit"),
        "cuit": find("session_cuit"),
        "nombrecuit": find("session_nombrecuit"),
    }


def parse_response_body(raw: bytes, headers: dict[str, str]) -> Any:
    if not raw:
        return None
    content_type = headers.get("Content-Type", "").lower()
    if "application/json" in content_type:
        try:
            return json.loads(raw.decode("utf-8"))
        except json.JSONDecodeError:
            return raw.decode("utf-8", errors="replace")
    try:
        text = raw.decode("utf-8")
    except UnicodeDecodeError:
        return raw
    stripped = text.strip()
    if stripped.startswith("{") or stripped.startswith("["):
        try:
            return json.loads(text)
        except json.JSONDecodeError:
            return text
    return text


def response_looks_like_pdf(raw: bytes, headers: dict[str, str]) -> bool:
    content_type = headers.get("Content-Type", "").lower()
    if "application/pdf" in content_type:
        return True
    return raw.lstrip().startswith(b"%PDF-")


def extract_ofinube_lote_id(payload: Any) -> str:
    if not isinstance(payload, dict):
        raise CLIError("La respuesta del agente de impresion no tiene el formato esperado.")
    resultados = payload.get("resultados")
    if not isinstance(resultados, list):
        raise CLIError("La respuesta del agente de impresion no devolvio resultados.")
    for item in resultados:
        if not isinstance(item, dict):
            continue
        impresion = item.get("resultados_impresion")
        if not isinstance(impresion, dict):
            continue
        lote = impresion.get("lote")
        if not isinstance(lote, dict):
            continue
        lote_id = str(lote.get("id") or "").strip()
        if lote_id:
            return lote_id
    raise CLIError("No se pudo obtener el lote de impresion para la venta solicitada.")


def extract_token(payload: Any, preferred_keys: tuple[str, ...]) -> str | None:
    token = search_token_by_keys(payload, preferred_keys)
    if token:
        return sanitize_token(token)
    if isinstance(payload, str):
        return sanitize_token(payload)
    return None


def search_token_by_keys(payload: Any, preferred_keys: tuple[str, ...]) -> str | None:
    if isinstance(payload, dict):
        for preferred in preferred_keys:
            for key, value in payload.items():
                if key.lower() == preferred.lower() and isinstance(value, str):
                    return value
        for value in payload.values():
            token = search_token_by_keys(value, preferred_keys)
            if token:
                return token
    elif isinstance(payload, list):
        for item in payload:
            token = search_token_by_keys(item, preferred_keys)
            if token:
                return token
    return None


def sanitize_token(value: str) -> str | None:
    token = value.strip()
    if token.lower().startswith("bearer "):
        token = token[7:].strip()
    if token and "." in token and " " not in token:
        return token
    if token and " " not in token and len(token) >= 8:
        return token
    return None


def resolve_cuit_from_login(
    payload: Any,
    explicit_id: str | None,
    explicit_cuit: str | None,
    *,
    require_explicit: bool = False,
) -> tuple[str, str | None]:
    candidates = list(find_cuit_candidates(payload))
    by_id = {candidate["id"]: candidate for candidate in candidates if candidate["id"]}
    if explicit_id:
        chosen = by_id.get(str(explicit_id))
        if chosen is None:
            raise CLIError(f"La CUIT de trabajo con id {explicit_id} no esta disponible para este usuario.")
        return str(explicit_id), chosen["cuit"]

    if explicit_cuit:
        target = digits_only(explicit_cuit)
        matches = [candidate for candidate in candidates if digits_only(candidate["cuit"] or "") == target]
        unique_ids = sorted({candidate["id"] for candidate in matches if candidate["id"]})
        if len(unique_ids) == 1:
            match = next(candidate for candidate in matches if candidate["id"] == unique_ids[0])
            return unique_ids[0], match["cuit"]
        if not matches:
            raise CLIError(f"La CUIT de trabajo {target} no esta disponible para este usuario.")
        raise CLIError(
            f"La CUIT de trabajo {target} coincide con multiples entradas accesibles. "
            "Indique --cuit-trabajo-id para elegir una unica coincidencia."
        )

    if require_explicit:
        raise CLIError(missing_work_cuit_error())

    unique_ids = sorted({candidate["id"] for candidate in candidates if candidate["id"]})
    if len(unique_ids) == 1:
        match = next(candidate for candidate in candidates if candidate["id"] == unique_ids[0])
        return unique_ids[0], match["cuit"]
    raise CLIError(missing_work_cuit_error())


def find_cuit_candidates(payload: Any) -> Iterable[dict[str, str | None]]:
    for item in iter_dicts(payload):
        cuit_key = find_key(item, ("cuit",))
        id_key = find_key(item, ("idcuit", "id_cuit", "id"))
        if not cuit_key or not id_key:
            continue
        cuit_value = item.get(cuit_key)
        id_value = item.get(id_key)
        if cuit_value is None or id_value is None:
            continue
        normalized = digits_only(str(cuit_value))
        if len(normalized) < 8:
            continue
        razon_social_key = find_key(item, ("razonsocial", "razon_social", "razon social", "razon", "nombre"))
        razon_social = None
        if razon_social_key:
            raw_name = item.get(razon_social_key)
            if raw_name is not None:
                text = str(raw_name).strip()
                razon_social = text or None
        yield {"id": str(id_value), "cuit": str(cuit_value), "razon_social": razon_social}


def collect_cuit_entries(payload: Any) -> list[dict[str, str | None]]:
    unique: dict[tuple[str | None, str | None], dict[str, str | None]] = {}
    for candidate in find_cuit_candidates(payload):
        key = (candidate.get("id"), candidate.get("cuit"))
        current = unique.get(key)
        if current is None:
            unique[key] = dict(candidate)
            continue
        if not current.get("razon_social") and candidate.get("razon_social"):
            current["razon_social"] = candidate["razon_social"]
    return sorted(
        unique.values(),
        key=lambda item: (digits_only(item.get("cuit") or ""), item.get("id") or ""),
    )


def iter_dicts(payload: Any) -> Iterable[dict[str, Any]]:
    if isinstance(payload, dict):
        yield payload
        for value in payload.values():
            yield from iter_dicts(value)
    elif isinstance(payload, list):
        for item in payload:
            yield from iter_dicts(item)


def find_key(mapping: dict[str, Any], aliases: tuple[str, ...]) -> str | None:
    lowered = {key.lower(): key for key in mapping.keys()}
    for alias in aliases:
        if alias.lower() in lowered:
            return lowered[alias.lower()]
    return None


def digits_only(value: str) -> str:
    return "".join(ch for ch in value if ch.isdigit())


def is_valid_argentina_cuit(value: Any) -> bool:
    digits = digits_only(str(value or ""))
    if (
        len(digits) != 11
        or len(set(digits)) == 1
        or digits[:2] not in CUIT_ALLOWED_PREFIXES
    ):
        return False
    total = sum(int(digit) * weight for digit, weight in zip(digits[:10], CUIT_CHECK_WEIGHTS))
    expected = 11 - (total % 11)
    if expected == 11:
        expected = 0
    elif expected == 10:
        expected = 9
    return int(digits[-1]) == expected


def require_valid_argentina_cuit(value: Any, *, label: str = "CUIT") -> str:
    digits = digits_only(str(value or ""))
    if not is_valid_argentina_cuit(digits):
        raise CLIError(
            f"{label} inválido: debe tener 11 dígitos y un dígito verificador correcto."
        )
    return digits


def parse_key_value(value: str) -> tuple[str, str]:
    if "=" not in value:
        raise CLIError(f"Parametro invalido '{value}'. Use formato clave=valor.")
    key, raw = value.split("=", 1)
    if not key:
        raise CLIError(f"Parametro invalido '{value}'. La clave no puede estar vacia.")
    return key, raw


def parse_unique_key_values(values: Iterable[str], *, label: str) -> dict[str, str]:
    result: dict[str, str] = {}
    for value in values:
        key, raw = parse_key_value(value)
        if key in result:
            raise CLIError(f"{label} duplicado: {key}")
        result[key] = raw
    return result


def operation_path_parameter_names(path_template: str) -> dict[str, bool]:
    return {
        match.group("name"): bool(match.group("optional"))
        for match in re.finditer(r":(?P<name>[A-Za-z_][A-Za-z0-9_]*)(?P<optional>\?)?", path_template)
    }


def render_operation_path(path_template: str, parameters: dict[str, str]) -> str:
    expected = operation_path_parameter_names(path_template)
    unexpected = sorted(set(parameters) - set(expected))
    if unexpected:
        raise CLIError(f"Parametros de path no admitidos: {', '.join(unexpected)}")

    def replace(match: re.Match[str]) -> str:
        name = match.group("name")
        optional = bool(match.group("optional"))
        value = str(parameters.get(name) or "").strip()
        if not value:
            if optional:
                return ""
            raise CLIError(f"Falta --param {name}=<valor> para completar el path.")
        return urllib.parse.quote(value, safe="")

    rendered = re.sub(
        r":(?P<name>[A-Za-z_][A-Za-z0-9_]*)(?P<optional>\?)?",
        replace,
        path_template.strip().strip("/"),
    )
    rendered = re.sub(r"/{2,}", "/", rendered).strip("/")
    if not rendered:
        raise CLIError("El path resultante no puede quedar vacio.")
    return rendered


def validate_operation_query(operation: dict[str, Any], query: list[tuple[str, str]]) -> None:
    allowed = {str(key) for key in operation.get("query", [])}
    unexpected = sorted({key for key, _ in query if key not in allowed})
    if unexpected:
        raise CLIError(
            f"Query no admitida para {operation['id']}: {', '.join(unexpected)}. "
            f"Admitidas: {', '.join(sorted(allowed)) or '(ninguna)'}."
        )


def validate_operation_body(operation: dict[str, Any], body: Any) -> None:
    body_mode = str(operation.get("body") or "none")
    if body_mode == "required" and body is None:
        raise CLIError(f"{operation['id']} requiere --body-json o --body-file.")
    if body_mode == "none" and body is not None:
        raise CLIError(f"{operation['id']} no admite un body.")


def redact_sensitive_payload(payload: Any) -> Any:
    if isinstance(payload, dict):
        result: dict[str, Any] = {}
        for raw_key, value in payload.items():
            key = str(raw_key)
            normalized = normalize_search_text(key).replace(" ", "_")
            if normalized in SENSITIVE_PREVIEW_KEYS or any(
                hint in normalized for hint in ("password", "secret", "token")
            ):
                result[key] = "<redacted>"
            else:
                result[key] = redact_sensitive_payload(value)
        return result
    if isinstance(payload, list):
        return [redact_sensitive_payload(item) for item in payload]
    return payload


def load_json_source(raw_json: str | None, file_path: str | None) -> Any:
    if raw_json and file_path:
        raise CLIError("Use solo una de estas opciones: --body-json o --body-file.")
    if raw_json:
        try:
            return json.loads(raw_json)
        except json.JSONDecodeError as exc:
            raise CLIError(f"JSON invalido en --body-json: {exc}") from exc
    if file_path:
        try:
            return json.loads(Path(file_path).read_text(encoding="utf-8"))
        except FileNotFoundError as exc:
            raise CLIError(f"No existe el archivo JSON: {file_path}") from exc
        except json.JSONDecodeError as exc:
            raise CLIError(f"JSON invalido en archivo {file_path}: {exc}") from exc
    return None


def bool_to_str(value: bool) -> str:
    return "true" if value else "false"


def format_output(payload: Any) -> str:
    if payload is None:
        return "null"
    if isinstance(payload, (dict, list)):
        return json.dumps(payload, ensure_ascii=False, indent=2, sort_keys=True)
    if isinstance(payload, bytes):
        return f"<{len(payload)} bytes binarios>"
    return str(payload)


def redacted_preview(method: str, path: str, query: list[tuple[str, str]] | None, body: Any, auth_mode: str) -> dict[str, Any]:
    return {
        "method": method.upper(),
        "path": path,
        "query": query or [],
        "auth_mode": auth_mode,
        "body": redact_sensitive_payload(body),
    }


def set_mutation_preview(
    args: argparse.Namespace,
    *,
    business_preview: dict[str, Any] | None = None,
    preview_markdown: str | None = None,
) -> None:
    setattr(args, "_business_preview", business_preview)
    setattr(args, "_preview_markdown", preview_markdown)


def mutation_preview_payload(
    method: str,
    path: str,
    query: list[tuple[str, str]] | None,
    body: Any,
    args: argparse.Namespace,
) -> dict[str, Any]:
    payload = {
        "technical_preview": redacted_preview(method, path, query, body, getattr(args, "auth_mode", "jwtc"))
    }
    business_preview = getattr(args, "_business_preview", None)
    preview_markdown = getattr(args, "_preview_markdown", None)
    if business_preview is not None:
        payload["business_preview"] = business_preview
    if preview_markdown:
        payload["preview_markdown"] = preview_markdown
    return payload


def is_comprobante_mutation_path(path: str | None) -> bool:
    normalized_path = normalize_search_text(path or "")
    if not normalized_path:
        return False
    return any(normalize_search_text(hint) in normalized_path for hint in COMPROBANTE_MUTATION_PATH_HINTS)


def payload_cancellation_hints(payload: Any, *, prefix: str = "body") -> list[str]:
    hints: list[str] = []
    if isinstance(payload, dict):
        for raw_key, value in payload.items():
            key = str(raw_key)
            normalized_key = normalize_search_text(key)
            current_prefix = f"{prefix}.{key}"
            if any(hint in normalized_key for hint in COMPROBANTE_CANCELLATION_KEY_HINTS):
                hints.append(current_prefix)
            if normalized_key in {"estado", "accion", "action", "operacion", "operation", "modo", "mode"}:
                normalized_value = normalize_search_text(value if isinstance(value, str) else str(value or ""))
                if normalized_value in COMPROBANTE_CANCELLATION_VALUE_HINTS:
                    hints.append(f"{current_prefix}={value}")
            hints.extend(payload_cancellation_hints(value, prefix=current_prefix))
    elif isinstance(payload, list):
        for index, item in enumerate(payload):
            hints.extend(payload_cancellation_hints(item, prefix=f"{prefix}[{index}]"))
    return hints


def ensure_mutation_allowed(method: str, path: str, query: list[tuple[str, str]] | None, body: Any, args: argparse.Namespace) -> None:
    if method.upper() not in MUTATING_METHODS:
        return
    if is_comprobante_mutation_path(path):
        if method.upper() == "DELETE":
            raise CLIError(
                "Las anulaciones o bajas de comprobantes estan bloqueadas por defecto. "
                "Solo pueden habilitarse despues de un pedido explicito del usuario y una validacion previa del mecanismo exacto."
            )
        cancellation_hints = payload_cancellation_hints(body)
        if cancellation_hints:
            joined = ", ".join(cancellation_hints[:6])
            raise CLIError(
                "Se bloquearon campos con semantica de anulacion/cancelacion en un comprobante: "
                f"{joined}. Este skill no puede ejecutar anulaciones sin autorizacion explicita del usuario."
            )
    preview = mutation_preview_payload(method, path, query, body, args)
    if getattr(args, "dry_run", False):
        print(format_output(preview))
        raise SystemExit(0)
    if not getattr(args, "confirm", False):
        print(format_output(preview), file=sys.stderr)
        raise CLIError("Operacion de escritura bloqueada. Revise el preview y vuelva a ejecutar con --confirm.")
    print(format_output(preview), file=sys.stderr)


def pick_first_list(payload: Any) -> list[Any]:
    if isinstance(payload, list):
        return payload
    if isinstance(payload, dict):
        priority = (
            "clientes",
            "productos",
            "cobros",
            "pagos",
            "ventas",
            "data",
            "items",
            "rows",
            "resultados",
            "listado",
            "libro",
        )
        for key in priority:
            value = payload.get(key)
            if isinstance(value, list):
                return value
        for value in payload.values():
            if isinstance(value, list):
                return value
    return []


def search_paginated(
    client: SOSContadorClient,
    *,
    path: str,
    query: list[tuple[str, str]],
    predicate,
    max_pages: int = 20,
) -> dict[str, Any] | None:
    query_map = dict(query)
    page = int(query_map.get("pagina", "1"))
    per_page = int(query_map.get("registros", "200"))
    for _ in range(max_pages):
        current_query = [(key, value) for key, value in query if key not in {"pagina", "registros"}]
        current_query.append(("pagina", str(page)))
        current_query.append(("registros", str(per_page)))
        payload = client.request("GET", path, query=current_query)
        items = pick_first_list(payload)
        for item in items:
            if isinstance(item, dict) and predicate(item):
                return item
        total_pages = extract_total_pages(payload)
        if total_pages is not None:
            if page >= total_pages:
                return None
        elif len(items) < per_page:
            return None
        page += 1
    return None


def extract_total_pages(payload: Any) -> int | None:
    for item in iter_dicts(payload):
        key = find_key(item, ("totalpaginas", "totalPages", "total_pages", "paginas"))
        if key:
            value = item.get(key)
            try:
                return int(value)
            except (TypeError, ValueError):
                return None
    return None


def normalize_name(value: str) -> str:
    return " ".join(value.strip().lower().split())


def find_value(item: dict[str, Any], aliases: tuple[str, ...]) -> Any:
    key = find_key(item, aliases)
    if key:
        return item.get(key)
    return None


def refresh_client_catalog(
    client: SOSContadorClient,
    *,
    include_clientes: bool = True,
    include_proveedores: bool = True,
    per_page: int = 1000,
    max_pages: int = 200,
) -> list[dict[str, Any]]:
    catalog: list[dict[str, Any]] = []
    seen_ids: set[str] = set()
    for page in range(1, max_pages + 1):
        payload = client.request(
            "GET",
            "cliente/listado",
            query=[
                ("cliente", bool_to_str(include_clientes)),
                ("proveedor", bool_to_str(include_proveedores)),
                ("pagina", str(page)),
                ("registros", str(per_page)),
            ],
        )
        items = [item for item in pick_first_list(payload) if isinstance(item, dict)]
        if not items:
            break
        added = 0
        for item in items:
            item_id = str(find_value(item, ("idclipro", "id", "idcliente")) or "").strip()
            if not item_id or item_id in seen_ids:
                continue
            seen_ids.add(item_id)
            enriched = dict(item)
            candidate_name = str(find_value(item, ("clipro", "cliente", "nombre")) or "").strip()
            candidate_cuit = digits_only(str(find_value(item, ("cuit",)) or ""))
            enriched["normalized_name"] = normalize_name(candidate_name) if candidate_name else ""
            enriched["cuit_digits"] = candidate_cuit
            catalog.append(enriched)
            added += 1
        total_pages = extract_total_pages(payload)
        if total_pages is not None:
            if page >= total_pages:
                break
        elif len(items) < per_page:
            break
        if added == 0:
            break
    return catalog


def get_client_catalog(client: SOSContadorClient) -> list[dict[str, Any]]:
    cuit = client.get_session().cuit
    payload = load_document_catalog_cache()
    entry = payload.get(catalog_cache_key("clientes", cuit))
    if isinstance(entry, dict):
        fetched_at = parse_iso_datetime(entry.get("fetched_at"))
        age = datetime.datetime.now(datetime.timezone.utc) - fetched_at if fetched_at is not None else None
        items = entry.get("items")
        if (
            fetched_at is not None
            and age is not None
            and age.total_seconds() <= DOCUMENT_CATALOG_MAX_AGE_SECONDS
            and entry.get("complete") is True
            and int(entry.get("version") or 0) >= 3
            and isinstance(items, list)
        ):
            return [item for item in items if isinstance(item, dict)]
    catalog = refresh_client_catalog(client)
    payload[catalog_cache_key("clientes", cuit)] = {
        "fetched_at": utc_now_iso(),
        "items": catalog,
        "complete": True,
        "version": 3,
    }
    save_document_catalog_cache(payload)
    return catalog


def resolve_cliente_match(
    client: SOSContadorClient,
    *,
    explicit_id: str | None,
    cuit: str | None,
    nombre: str | None,
) -> dict[str, Any]:
    if explicit_id:
        item = command_cliente_get(argparse.Namespace(id=str(explicit_id), cuit=None, nombre=None), client)
        if isinstance(item, dict):
            return item
        raise CLIError("Se encontro el cliente por ID, pero la respuesta no tiene un formato utilizable.")
    if not cuit and not nombre:
        raise CLIError("Debe indicar --idclipro o bien --cliente-cuit / --cliente-nombre.")

    target_cuit = digits_only(cuit or "")
    target_name = normalize_name(nombre or "") if nombre else ""
    catalog = get_client_catalog(client)

    cuit_matches = [item for item in catalog if target_cuit and item.get("cuit_digits") == target_cuit]
    if len(cuit_matches) == 1:
        return cuit_matches[0]
    if len(cuit_matches) > 1 and target_name:
        name_matches = [item for item in cuit_matches if item.get("normalized_name") == target_name]
        if len(name_matches) == 1:
            return name_matches[0]
        raise CLIError("Se encontraron multiples clientes con el mismo CUIT y nombre ambiguo.")

    if target_name:
        name_matches = [item for item in catalog if item.get("normalized_name") == target_name]
        if len(name_matches) == 1:
            return name_matches[0]
        if len(name_matches) > 1 and target_cuit:
            cuit_name_matches = [item for item in name_matches if item.get("cuit_digits") == target_cuit]
            if len(cuit_name_matches) == 1:
                return cuit_name_matches[0]
            raise CLIError("Se encontraron multiples clientes con el mismo nombre. Indique el CUIT o el idclipro.")
        if len(name_matches) > 1:
            raise CLIError("Se encontraron multiples clientes con el mismo nombre. Indique el CUIT o el idclipro.")

    raise CLIError("No se pudo resolver idclipro a partir del cliente indicado.")


def resolve_cliente_id(
    client: SOSContadorClient,
    *,
    explicit_id: str | None,
    cuit: str | None,
    nombre: str | None,
) -> str:
    match = resolve_cliente_match(client, explicit_id=explicit_id, cuit=cuit, nombre=nombre)
    match_id = find_value(match, ("idclipro", "id", "idcliente"))
    if match_id is None:
        raise CLIError("Se encontro el cliente, pero la respuesta no incluye un id utilizable.")
    return str(match_id)


def resolve_provincia_id(
    client: SOSContadorClient,
    *,
    explicit_id: str | None,
    provincia: str | None,
) -> str:
    if explicit_id:
        return explicit_id
    if not provincia:
        raise CLIError("Debe indicar --idprovinciaiibb o bien --provincia.")
    payload = client.request("GET", "provincia/listado")
    items = pick_first_list(payload)
    normalized_target = normalize_name(provincia)
    for item in items:
        if not isinstance(item, dict):
            continue
        province_name = find_value(item, ("provincia", "nombre", "descripcion"))
        if province_name and normalize_name(str(province_name)) == normalized_target:
            province_id = find_value(item, ("idprovincia", "id", "idprovinciaiibb"))
            if province_id is not None:
                return str(province_id)
    raise CLIError("No se pudo resolver idprovinciaiibb a partir de la provincia indicada.")


def validate_imputaciones(imputaciones: Any) -> None:
    if not isinstance(imputaciones, list) or not imputaciones:
        raise CLIError("Las imputaciones deben ser una lista JSON no vacia.")
    for index, item in enumerate(imputaciones, start=1):
        if not isinstance(item, dict):
            raise CLIError(f"La imputacion #{index} debe ser un objeto JSON.")
        if "fv" not in item or "cuid" not in item:
            raise CLIError(f"La imputacion #{index} debe incluir 'fv' y 'cuid'.")


def resolve_centrocosto_id(
    web_client: SOSContadorWebClient,
    *,
    explicit_id: str | None,
    centrocosto: str | None,
) -> str:
    if explicit_id:
        return str(explicit_id)
    target_name = centrocosto or "General"
    normalized_target = normalize_name(target_name)
    for item in web_client.get_centros_costo():
        centro_name = find_value(item, ("centrocosto", "nombre", "descripcion"))
        centro_id = find_value(item, ("idcentrocosto", "id"))
        if centro_id is None or not centro_name:
            continue
        if normalize_name(str(centro_name)) == normalized_target:
            return str(centro_id)
    raise CLIError(f"No se pudo resolver el centro de costo '{target_name}'.")


def resolve_work_cuit_context(client: SOSContadorClient, cuit_value: str) -> dict[str, Any]:
    target = digits_only(cuit_value)
    cached = load_cuit_catalog_cache()
    if not cached or is_cuit_catalog_stale(cached):
        cached = refresh_cuit_catalog(client)
    for item in cached:
        if digits_only(str(item.get("cuit") or "")) == target:
            return {
                "cuit": target,
                "cuit_id": str(item.get("id") or ""),
                "nombre": item.get("razon_social"),
            }
    raise CLIError(f"No se pudo resolver la CUIT de trabajo {cuit_value}.")


def get_bank_catalog(client: SOSContadorClient) -> list[dict[str, Any]]:
    cuit = client.get_session().cuit
    cached = get_cached_document_catalog("bancos", cuit=cuit)
    if cached is not None:
        return cached
    if not hasattr(client, "request"):
        return []
    payload = client.request("GET", "tipo/listado/bancos")
    items = pick_first_list(payload)
    catalog = [item for item in items if isinstance(item, dict)]
    set_cached_document_catalog("bancos", catalog, cuit=cuit)
    return catalog


def simplify_account_label(value: str | None) -> str:
    text = str(value or "").strip()
    text = re.sub(r"^[0-9.]+\s*", "", text)
    text = re.sub(r"^[0-9.]+\s*", "", text)
    return " ".join(text.split())


def get_account_catalog(client: SOSContadorClient) -> list[dict[str, Any]]:
    cuit = client.get_session().cuit
    cached = get_cached_document_catalog("cuentas", cuit=cuit)
    if cached is not None:
        return cached
    payload = client.request("GET", "cuentacontable/listado")
    items = pick_first_list(payload)
    catalog: list[dict[str, Any]] = []
    for item in items:
        if not isinstance(item, dict):
            continue
        if find_value(item, ("id",)) in (None, ""):
            continue
        label = simplify_account_label(find_value(item, ("cuenta",)) or find_value(item, ("rubro",)))
        if not label:
            continue
        enriched = dict(item)
        enriched["label"] = label
        enriched["normalized_label"] = normalize_search_text(label)
        catalog.append(enriched)
    set_cached_document_catalog("cuentas", catalog, cuit=cuit)
    return catalog


def get_centrocosto_catalog(web_client: SOSContadorWebClient) -> list[dict[str, Any]]:
    cuit = web_client.get_session().cuit
    cached = get_cached_document_catalog("centros_costo", cuit=cuit)
    if cached is not None:
        return cached
    catalog = []
    for item in web_client.get_centros_costo():
        centro_name = find_value(item, ("centrocosto", "nombre", "descripcion"))
        centro_id = find_value(item, ("idcentrocosto", "id"))
        if centro_id is None or not centro_name:
            continue
        catalog.append(
            {
                "id": str(centro_id),
                "centrocosto": str(centro_name),
                "normalized_label": normalize_search_text(str(centro_name)),
            }
        )
    set_cached_document_catalog("centros_costo", catalog, cuit=cuit)
    return catalog


def build_candidate_aliases(text: str) -> set[str]:
    normalized = normalize_search_text(text)
    aliases = {normalized}
    if normalized.startswith("banco "):
        aliases.add(normalized[6:].strip())
    aliases.add(normalized.replace(" buenos aires", " bs as"))
    aliases.add(normalized.replace(" bs as", " buenos aires"))
    aliases.add(normalized.replace(" s a ", " ").replace(" s a", "").strip())
    aliases.add(normalized.replace(" sociedad anonima", "").strip())
    return {alias for alias in aliases if alias}


def match_catalog_entry(
    target: str | None,
    catalog: list[dict[str, Any]],
    *,
    label_keys: tuple[str, ...],
) -> dict[str, Any] | None:
    normalized_target = normalize_search_text(target)
    if not normalized_target:
        return None
    aliases = build_candidate_aliases(normalized_target)
    scored: list[tuple[int, dict[str, Any]]] = []
    for item in catalog:
        label = ""
        for key in label_keys:
            if item.get(key):
                label = str(item.get(key))
                break
        normalized_label = normalize_search_text(label)
        if not normalized_label:
            continue
        item_aliases = build_candidate_aliases(normalized_label)
        score = 0
        if aliases & item_aliases:
            score = 100
        elif any(alias in normalized_label for alias in aliases):
            score = 80
        else:
            alias_tokens = set(" ".join(aliases).split())
            label_tokens = set(normalized_label.split())
            overlap = len(alias_tokens & label_tokens)
            if overlap:
                score = overlap * 10
        if score:
            scored.append((score, item))
    if not scored:
        return None
    scored.sort(key=lambda entry: entry[0], reverse=True)
    best_score = scored[0][0]
    best = [item for score, item in scored if score == best_score]
    if len(best) != 1:
        return None
    return best[0]


def resolve_bank_match(client: SOSContadorClient, bank_name: str | None) -> dict[str, Any] | None:
    if not bank_name:
        return None
    return match_catalog_entry(bank_name, get_bank_catalog(client), label_keys=("bancos",))


def resolve_account_match(client: SOSContadorClient, account_name: str | None) -> dict[str, Any] | None:
    if not account_name:
        return None
    return match_catalog_entry(account_name, get_account_catalog(client), label_keys=("label", "cuenta", "rubro"))


def resolve_centrocosto_match(web_client: SOSContadorWebClient, centro_name: str | None) -> dict[str, Any] | None:
    target = centro_name or "General"
    return match_catalog_entry(target, get_centrocosto_catalog(web_client), label_keys=("centrocosto",))


def parse_web_date(value: str | None, *, field_name: str) -> str:
    if value is None or not str(value).strip():
        raise CLIError(f"Falta {field_name}.")
    text = str(value).strip()
    for fmt in ("%d/%m/%Y", "%d/%m/%y", "%Y-%m-%d", "%Y-%m-%dT%H:%M:%S.%f"):
        try:
            return datetime.datetime.strptime(text, fmt).strftime("%d/%m/%Y")
        except ValueError:
            continue
    raise CLIError(f"{field_name} invalida '{value}'. Use DD/MM/YYYY o YYYY-MM-DD.")


def parse_required_decimal_text(value: Any, *, field_name: str) -> str:
    if value in (None, ""):
        raise CLIError(f"Falta {field_name}.")
    text = str(value).strip()
    normalized = text.replace(" ", "")
    if "," in normalized and "." not in normalized:
        normalized = normalized.replace(",", ".")
    try:
        decimal_value = Decimal(normalized)
    except (InvalidOperation, ValueError) as exc:
        raise CLIError(f"{field_name} invalido '{value}'.") from exc
    return format_decimal_string(decimal_value)


def parse_numeric_text(value: Any, *, field_name: str, required: bool = True) -> str:
    text = str(value or "").strip()
    if not text:
        if required:
            raise CLIError(f"Falta {field_name}.")
        return ""
    if not text.isdigit():
        raise CLIError(f"{field_name} invalido '{value}'. Debe contener solo digitos.")
    return text


def coalesce_cobro_observaciones(args: argparse.Namespace) -> str:
    parts: list[str] = []
    for value in (getattr(args, "comentarios", None), getattr(args, "memo", None)):
        text = str(value or "").strip()
        if text and text not in parts:
            parts.append(text)
    referencia = str(getattr(args, "referencia", "") or "").strip()
    if referencia:
        ref_text = f"Ref: {referencia}"
        if ref_text not in parts:
            parts.append(ref_text)
    return " | ".join(parts)


def detect_source_kind(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return "pdf"
    if suffix in {".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff"}:
        return "image"
    if suffix == ".txt":
        return "txt"
    if suffix == ".csv":
        return "csv"
    if suffix in {".xlsx", ".xlsm"}:
        return "xlsx"
    if suffix == ".xls":
        return "xls"
    return "text"


def clean_text_lines(text: str) -> list[str]:
    lines: list[str] = []
    for raw_line in text.splitlines():
        line = " ".join(str(raw_line).replace("\x00", " ").split())
        if line:
            lines.append(line)
    return lines


def rows_to_text(rows: list[list[str]]) -> str:
    return "\n".join(" | ".join(cell for cell in row if str(cell).strip()) for row in rows if any(str(cell).strip() for cell in row))


def read_csv_source(path: Path) -> tuple[str, list[list[str]]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        rows = [[str(cell).strip() for cell in row] for row in csv.reader(handle)]
    return rows_to_text(rows), rows


def read_xlsx_source(path: Path) -> tuple[str, list[list[str]]]:
    if load_workbook is None:
        raise CLIError("Falta openpyxl para leer archivos XLSX.")
    workbook = load_workbook(path, read_only=True, data_only=True)
    rows: list[list[str]] = []
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows(values_only=True):
            values = [str(cell).strip() if cell not in (None, "") else "" for cell in row]
            if any(values):
                rows.append(values)
    return rows_to_text(rows), rows


def read_xls_source(path: Path) -> tuple[str, list[list[str]]]:
    if xlrd is None:
        raise CLIError("Falta xlrd para leer archivos XLS.")
    workbook = xlrd.open_workbook(path)
    rows: list[list[str]] = []
    for sheet in workbook.sheets():
        for row_index in range(sheet.nrows):
            row = [str(sheet.cell_value(row_index, col_index)).strip() for col_index in range(sheet.ncols)]
            if any(row):
                rows.append(row)
    return rows_to_text(rows), rows


def read_pdf_text(path: Path) -> str:
    parts: list[str] = []
    if PdfReader is not None:
        try:
            reader = PdfReader(str(path))
            for page in reader.pages:
                parts.append(page.extract_text() or "")
        except Exception:
            parts = []
    if not "".join(parts).strip() and fitz is not None:
        try:
            with fitz.open(path) as document:
                parts = [page.get_text("text") or "" for page in document]
        except Exception:
            parts = []
    return "\n".join(part for part in parts if part)


def read_pdf_via_ocr(path: Path) -> str:
    if fitz is None:
        raise CLIError("Falta PyMuPDF para OCR de PDFs escaneados.")
    if not ocr_is_available():
        raise CLIError("OCR no disponible para PDFs escaneados.")
    parts: list[str] = []
    with fitz.open(path) as document:
        for page in document:
            pix = page.get_pixmap(dpi=200)
            image = Image.open(io.BytesIO(pix.tobytes("png")))
            parts.append(pytesseract.image_to_string(image, lang="spa+eng"))
    return "\n".join(parts)


def read_image_via_ocr(path: Path) -> str:
    if Image is None or pytesseract is None:
        raise CLIError("Faltan dependencias para OCR de imágenes.")
    if not ocr_is_available():
        raise CLIError("OCR no disponible para imágenes.")
    with Image.open(path) as image:
        return pytesseract.image_to_string(image, lang="spa+eng")


def extract_source_document(path: Path) -> dict[str, Any]:
    if not path.exists():
        raise CLIError(f"No existe el archivo fuente: {path}")
    kind = detect_source_kind(path)
    text = ""
    rows: list[list[str]] = []
    warnings: list[str] = []
    capabilities: list[str] = []

    if kind == "txt" or kind == "text":
        text = path.read_text(encoding="utf-8", errors="replace")
    elif kind == "csv":
        text, rows = read_csv_source(path)
    elif kind == "xlsx":
        text, rows = read_xlsx_source(path)
    elif kind == "xls":
        text, rows = read_xls_source(path)
    elif kind == "pdf":
        text = read_pdf_text(path)
        if not text.strip():
            try:
                text = read_pdf_via_ocr(path)
                warnings.append("PDF escaneado procesado por OCR.")
            except CLIError as exc:
                capabilities.append(str(exc))
    elif kind == "image":
        try:
            text = read_image_via_ocr(path)
        except CLIError as exc:
            capabilities.append(str(exc))

    return {
        "path": str(path),
        "kind": kind,
        "text": text,
        "lines": clean_text_lines(text),
        "rows": rows,
        "warnings": warnings,
        "capabilities": capabilities,
    }


def collect_counterparty_cuits(
    lines: list[str],
    *,
    work_cuit: str | None,
    explicit_cuit: str | None,
) -> list[str]:
    candidates: list[str] = []
    explicit = digits_only(explicit_cuit or "")
    if explicit:
        candidates.append(explicit)
    work_digits = digits_only(work_cuit or "")
    for item in extract_cuit_candidates(lines):
        candidate = digits_only(item.get("cuit"))
        if not candidate or candidate == work_digits or candidate in candidates:
            continue
        candidates.append(candidate)
    return candidates


def candidate_profile_parent_dirs(
    *,
    work_cuit: str,
    counterparty_cuits: list[str],
    document_kind: str,
) -> list[Path]:
    parents: list[Path] = []
    work_dir = DOCUMENT_PROFILE_ROOT / digits_only(work_cuit)
    for counterparty in counterparty_cuits:
        parents.append(work_dir / profile_counterparty_key(counterparty) / document_kind)
    parents.append(work_dir / "_unknown" / document_kind)
    return parents


def iter_local_document_profiles(
    *,
    work_cuit: str,
    counterparty_cuits: list[str],
    document_kind: str,
) -> list[dict[str, Any]]:
    profiles: list[dict[str, Any]] = []
    seen_paths: set[str] = set()
    for parent in candidate_profile_parent_dirs(
        work_cuit=work_cuit,
        counterparty_cuits=counterparty_cuits,
        document_kind=document_kind,
    ):
        if not parent.exists():
            continue
        for child in sorted(parent.iterdir()):
            if not child.is_dir():
                continue
            manifest_path = child / "manifest.json"
            rules_path = child / "rules.json"
            if not manifest_path.exists() or not rules_path.exists():
                continue
            key = str(child.resolve())
            if key in seen_paths:
                continue
            seen_paths.add(key)
            try:
                manifest = load_profile_json(manifest_path)
                rules = load_profile_json(rules_path)
            except CLIError:
                continue
            profiles.append(
                {
                    "path": child,
                    "manifest": manifest,
                    "rules": rules,
                }
            )
    return profiles


def profile_string_list(value: Any) -> list[str]:
    if isinstance(value, str):
        return [value]
    if isinstance(value, list):
        return [str(item) for item in value if str(item).strip()]
    return []


def profile_matches_source(
    profile: dict[str, Any],
    *,
    work_cuit: str,
    counterparty_cuits: list[str],
    file_names: list[str],
    normalized_text: str,
) -> tuple[bool, int, list[str]]:
    manifest = profile.get("manifest", {})
    reasons: list[str] = []
    if digits_only(manifest.get("work_cuit")) not in {"", digits_only(work_cuit)}:
        return False, -1, reasons
    manifest_counterparty = digits_only(manifest.get("counterparty_cuit"))
    if manifest_counterparty and manifest_counterparty not in counterparty_cuits:
        return False, -1, reasons
    filename_rules = [normalize_search_text(item) for item in profile_string_list(manifest.get("filename_contains"))]
    if filename_rules and not all(any(rule in file_name for file_name in file_names) for rule in filename_rules):
        return False, -1, reasons
    required_text = [normalize_search_text(item) for item in profile_string_list(manifest.get("required_text"))]
    if any(item and item not in normalized_text for item in required_text):
        return False, -1, reasons
    forbidden_text = [normalize_search_text(item) for item in profile_string_list(manifest.get("forbidden_text"))]
    if any(item and item in normalized_text for item in forbidden_text):
        return False, -1, reasons
    score = int(manifest.get("priority") or 0)
    if manifest_counterparty:
        score += 100
        reasons.append(f"contra={manifest_counterparty}")
    score += len(required_text) * 10
    score += len(filename_rules) * 5
    if manifest.get("name"):
        reasons.append(str(manifest.get("name")))
    return True, score, reasons


def select_local_document_profile(
    *,
    work_cuit: str | None,
    counterparty_cuits: list[str],
    document_kind: str,
    extracted_sources: list[dict[str, Any]],
) -> dict[str, Any] | None:
    if not work_cuit:
        return None
    file_names = [normalize_search_text(Path(item.get("path") or "").name) for item in extracted_sources]
    combined_text = normalize_search_text("\n".join(item.get("text") or "" for item in extracted_sources))
    scored: list[tuple[int, dict[str, Any], list[str]]] = []
    for profile in iter_local_document_profiles(
        work_cuit=work_cuit,
        counterparty_cuits=counterparty_cuits,
        document_kind=document_kind,
    ):
        matches, score, reasons = profile_matches_source(
            profile,
            work_cuit=work_cuit,
            counterparty_cuits=counterparty_cuits,
            file_names=file_names,
            normalized_text=combined_text,
        )
        if matches:
            scored.append((score, profile, reasons))
    if not scored:
        return None
    scored.sort(key=lambda entry: entry[0], reverse=True)
    best_score, best_profile, reasons = scored[0]
    selected = dict(best_profile)
    selected["match_score"] = best_score
    selected["match_reasons"] = reasons
    return selected


def profile_compiled_pattern(rule: dict[str, Any]) -> re.Pattern[str]:
    pattern = str(rule.get("pattern") or "").strip()
    if not pattern:
        raise CLIError("Regla de perfil sin pattern.")
    flags = 0
    for item in profile_string_list(rule.get("flags")):
        if str(item).lower() == "ignorecase":
            flags |= re.IGNORECASE
        elif str(item).lower() == "multiline":
            flags |= re.MULTILINE
        elif str(item).lower() == "dotall":
            flags |= re.DOTALL
    return re.compile(pattern, flags)


def apply_profile_value_transform(value: str | None, transform: str | None) -> str:
    text = str(value or "").strip()
    if not text:
        return ""
    transform_name = str(transform or "").strip().lower()
    if transform_name == "date":
        return guess_date_from_text(text) or ""
    if transform_name == "decimal":
        return parse_optional_decimal_text(text)
    if transform_name == "digits":
        return digits_only(text)
    if transform_name == "comprobante":
        return normalize_comprobante_reference(text) or ""
    if transform_name == "normalize-space":
        return " ".join(text.split())
    return text


def extract_profile_rule_matches(
    corpus: str,
    lines: list[str],
    rule: dict[str, Any],
) -> list[tuple[re.Match[str], str]]:
    pattern = profile_compiled_pattern(rule)
    scope = str(rule.get("scope") or PROFILE_SCOPE_LINE).strip().lower()
    results: list[tuple[re.Match[str], str]] = []
    if scope == PROFILE_SCOPE_TEXT:
        for match in pattern.finditer(corpus):
            results.append((match, match.group(0)))
        return results
    for line in lines:
        for match in pattern.finditer(line):
            results.append((match, line))
    return results


def build_profile_field_status(
    *,
    match: re.Match[str],
    rule: dict[str, Any],
    source_text: str,
) -> dict[str, Any]:
    group_name = str(rule.get("value_group") or rule.get("group") or "value")
    raw_value = match.groupdict().get(group_name) if match.groupdict() else None
    if raw_value is None:
        raw_value = match.group(1) if match.groups() else match.group(0)
    value = apply_profile_value_transform(raw_value, str(rule.get("transform") or ""))
    return {
        "value": value,
        "estado": "inferido",
        "source_text": source_text,
    }


def extract_profile_fields(
    rules: dict[str, Any],
    *,
    corpus: str,
    lines: list[str],
) -> dict[str, dict[str, Any]]:
    extracted: dict[str, dict[str, Any]] = {}
    fields = rules.get("fields")
    if not isinstance(fields, dict):
        return extracted
    for field_name, field_rules in fields.items():
        if not isinstance(field_rules, list):
            continue
        for rule in field_rules:
            if not isinstance(rule, dict):
                continue
            matches = extract_profile_rule_matches(corpus, lines, rule)
            if not matches:
                continue
            first_match, source_text = matches[0]
            extracted[str(field_name)] = build_profile_field_status(match=first_match, rule=rule, source_text=source_text)
            break
    return extracted


def extract_profile_items(
    item_rules: Any,
    *,
    corpus: str,
    lines: list[str],
) -> list[dict[str, Any]]:
    items: list[dict[str, Any]] = []
    if not isinstance(item_rules, list):
        return items
    for rule in item_rules:
        if not isinstance(rule, dict):
            continue
        groups = rule.get("groups")
        defaults = rule.get("defaults")
        if not isinstance(groups, dict):
            groups = {}
        if not isinstance(defaults, dict):
            defaults = {}
        for match, source_text in extract_profile_rule_matches(corpus, lines, rule):
            item: dict[str, Any] = {
                "source_text": source_text,
                "estado": "inferido",
            }
            for key, default_value in defaults.items():
                if (
                    str(key) == "fecha"
                    and str(rule.get("tipo") or "") == "retencion"
                    and "fecha" not in groups
                ):
                    item[str(key)] = PROFILE_DOCUMENT_DATE_TOKEN
                else:
                    item[str(key)] = default_value
            for target_key, group_name in groups.items():
                raw_value = match.groupdict().get(str(group_name)) if match.groupdict() else None
                if raw_value is None:
                    continue
                transform_name = None
                transforms = rule.get("transforms")
                if isinstance(transforms, dict):
                    transform_name = transforms.get(str(target_key))
                item[str(target_key)] = apply_profile_value_transform(raw_value, str(transform_name or ""))
            if "tipo" in rule and not item.get("tipo"):
                item["tipo"] = rule.get("tipo")
            items.append(item)
    return items


def extract_document_profile_payload(
    profile: dict[str, Any],
    *,
    extracted_sources: list[dict[str, Any]],
) -> dict[str, Any]:
    lines: list[str] = []
    for item in extracted_sources:
        lines.extend(item.get("lines") or [])
    corpus = normalize_lines_corpus(lines)
    rules = profile.get("rules", {})
    fields = extract_profile_fields(rules, corpus=corpus, lines=lines)
    document_date = str(fields.get("fecha", {}).get("value") or "").strip()
    profile_movements = []
    for item in extract_profile_items(rules.get("movimientos"), corpus=corpus, lines=lines):
        if not parse_optional_decimal_text(item.get("monto")):
            continue
        if item.get("fecha") == PROFILE_DOCUMENT_DATE_TOKEN:
            item["fecha"] = document_date
        profile_movements.append(item)
    profile_facturas = [
        item
        for item in extract_profile_items(rules.get("facturas"), corpus=corpus, lines=lines)
        if normalize_comprobante_reference(item.get("comprobante")) or str(item.get("venta_id") or "").strip()
    ]
    return {
        "metadata": {
            "profile_id": str(profile.get("manifest", {}).get("profile_id") or profile.get("path", Path()).name),
            "name": profile.get("manifest", {}).get("name"),
            "path": str(profile.get("path")),
            "match_score": profile.get("match_score"),
            "match_reasons": profile.get("match_reasons", []),
        },
        "fecha": fields.get("fecha"),
        "comentarios": fields.get("comentarios"),
        "cliente": {
            "nombre": fields.get("cliente_nombre", {}).get("value"),
            "cuit": digits_only(fields.get("cliente_cuit", {}).get("value") or ""),
            "estado": max(
                [
                    fields.get("cliente_nombre", {}).get("estado", "falta"),
                    fields.get("cliente_cuit", {}).get("estado", "falta"),
                ],
                key=lambda item: {"confirmado": 3, "inferido": 2, "falta": 1}.get(str(item), 0),
            ),
            "source_text": fields.get("cliente_nombre", {}).get("source_text")
            or fields.get("cliente_cuit", {}).get("source_text"),
        },
        "movimientos": profile_movements,
        "facturas": profile_facturas,
    }


def header_key(value: str) -> str:
    normalized = normalize_search_text(value)
    aliases = {
        "tipo": "tipo",
        "movimiento": "tipo",
        "medio": "tipo",
        "cuenta": "cuenta",
        "cuenta destino": "cuenta",
        "metodo": "cuenta",
        "metodo de cobranza": "cuenta",
        "banco": "banco",
        "bancos": "banco",
        "numero": "numero",
        "num": "numero",
        "nro": "numero",
        "num cheque": "numero",
        "num retencion": "numero",
        "fecha": "fecha",
        "vencimiento": "fecha",
        "fecha cobro": "fecha",
        "monto": "monto",
        "importe": "monto",
        "valor": "monto",
        "regimen": "regimen",
        "reg": "regimen",
        "cliente": "cliente",
        "razon social": "cliente",
        "cuit": "cuit",
        "comentarios": "comentarios",
        "comentario": "comentarios",
        "factura": "factura",
        "comprobante": "factura",
    }
    return aliases.get(normalized, normalized)


def extract_structured_rows(source_rows: list[list[str]]) -> list[dict[str, str]]:
    if not source_rows:
        return []
    header = [header_key(cell) for cell in source_rows[0]]
    recognized = {"tipo", "cuenta", "banco", "numero", "fecha", "monto", "regimen", "cliente", "cuit", "comentarios", "factura"}
    if len(set(header) & recognized) < 2:
        return []
    rows: list[dict[str, str]] = []
    for raw_row in source_rows[1:]:
        item = {}
        for index, key in enumerate(header):
            if not key or index >= len(raw_row):
                continue
            value = str(raw_row[index]).strip()
            if value:
                item[key] = value
        if item:
            rows.append(item)
    return rows


def extract_cuit_candidates(lines: list[str]) -> list[dict[str, Any]]:
    results: list[dict[str, Any]] = []
    pattern = re.compile(r"\b(\d{2})[- ]?(\d{8})[- ]?(\d)\b")
    for line_index, line in enumerate(lines):
        for match in pattern.finditer(line):
            cuit = "".join(match.groups())
            results.append({"cuit": cuit, "line": line, "line_index": line_index})
    unique: dict[str, dict[str, Any]] = {}
    for item in results:
        unique.setdefault(item["cuit"], item)
    return list(unique.values())


def extract_document_work_cuit(lines: list[str], client: SOSContadorClient) -> dict[str, Any]:
    candidates = extract_cuit_candidates(lines)
    if not candidates:
        return {"cuit": "", "estado": "falta"}
    catalog = ensure_cuit_catalog(client)
    by_cuit = {
        digits_only(str(item.get("cuit") or "")): item
        for item in catalog
        if digits_only(str(item.get("cuit") or ""))
    }
    scored: list[tuple[int, dict[str, Any], dict[str, Any], list[str]]] = []
    for candidate in candidates:
        entry = by_cuit.get(digits_only(str(candidate.get("cuit") or "")))
        if entry is None:
            continue
        line_index = int(candidate.get("line_index") or 0)
        window = lines[max(0, line_index - 2): min(len(lines), line_index + 3)]
        window_text = normalize_search_text("\n".join(window))
        score = 0
        reasons: list[str] = []
        if any(hint in window_text for hint in WORK_CUIT_POSITIVE_HINTS):
            score += 50
            reasons.append("buyer_section")
        if any(hint in window_text for hint in WORK_CUIT_NEGATIVE_HINTS):
            score -= 40
            reasons.append("seller_section")
        normalized_name = normalize_entity_name(str(entry.get("razon_social") or ""))
        if normalized_name and normalized_name in window_text:
            score += 20
            reasons.append("matched_name")
        if score > 0:
            scored.append((score, candidate, entry, reasons))
    if not scored:
        return {
            "cuit": "",
            "estado": "falta",
            "warning": "No se pudo inferir una CUIT de trabajo unica desde una seccion clara de comprador/receptor.",
        }
    scored.sort(key=lambda item: item[0], reverse=True)
    best_score = scored[0][0]
    best_matches = [item for item in scored if item[0] == best_score]
    if len(best_matches) != 1:
        return {
            "cuit": "",
            "estado": "ambiguo",
            "warning": "El documento contiene multiples CUITs de trabajo posibles. Confirme el CUIT de trabajo explicitamente.",
        }
    _, candidate, entry, reasons = best_matches[0]
    return {
        "cuit": digits_only(str(entry.get("cuit") or "")),
        "cuit_id": str(entry.get("id") or ""),
        "nombre": entry.get("razon_social"),
        "estado": "inferido",
        "source_text": candidate.get("line"),
        "matched_by": ",".join(reasons),
    }


def pick_document_client(
    lines: list[str],
    *,
    work_cuit: str | None,
    explicit_cuit: str | None,
    explicit_name: str | None,
) -> dict[str, Any]:
    if explicit_cuit or explicit_name:
        return {
            "nombre": explicit_name,
            "cuit": digits_only(explicit_cuit or "") or None,
            "estado": "confirmado",
            "source_text": "hint del usuario",
        }
    cuit_candidates = extract_cuit_candidates(lines)
    if work_cuit:
        for candidate in cuit_candidates:
            if candidate["cuit"] != digits_only(work_cuit):
                return {
                    "nombre": None,
                    "cuit": candidate["cuit"],
                    "estado": "inferido",
                    "source_text": candidate["line"],
                }
    if len(cuit_candidates) == 1:
        return {
            "nombre": None,
            "cuit": cuit_candidates[0]["cuit"],
            "estado": "inferido",
            "source_text": cuit_candidates[0]["line"],
        }
    for line in lines:
        if "cliente" in normalize_search_text(line):
            candidate_name = re.split(r":", line, maxsplit=1)
            if len(candidate_name) == 2:
                return {
                    "nombre": candidate_name[1].strip(),
                    "cuit": None,
                    "estado": "inferido",
                    "source_text": line,
                }
    return {"nombre": None, "cuit": None, "estado": "falta"}


def extract_document_comment(lines: list[str]) -> dict[str, Any]:
    for line in lines:
        if "orden de pago" in normalize_search_text(line):
            candidate = re.sub(r"(?i)^.*?(orden de pago)", r"\1", line).strip()
            candidate = re.sub(r"\s+", " ", candidate).strip(" :")
            if candidate:
                return {
                    "value": candidate,
                    "estado": "inferido",
                    "source_text": line,
                }
    patterns = [
        re.compile(r"(?i)\bop\b[:\s#-]*(-?[A-Z0-9/]+)"),
    ]
    for line in lines:
        for pattern in patterns:
            match = pattern.search(line)
            if match:
                return {
                    "value": f"Orden de pago {match.group(1)}",
                    "estado": "inferido",
                    "source_text": line,
                }
    return {"value": "", "estado": "falta"}


def extract_document_date(lines: list[str], explicit_date: str | None = None) -> dict[str, Any]:
    if explicit_date:
        return {"value": parse_web_date(explicit_date, field_name="fecha"), "estado": "confirmado", "source_text": "hint del usuario"}
    pattern = re.compile(r"\b(?:\d{1,2}[/-]\d{1,2}[/-]\d{2,4}|\d{4}[/-]\d{1,2}[/-]\d{1,2})\b")
    for line in lines:
        match = pattern.search(line)
        if match:
            date_value = parse_web_date(match.group(0), field_name="fecha")
            return {"value": date_value, "estado": "inferido", "source_text": line}
    return {"value": "", "estado": "falta"}


def amount_candidates_in_line(line: str) -> list[str]:
    pattern = re.compile(r"(?<!\d)(?:\d{1,3}(?:[.,]\d{3})*(?:[.,]\d{2})|\d+(?:[.,]\d{2}))(?!\d)")
    return [match.group(0) for match in pattern.finditer(line)]


def date_candidates_in_line(line: str) -> list[str]:
    pattern = re.compile(r"\b\d{1,2}[/-]\d{1,2}[/-]\d{2,4}\b")
    return [match.group(0) for match in pattern.finditer(line)]


def long_numeric_tokens(line: str) -> list[str]:
    pattern = re.compile(r"\b\d{5,12}\b")
    return [match.group(0) for match in pattern.finditer(line)]


def extract_bank_hint_from_line(line: str) -> str | None:
    date_matches = date_candidates_in_line(line)
    if date_matches:
        prefix = line.split(date_matches[0], 1)[0].strip(" -|")
        prefix = re.sub(r"(?i)\b(chp|cheque|e-?cheq)\b\s*(de)?\s*", "", prefix).strip(" -|.")
        if prefix:
            return prefix
    match = re.search(r"(?i)\b(banco\s+[a-záéíóúñ.& ]+)", line)
    if match:
        return match.group(1).strip()
    return None


def parse_structured_movements(rows: list[dict[str, str]]) -> list[dict[str, Any]]:
    movements: list[dict[str, Any]] = []
    for row in rows:
        amount = row.get("monto") or row.get("importe") or row.get("valor")
        if not amount:
            continue
        tipo_raw = normalize_search_text(row.get("tipo") or "")
        bank_name = row.get("banco")
        regimen = row.get("regimen")
        account = row.get("cuenta")
        movement_type = "simple"
        if bank_name or "cheque" in tipo_raw:
            movement_type = "cheque"
        elif regimen or "ret" in tipo_raw:
            movement_type = "retencion"
        movements.append(
            {
                "tipo": movement_type,
                "cuenta_nombre": account,
                "banco_nombre": bank_name,
                "fecha": row.get("fecha"),
                "numero": row.get("numero"),
                "regimen": regimen,
                "monto": amount,
                "source_text": " | ".join(f"{key}={value}" for key, value in row.items()),
            }
        )
    return movements


def infer_retention_account_name(line: str) -> str | None:
    normalized = normalize_search_text(line)
    if "gan" in normalized or re.search(r"(?i)\brg[.\s-]*830\b", line) or "enaj bs mbles" in normalized:
        return "Retencion Ganancias Sufrida"
    if "iva" in normalized:
        return "Retencion IVA Sufrida"
    if "iibb" in normalized:
        provinces = [
            "buenos aires",
            "caba",
            "cordoba",
            "mendoza",
            "santa fe",
            "tucuman",
            "salta",
            "misiones",
            "neuquen",
            "rio negro",
        ]
        for province in provinces:
            if province in normalized:
                return f"Retención IIBB Sufrida {province.title()}"
        return "Retencion IIBB Sufrida"
    return None


def looks_like_retention_line(line: str) -> bool:
    normalized = normalize_search_text(line)
    if any(token in normalized for token in ("ret", "gan", "iva", "iibb", "sicore")):
        return True
    if re.search(r"(?i)\brg[.\s-]*\d{2,4}\b", line):
        return True
    if "enaj bs mbles" in normalized:
        return True
    return False


def extract_retention_regimen(line: str) -> str:
    patterns = (
        re.compile(r"(?i)\brg[.\s-]*(\d{2,4})\b"),
        re.compile(r"(?i)\b(?:reg|regimen|sicore)\D{0,5}(\d{2,4})\b"),
    )
    for pattern in patterns:
        match = pattern.search(line)
        if match:
            return match.group(1)
    return ""


def extract_movement_source_line(line: str) -> str:
    marker_match = re.search(r"(?i)\b(chp|cheque|e-?cheq)\b", line)
    if marker_match:
        return line[marker_match.start():].strip()
    return line


def parse_line_movements(lines: list[str], client: SOSContadorClient | None) -> list[dict[str, Any]]:
    movements: list[dict[str, Any]] = []
    seen: set[tuple[str, str, str, str]] = set()
    bank_catalog = get_bank_catalog(client) if client is not None else []
    for line in lines:
        working_line = extract_movement_source_line(line)
        normalized = normalize_search_text(working_line)
        amounts = amount_candidates_in_line(working_line)
        if not amounts:
            continue
        amount = amounts[-1]
        dates = date_candidates_in_line(working_line)
        numeric_tokens = long_numeric_tokens(working_line)
        comprobante_ref = normalize_comprobante_reference(working_line)
        if comprobante_ref and is_sale_comprobante(comprobante_ref):
            continue
        account_name = None
        bank_hint = extract_bank_hint_from_line(working_line)
        bank_match = match_catalog_entry(bank_hint, bank_catalog, label_keys=("bancos",)) if bank_hint else None
        movement_type = "simple"
        regimen = ""
        if looks_like_retention_line(working_line):
            movement_type = "retencion"
            account_name = infer_retention_account_name(working_line)
            regimen = extract_retention_regimen(working_line)
        elif bank_match is not None or "cheque" in normalized or re.search(r"(?i)\bchp\b", working_line):
            movement_type = "cheque"
            account_name = "Valores A Depositar"
        elif dates and numeric_tokens and len(amounts) == 1:
            movement_type = "cheque"
            account_name = "Valores A Depositar"
        elif "banco" in normalized and len(amounts) == 1:
            movement_type = "simple"
            account_name = "Banco"
        else:
            continue

        date_value = dates[0] if dates else ""
        number_candidates = [token for token in numeric_tokens if token not in digits_only(amount)]
        number = number_candidates[0] if number_candidates else ""
        key = (movement_type, date_value, number, parse_optional_decimal_text(amount))
        if key in seen:
            continue
        seen.add(key)
        movements.append(
            {
                "tipo": movement_type,
                "cuenta_nombre": account_name,
                "banco_nombre": bank_match.get("bancos") if bank_match else None,
                "fecha": date_value,
                "numero": number,
                "regimen": regimen,
                "monto": amount,
                "source_text": working_line,
            }
        )
    return movements


def extract_invoice_candidates(lines: list[str], structured_rows: list[dict[str, str]]) -> list[dict[str, Any]]:
    matches: dict[str, dict[str, Any]] = {}
    for row in structured_rows:
        comprobante = normalize_comprobante_reference(row.get("factura") or row.get("comprobante"))
        if not comprobante or not is_sale_comprobante(comprobante):
            continue
        matches.setdefault(
            comprobante,
            {
                "comprobante": comprobante,
                "fecha": guess_date_from_text(row.get("fecha")),
                "total": parse_optional_decimal_text(row.get("monto") or row.get("importe")),
                "estado": "inferido",
                "source_text": " | ".join(f"{key}={value}" for key, value in row.items()),
            },
        )
    for line in lines:
        comprobante = normalize_comprobante_reference(line)
        if not comprobante or not is_sale_comprobante(comprobante):
            continue
        matches.setdefault(
            comprobante,
            {
                "comprobante": comprobante,
                "fecha": guess_date_from_text(next(iter(date_candidates_in_line(line)), "")),
                "total": parse_optional_decimal_text(next(iter(amount_candidates_in_line(line)), "")),
                "estado": "inferido",
                "source_text": line,
            },
        )
    return list(matches.values())


def status_rank(state: str | None) -> int:
    return {"confirmado": 3, "inferido": 2, "falta": 1}.get(str(state or ""), 0)


def coalesce_status_info(
    primary: dict[str, Any] | None,
    secondary: dict[str, Any] | None,
) -> dict[str, Any]:
    primary = primary or {}
    secondary = secondary or {}
    return {
        "value": primary.get("value") or secondary.get("value") or "",
        "estado": primary.get("estado")
        if status_rank(primary.get("estado")) >= status_rank(secondary.get("estado"))
        else secondary.get("estado"),
        "source_text": primary.get("source_text") or secondary.get("source_text"),
    }


def merge_client_candidates(
    primary: dict[str, Any] | None,
    secondary: dict[str, Any] | None,
) -> dict[str, Any]:
    primary = primary or {}
    secondary = secondary or {}
    return {
        "nombre": primary.get("nombre") or secondary.get("nombre"),
        "cuit": digits_only(primary.get("cuit") or "") or digits_only(secondary.get("cuit") or "") or None,
        "estado": primary.get("estado")
        if status_rank(primary.get("estado")) >= status_rank(secondary.get("estado"))
        else secondary.get("estado"),
        "source_text": primary.get("source_text") or secondary.get("source_text"),
    }


def movement_identity(item: dict[str, Any]) -> tuple[str, str, str, str]:
    movement_type = str(item.get("tipo") or "")
    amount = parse_optional_decimal_text(item.get("monto")) or ""
    number = digits_only(item.get("numero") or "")
    if movement_type == "cheque":
        return (movement_type, number, amount, "")
    if movement_type == "retencion":
        return (movement_type, digits_only(item.get("regimen") or ""), number, amount)
    return (
        movement_type,
        simplify_account_label(item.get("cuenta_nombre")),
        guess_date_from_text(item.get("fecha")) or "",
        amount,
    )


def invoice_identity(item: dict[str, Any]) -> tuple[str, str]:
    comprobante = normalize_comprobante_reference(item.get("comprobante")) or ""
    venta_id = str(item.get("venta_id") or item.get("id") or "").strip()
    return comprobante, venta_id


def merge_item_lists(
    primary: list[dict[str, Any]],
    secondary: list[dict[str, Any]],
    *,
    identity,
) -> list[dict[str, Any]]:
    merged: list[dict[str, Any]] = []
    seen: set[Any] = set()
    for collection in (primary, secondary):
        for item in collection:
            if not isinstance(item, dict):
                continue
            key = identity(item)
            if key in seen:
                continue
            seen.add(key)
            merged.append(item)
    return merged


def draft_missing_field(section: str, field: str, *, source_text: str | None = None) -> dict[str, Any]:
    payload = {"section": section, "field": field}
    if source_text:
        payload["source_text"] = source_text
    return payload


def build_cobro_draft(args: argparse.Namespace, client: SOSContadorClient) -> dict[str, Any]:
    start = time.perf_counter()
    sources = [Path(path) for path in (getattr(args, "source", None) or [])]
    if not sources:
        raise CLIError("cobro draft requiere al menos un --source.")
    explicit_work_target = resolve_work_cuit_target_from_args(client, args, required=False)
    extracted_sources = [extract_source_document(path) for path in sources]
    all_lines: list[str] = []
    all_rows: list[dict[str, str]] = []
    warnings: list[str] = []
    capabilities: list[str] = []
    for source in extracted_sources:
        all_lines.extend(source["lines"])
        all_rows.extend(extract_structured_rows(source["rows"]))
        warnings.extend(source["warnings"])
        capabilities.extend(source["capabilities"])

    work_context = {"cuit": "", "estado": "falta"}
    if explicit_work_target is not None:
        work_context.update(
            {
                "cuit": explicit_work_target.get("cuit") or "",
                "cuit_id": explicit_work_target.get("cuit_id") or "",
                "nombre": explicit_work_target.get("nombre"),
                "estado": "confirmado",
            }
        )
    else:
        inferred_work_context = extract_document_work_cuit(all_lines, client)
        work_context.update(inferred_work_context)
    final_work_cuit = digits_only(str(work_context.get("cuit") or ""))
    bound_client = ensure_bound_client(client, cuit=final_work_cuit or None, cuit_id=work_context.get("cuit_id") or None) if final_work_cuit else None
    if final_work_cuit and bound_client is not None:
        resolved_context = resolve_work_cuit_context(bound_client, final_work_cuit)
        work_context.update(resolved_context)
        if explicit_work_target is not None:
            work_context["estado"] = "confirmado"
    if work_context.get("warning"):
        warnings.append(str(work_context.get("warning")))

    counterparty_cuits = collect_counterparty_cuits(
        all_lines,
        work_cuit=final_work_cuit or None,
        explicit_cuit=getattr(args, "cliente_cuit", None),
    )
    matched_profile = select_local_document_profile(
        work_cuit=final_work_cuit or None,
        counterparty_cuits=counterparty_cuits,
        document_kind=PROFILE_DOCUMENT_KIND_COBRO,
        extracted_sources=extracted_sources,
    )
    profile_payload = (
        extract_document_profile_payload(matched_profile, extracted_sources=extracted_sources)
        if matched_profile is not None
        else {}
    )

    generic_date_info = extract_document_date(all_lines, explicit_date=getattr(args, "fecha", None))
    date_info = coalesce_status_info(profile_payload.get("fecha"), generic_date_info)
    if getattr(args, "fecha", None):
        date_info = {"value": parse_web_date(args.fecha, field_name="fecha"), "estado": "confirmado", "source_text": "hint del usuario"}

    generic_comment_info = extract_document_comment(all_lines)
    comment_info = coalesce_status_info(profile_payload.get("comentarios"), generic_comment_info)
    if getattr(args, "comentarios", None):
        comment_info = {"value": str(args.comentarios), "estado": "confirmado", "source_text": "hint del usuario"}

    generic_client_info = pick_document_client(
        all_lines,
        work_cuit=final_work_cuit or None,
        explicit_cuit=getattr(args, "cliente_cuit", None),
        explicit_name=getattr(args, "cliente_nombre", None),
    )
    client_info = merge_client_candidates(profile_payload.get("cliente"), generic_client_info)
    if getattr(args, "cliente_cuit", None) or getattr(args, "cliente_nombre", None):
        client_info["estado"] = "confirmado"
        client_info["source_text"] = "hint del usuario"
    if bound_client is not None and (client_info.get("cuit") or client_info.get("nombre")):
        try:
            client_match = resolve_cliente_match(
                bound_client,
                explicit_id=None,
                cuit=client_info.get("cuit"),
                nombre=client_info.get("nombre"),
            )
            client_info["idclipro"] = str(find_value(client_match, ("idclipro", "id", "idcliente")) or "")
            matched_name = find_value(client_match, ("clipro", "cliente", "nombre"))
            matched_cuit = find_value(client_match, ("cuit",))
            if matched_name:
                client_info["nombre"] = str(matched_name)
            if matched_cuit:
                client_info["cuit"] = digits_only(str(matched_cuit)) or client_info.get("cuit")
            if client_info.get("estado") != "confirmado":
                client_info["estado"] = "inferido"
        except CLIError as exc:
            warnings.append(str(exc))
    else:
        client_info["idclipro"] = None

    generic_movements = parse_structured_movements(all_rows)
    if not generic_movements:
        generic_movements = parse_line_movements(all_lines, bound_client)
    parsed_movements = merge_item_lists(
        profile_payload.get("movimientos", []),
        generic_movements,
        identity=movement_identity,
    )

    centro_match = None
    if bound_client is not None:
        web_client = SOSContadorWebClient(api_client=bound_client, explicit_cuit=final_work_cuit or None)
        centro_match = resolve_centrocosto_match(web_client, getattr(args, "centrocosto", None) or "General")
    centro_info = {
        "id": str(centro_match.get("id")) if centro_match else "",
        "value": centro_match.get("centrocosto") if centro_match else (getattr(args, "centrocosto", None) or "General"),
        "estado": "inferido" if centro_match else ("confirmado" if getattr(args, "centrocosto", None) else "falta"),
    }

    normalized_movements: list[dict[str, Any]] = []
    missing_fields: list[dict[str, Any]] = []
    total = Decimal("0")
    for index, movement in enumerate(parsed_movements, start=1):
        amount_text = parse_optional_decimal_text(movement.get("monto"))
        if not amount_text:
            missing_fields.append(draft_missing_field(f"movimientos[{index}]", "monto", source_text=movement.get("source_text")))
            continue
        total += Decimal(amount_text)
        account_match = resolve_account_match(bound_client, movement.get("cuenta_nombre")) if bound_client is not None else None
        bank_match = resolve_bank_match(bound_client, movement.get("banco_nombre")) if bound_client is not None else None
        movement_type = movement.get("tipo") or "simple"
        normalized = {
            "tipo": movement_type,
            "cuenta_id": str(find_value(account_match, ("id",)) or "") if account_match else "",
            "cuenta_nombre": simplify_account_label(find_value(account_match, ("cuenta", "label", "rubro")) if account_match else movement.get("cuenta_nombre")),
            "banco_id": str(find_value(bank_match, ("id",)) or "") if bank_match else "",
            "banco_nombre": find_value(bank_match, ("bancos",)) if bank_match else movement.get("banco_nombre"),
            "fecha": guess_date_from_text(movement.get("fecha")) or "",
            "numero": digits_only(str(movement.get("numero") or "")),
            "regimen": digits_only(str(movement.get("regimen") or "")),
            "monto": amount_text,
            "source_text": movement.get("source_text"),
            "estado": "inferido",
        }
        if movement_type == "simple" and not normalized["cuenta_id"]:
            missing_fields.append(draft_missing_field(f"movimientos[{index}]", "cuenta_id", source_text=movement.get("source_text")))
            normalized["estado"] = "falta"
        if movement_type == "cheque":
            if not normalized["cuenta_id"]:
                normalized["cuenta_nombre"] = normalized["cuenta_nombre"] or "Valores A Depositar"
                account_match = resolve_account_match(bound_client, normalized["cuenta_nombre"]) if bound_client is not None else None
                normalized["cuenta_id"] = str(find_value(account_match, ("id",)) or "") if account_match else ""
            if not normalized["cuenta_id"]:
                missing_fields.append(draft_missing_field(f"movimientos[{index}]", "cuenta_id", source_text=movement.get("source_text")))
            if not normalized["banco_id"]:
                missing_fields.append(draft_missing_field(f"movimientos[{index}]", "banco_id", source_text=movement.get("source_text")))
            if not normalized["fecha"]:
                missing_fields.append(draft_missing_field(f"movimientos[{index}]", "fecha", source_text=movement.get("source_text")))
            if not normalized["numero"]:
                missing_fields.append(draft_missing_field(f"movimientos[{index}]", "numero", source_text=movement.get("source_text")))
        if movement_type == "retencion":
            if not normalized["fecha"]:
                normalized["fecha"] = guess_date_from_text(date_info.get("value"))
            if not normalized["cuenta_id"]:
                missing_fields.append(draft_missing_field(f"movimientos[{index}]", "cuenta_id", source_text=movement.get("source_text")))
            if not normalized["fecha"]:
                missing_fields.append(draft_missing_field(f"movimientos[{index}]", "fecha", source_text=movement.get("source_text")))
            if not normalized["regimen"]:
                missing_fields.append(draft_missing_field(f"movimientos[{index}]", "regimen", source_text=movement.get("source_text")))
            if movement.get("numero") and not normalized["numero"]:
                missing_fields.append(draft_missing_field(f"movimientos[{index}]", "numero", source_text=movement.get("source_text")))
        if any(item["section"] == f"movimientos[{index}]" for item in missing_fields):
            normalized["estado"] = "falta"
        normalized_movements.append(normalized)

    facturas = merge_item_lists(
        profile_payload.get("facturas", []),
        extract_invoice_candidates(all_lines, all_rows),
        identity=invoice_identity,
    )
    for factura in getattr(args, "factura", None) or []:
        normalized = normalize_comprobante_reference(factura)
        if normalized:
            facturas.append({"comprobante": normalized, "fecha": None, "total": "", "estado": "confirmado", "source_text": "hint del usuario"})
    for venta_id in getattr(args, "venta_id", None) or []:
        facturas.append({"venta_id": str(venta_id), "estado": "confirmado", "source_text": "hint del usuario"})

    if not final_work_cuit:
        missing_fields.append(draft_missing_field("contexto", "cuit_trabajo"))
    if not client_info.get("idclipro"):
        missing_fields.append(draft_missing_field("cliente", "idclipro", source_text=client_info.get("source_text")))
    if not date_info.get("value"):
        missing_fields.append(draft_missing_field("recibo", "fecha", source_text=date_info.get("source_text")))
    if not normalized_movements:
        missing_fields.append(draft_missing_field("recibo", "movimientos"))
    if capabilities:
        warnings.extend(capabilities)

    total_text = format_decimal_string(total)
    total_asociacion = sum((decimal_or_zero(item.get("total")) for item in facturas), Decimal("0"))
    draft = {
        "draft_id": uuid.uuid4().hex[:12],
        "created_at": utc_now_iso(),
        "contexto": {
            "cuit_trabajo": work_context,
            "modo": "cobro_create_and_associate" if facturas else "cobro_create",
            "fuentes": [{"path": source["path"], "kind": source["kind"]} for source in extracted_sources],
            "extraccion": {
                "origen": "perfil_local" if matched_profile is not None else "estandar",
                "perfil": profile_payload.get("metadata"),
            },
        },
        "cliente": {
            "nombre": client_info.get("nombre"),
            "cuit": client_info.get("cuit"),
            "idclipro": client_info.get("idclipro"),
            "estado": client_info.get("estado"),
            "source_text": client_info.get("source_text"),
        },
        "recibo": {
            "fecha": date_info.get("value"),
            "fecha_estado": date_info.get("estado"),
            "fecha_source_text": date_info.get("source_text"),
            "comentarios": comment_info.get("value"),
            "comentarios_estado": comment_info.get("estado"),
            "comentarios_source_text": comment_info.get("source_text"),
            "centro_costo": centro_info.get("value"),
            "centro_costo_id": centro_info.get("id"),
            "centro_costo_estado": centro_info.get("estado"),
            "total": total_text,
            "movimientos": normalized_movements,
        },
        "asociacion": {
            "facturas": facturas,
            "exigir_cuadre_total": True,
            "total": format_decimal_string(total_asociacion),
            "diferencia": format_decimal_string(total - total_asociacion),
        },
        "validacion": {
            "missing_fields": missing_fields,
            "warnings": warnings,
            "resolved_catalogs": {
                "ocr_available": ocr_is_available(),
                "source_capabilities": capabilities,
                "perfil_local": profile_payload.get("metadata"),
            },
            "timings": {
                "draft_seconds": round(time.perf_counter() - start, 3),
            },
        },
    }
    saved_path = save_draft_payload(draft)
    draft["draft_file"] = str(saved_path)
    return draft


def build_cobro_draft_preview(draft: dict[str, Any]) -> dict[str, Any]:
    cliente = draft.get("cliente", {})
    recibo = draft.get("recibo", {})
    contexto = draft.get("contexto", {})
    return {
        "contexto": [
            {
                "CUIT de trabajo": contexto.get("cuit_trabajo", {}).get("cuit"),
                "Contribuyente": contexto.get("cuit_trabajo", {}).get("nombre"),
                "Estado": contexto.get("cuit_trabajo", {}).get("estado"),
            },
            {
                "Cliente": cliente.get("nombre"),
                "CUIT cliente": cliente.get("cuit"),
                "Estado": cliente.get("estado"),
            },
            {
                "Fecha": recibo.get("fecha"),
                "Centro de costo": recibo.get("centro_costo"),
                "Comentarios": recibo.get("comentarios"),
                "Estado fecha": recibo.get("fecha_estado"),
                "Estado centro": recibo.get("centro_costo_estado"),
                "Estado comentarios": recibo.get("comentarios_estado"),
            },
        ],
        "movimientos": [
            {
                "Tipo": item.get("tipo"),
                "Cuenta destino": item.get("cuenta_nombre"),
                "Banco/Régimen": item.get("banco_nombre") or item.get("regimen"),
                "Número": item.get("numero"),
                "Fecha": item.get("fecha"),
                "Monto": item.get("monto"),
                "Estado": item.get("estado"),
            }
            for item in recibo.get("movimientos", [])
            if isinstance(item, dict)
        ],
        "asociacion": [
            {
                "Comprobante": item.get("comprobante") or item.get("venta_id"),
                "Fecha": item.get("fecha"),
                "Total": item.get("total"),
                "Estado": item.get("estado"),
            }
            for item in draft.get("asociacion", {}).get("facturas", [])
            if isinstance(item, dict)
        ],
        "totales": {
            "total_recibo": recibo.get("total"),
            "total_asociacion": draft.get("asociacion", {}).get("total"),
            "diferencia": draft.get("asociacion", {}).get("diferencia"),
        },
    }


def markdown_cell(value: Any) -> str:
    if value in (None, ""):
        return ""
    return str(value).replace("|", r"\|").replace("\n", " ").strip()


def render_markdown_table(rows: list[dict[str, Any]]) -> str:
    if not rows:
        return ""
    columns = list(rows[0].keys())
    header = "| " + " | ".join(columns) + " |"
    separator = "| " + " | ".join("---" for _ in columns) + " |"
    body = [
        "| " + " | ".join(markdown_cell(row.get(column)) for column in columns) + " |"
        for row in rows
    ]
    return "\n".join([header, separator, *body])


def render_business_preview_markdown(preview: dict[str, Any]) -> str:
    sections: list[str] = []
    title = str(preview.get("titulo") or "").strip()
    if title:
        sections.append(f"Previsualizacion `{title}`.")
    contexto_table = render_markdown_table(
        [row for row in preview.get("contexto", []) if isinstance(row, dict)]
    )
    if contexto_table:
        sections.append("**Contexto**\n" + contexto_table)
    for table in preview.get("tablas", []):
        if not isinstance(table, dict):
            continue
        rows = [row for row in table.get("rows", []) if isinstance(row, dict)]
        table_markdown = render_markdown_table(rows)
        if table_markdown:
            sections.append(f"**{table.get('titulo') or 'Detalle'}**\n" + table_markdown)
    totals = preview.get("totales", {})
    if isinstance(totals, dict) and totals:
        sections.append(
            "\n".join(
                f"{key}: {value}"
                for key, value in totals.items()
                if value not in (None, "")
            )
        )
    warnings = preview.get("advertencias", [])
    if isinstance(warnings, list) and warnings:
        sections.append("**Advertencias**\n" + "\n".join(f"- {item}" for item in warnings))
    return "\n\n".join(part for part in sections if part.strip())


def render_cobro_draft_markdown(draft: dict[str, Any]) -> str:
    preview = build_cobro_draft_preview(draft)
    contexto = draft.get("contexto", {})
    cliente = draft.get("cliente", {})
    recibo = draft.get("recibo", {})
    sections: list[str] = []
    sections.append(
        f"Borrador `{draft.get('draft_id')}` para `{draft.get('contexto', {}).get('cuit_trabajo', {}).get('nombre') or draft.get('contexto', {}).get('cuit_trabajo', {}).get('cuit')}`."
    )
    contexto_table = render_markdown_table(
        [
            {
                "CUIT de trabajo": contexto.get("cuit_trabajo", {}).get("cuit"),
                "Contribuyente": contexto.get("cuit_trabajo", {}).get("nombre"),
                "Cliente": cliente.get("nombre"),
                "CUIT cliente": cliente.get("cuit"),
                "Fecha": recibo.get("fecha"),
                "Centro de costo": recibo.get("centro_costo"),
                "Comentarios": recibo.get("comentarios"),
            }
        ]
    )
    sections.append("**Contexto**\n" + contexto_table)
    movimientos_table = render_markdown_table(preview.get("movimientos", []))
    if movimientos_table:
        sections.append("**Movimientos**\n" + movimientos_table)
    asociacion_table = render_markdown_table(preview.get("asociacion", []))
    if asociacion_table:
        sections.append("**Asociacion**\n" + asociacion_table)
    totales = preview.get("totales", {})
    sections.append(
        "\n".join(
            [
                f"Total recibo: {totales.get('total_recibo') or ''}",
                f"Total asociacion: {totales.get('total_asociacion') or ''}",
                f"Diferencia: {totales.get('diferencia') or ''}",
            ]
        )
    )
    missing_fields = draft.get("validacion", {}).get("missing_fields", [])
    if isinstance(missing_fields, list) and missing_fields:
        sections.append(
            "**Campos faltantes**\n"
            + "\n".join(
                f"- {item.get('section')}.{item.get('field')}"
                for item in missing_fields
                if isinstance(item, dict)
            )
        )
    warnings = draft.get("validacion", {}).get("warnings", [])
    if isinstance(warnings, list) and warnings:
        sections.append("**Advertencias**\n" + "\n".join(f"- {item}" for item in warnings))
    return "\n\n".join(part for part in sections if part.strip())


def preview_work_context(bound_client: SOSContadorClient, work_target: dict[str, Any] | None) -> dict[str, Any]:
    session = bound_client.get_session()
    return {
        "CUIT de trabajo": session.cuit,
        "Contribuyente": (work_target or {}).get("nombre") or session.cuit,
    }


def preview_client_identity(client: SOSContadorClient, idclipro: Any) -> dict[str, Any]:
    target_id = str(idclipro or "").strip()
    if not target_id:
        return {"Cliente": "", "CUIT cliente": ""}
    try:
        catalog = get_client_catalog(client)
    except CLIError:
        return {"Cliente": "", "CUIT cliente": "", "idclipro": target_id}
    for item in catalog:
        candidate_id = find_value(item, ("idclipro", "id", "idcliente"))
        if candidate_id is None or str(candidate_id) != target_id:
            continue
        return {
            "Cliente": find_value(item, ("clipro", "cliente", "nombre")),
            "CUIT cliente": digits_only(str(find_value(item, ("cuit",)) or "")),
            "idclipro": target_id,
        }
    return {"Cliente": "", "CUIT cliente": "", "idclipro": target_id}


def preview_account_name(client: SOSContadorClient, account_id: Any) -> str:
    target_id = str(account_id or "").strip()
    if not target_id:
        return ""
    try:
        catalog = get_account_catalog(client)
    except CLIError:
        return target_id
    for item in catalog:
        candidate_id = find_value(item, ("id", "idcuenta", "idcuentacontable"))
        if candidate_id is None or str(candidate_id) != target_id:
            continue
        return simplify_account_label(str(find_value(item, ("label", "cuenta", "rubro")) or target_id))
    return target_id


def build_generic_mutation_preview(
    title: str,
    *,
    bound_client: SOSContadorClient,
    work_target: dict[str, Any] | None,
    detail_rows: list[dict[str, Any]] | None = None,
    detail_title: str = "Detalle",
    warnings: list[str] | None = None,
) -> tuple[dict[str, Any], str]:
    preview = {
        "titulo": title,
        "contexto": [preview_work_context(bound_client, work_target)],
        "tablas": [],
        "totales": {},
        "advertencias": warnings or [],
    }
    if detail_rows:
        preview["tablas"].append({"titulo": detail_title, "rows": detail_rows})
    return preview, render_business_preview_markdown(preview)


def build_payload_summary_rows(body: Any) -> list[dict[str, Any]]:
    if isinstance(body, dict):
        row: dict[str, Any] = {}
        for key, value in body.items():
            if isinstance(value, (dict, list)):
                continue
            row[str(key)] = value
        return [row] if row else []
    if body is None:
        return []
    return [{"Valor": body}]


def build_generic_payload_preview(
    title: str,
    *,
    bound_client: SOSContadorClient,
    work_target: dict[str, Any] | None,
    body: Any,
) -> tuple[dict[str, Any], str]:
    return build_generic_mutation_preview(
        title,
        bound_client=bound_client,
        work_target=work_target,
        detail_rows=build_payload_summary_rows(body),
        detail_title="Detalle",
    )


def build_cobro_or_pago_business_preview(
    title: str,
    *,
    bound_client: SOSContadorClient,
    work_target: dict[str, Any] | None,
    body: dict[str, Any],
) -> tuple[dict[str, Any], str]:
    client_identity = preview_client_identity(bound_client, body.get("idclipro"))
    context_row = preview_work_context(bound_client, work_target)
    context_row.update(
        {
            "Cliente": client_identity.get("Cliente"),
            "CUIT cliente": client_identity.get("CUIT cliente"),
            "Fecha": body.get("fecha"),
        }
    )
    preview = {
        "titulo": title,
        "contexto": [context_row],
        "tablas": [],
        "totales": {},
        "advertencias": [],
    }
    if isinstance(body.get("movimientos"), list):
        rows = [
            {
                "Tipo": item.get("tipo"),
                "Cuenta destino": preview_account_name(bound_client, item.get("cuenta_id")),
                "Banco/Régimen": item.get("banco_id") or item.get("regimen"),
                "Número": item.get("numero"),
                "Fecha": item.get("fecha"),
                "Monto": item.get("monto"),
            }
            for item in body.get("movimientos", [])
            if isinstance(item, dict)
        ]
        preview["tablas"].append({"titulo": "Movimientos", "rows": rows})
        total = sum((decimal_or_zero(item.get("monto")) for item in body.get("movimientos", []) if isinstance(item, dict)), Decimal("0"))
        preview["totales"] = {"Total": format_decimal_string(total)}
    elif isinstance(body.get("imputaciones"), list):
        rows = [
            {
                "Factura/Item": item.get("cuid"),
                "Monto": item.get("fv"),
            }
            for item in body.get("imputaciones", [])
            if isinstance(item, dict)
        ]
        preview["tablas"].append({"titulo": "Imputaciones", "rows": rows})
        total = sum((decimal_or_zero(item.get("fv")) for item in body.get("imputaciones", []) if isinstance(item, dict)), Decimal("0"))
        preview["totales"] = {"Total": format_decimal_string(total)}
    return preview, render_business_preview_markdown(preview)


def build_venta_business_preview(
    *,
    bound_client: SOSContadorClient,
    work_target: dict[str, Any] | None,
    body: dict[str, Any],
) -> tuple[dict[str, Any], str]:
    client_identity = preview_client_identity(bound_client, body.get("idclipro"))
    context_row = preview_work_context(bound_client, work_target)
    context_row.update(
        {
            "Cliente": client_identity.get("Cliente"),
            "CUIT cliente": client_identity.get("CUIT cliente"),
            "Fecha": body.get("fecha"),
            "Comprobante": format_comprobante_text(body.get("letra"), body.get("sucursal"), body.get("numero")),
        }
    )
    rows: list[dict[str, Any]] = []
    total = Decimal("0")
    for item in body.get("productos", []):
        if not isinstance(item, dict):
            continue
        cantidad = decimal_or_zero(item.get("fc"))
        unitario = decimal_or_zero(item.get("fu"))
        bonificacion = decimal_or_zero(item.get("fa"))
        subtotal = (cantidad * unitario) - bonificacion
        total += subtotal
        rows.append(
            {
                "Producto": item.get("id"),
                "Cantidad": format_decimal_string(cantidad),
                "Precio unitario": format_decimal_string(unitario),
                "Bonificación": format_decimal_string(bonificacion),
                "Cuenta": preview_account_name(bound_client, item.get("cuid")),
                "Total": format_decimal_string(subtotal),
            }
        )
    preview = {
        "titulo": "venta_create",
        "contexto": [context_row],
        "tablas": [{"titulo": "Items", "rows": rows}],
        "totales": {"Total estimado": format_decimal_string(total)},
        "advertencias": [],
    }
    return preview, render_business_preview_markdown(preview)


def replace_escaped_once(pattern_text: str, raw_fragment: str | None, replacement: str) -> tuple[str, bool]:
    if not raw_fragment:
        return pattern_text, False
    escaped = re.escape(str(raw_fragment))
    if escaped not in pattern_text:
        return pattern_text, False
    return pattern_text.replace(escaped, replacement, 1), True


def build_cuit_match_pattern(cuit_value: str | None, group_name: str) -> str:
    digits = digits_only(cuit_value or "")
    if len(digits) != 11:
        return rf"(?P<{group_name}>\d{{2}}[- ]?\d{{8}}[- ]?\d)"
    return rf"(?P<{group_name}>{digits[:2]}[- ]?{digits[2:10]}[- ]?{digits[10:]})"


def build_amount_match_pattern(group_name: str) -> str:
    return rf"(?P<{group_name}>\d{{1,3}}(?:[.,]\d{{3}})*(?:[.,]\d{{2}})|\d+(?:[.,]\d{{2}}))"


def build_date_match_pattern(group_name: str) -> str:
    return rf"(?P<{group_name}>\d{{1,2}}[/-]\d{{1,2}}[/-]\d{{2,4}})"


def build_digits_match_pattern(group_name: str, *, min_length: int = 1) -> str:
    return rf"(?P<{group_name}>\d{{{min_length},12}})"


def source_based_pattern(
    source_text: str | None,
    *,
    replacements: list[tuple[str | None, str]],
    fallback_pattern: str,
) -> str:
    if not source_text:
        return fallback_pattern
    pattern_text = re.escape(source_text)
    changed = False
    for raw_fragment, replacement in sorted(replacements, key=lambda item: len(str(item[0] or "")), reverse=True):
        pattern_text, applied = replace_escaped_once(pattern_text, raw_fragment, replacement)
        changed = changed or applied
    return pattern_text if changed else fallback_pattern


def build_profile_field_rules_from_draft(draft: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    recibo = draft.get("recibo", {})
    cliente = draft.get("cliente", {})
    field_rules: dict[str, list[dict[str, Any]]] = {}

    fecha_source = str(recibo.get("fecha_source_text") or "")
    fecha_candidate = next(iter(date_candidates_in_line(fecha_source)), None)
    field_rules["fecha"] = [
        {
            "scope": PROFILE_SCOPE_LINE,
            "pattern": source_based_pattern(
                fecha_source,
                replacements=[(fecha_candidate, build_date_match_pattern("fecha"))],
                fallback_pattern=rf"(?i)\bfecha\b[^\n]{{0,40}}{build_date_match_pattern('fecha')}",
            ),
            "value_group": "fecha",
            "transform": "date",
            "flags": ["ignorecase"],
        }
    ]

    comment_source = str(recibo.get("comentarios_source_text") or "")
    field_rules["comentarios"] = [
        {
            "scope": PROFILE_SCOPE_LINE,
            "pattern": source_based_pattern(
                comment_source,
                replacements=[(recibo.get("comentarios"), r"(?P<comentarios>.+)")],
                fallback_pattern=r"(?i)(?P<comentarios>orden de pago[^\n]*)",
            ),
            "value_group": "comentarios",
            "transform": "normalize-space",
            "flags": ["ignorecase"],
        }
    ]

    field_rules["cliente_cuit"] = [
        {
            "scope": PROFILE_SCOPE_LINE,
            "pattern": build_cuit_match_pattern(cliente.get("cuit"), "cliente_cuit"),
            "value_group": "cliente_cuit",
            "transform": "digits",
        }
    ]
    if cliente.get("nombre"):
        client_pattern = re.escape(str(cliente.get("nombre"))).replace(r"\ ", r"\s+")
        field_rules["cliente_nombre"] = [
            {
                "scope": PROFILE_SCOPE_LINE,
                "pattern": rf"(?P<cliente_nombre>{client_pattern})",
                "value_group": "cliente_nombre",
                "transform": "normalize-space",
                "flags": ["ignorecase"],
            }
        ]
    return field_rules


def build_profile_movement_rules_from_draft(draft: dict[str, Any]) -> list[dict[str, Any]]:
    rules: list[dict[str, Any]] = []
    seen: set[tuple[str, str]] = set()
    for movement in draft.get("recibo", {}).get("movimientos", []):
        if not isinstance(movement, dict):
            continue
        movement_type = str(movement.get("tipo") or "").strip() or "simple"
        source_text = str(movement.get("source_text") or "")
        date_fragment = next(iter(date_candidates_in_line(source_text)), None)
        amount_fragment = next(iter(amount_candidates_in_line(source_text)[-1:]), None)
        number_fragment = digits_only(movement.get("numero") or "") or next(iter(long_numeric_tokens(source_text)), None)
        regimen_fragment = digits_only(movement.get("regimen") or "")
        if movement_type == "cheque":
            fallback = (
                rf"(?i)(?:chp|cheque|e-?cheq).*?{build_date_match_pattern('fecha')}"
                rf".*?{build_amount_match_pattern('monto')}"
                rf".*?(?:nro\.?\s*)?{build_digits_match_pattern('numero', min_length=5)}"
            )
            pattern = source_based_pattern(
                source_text,
                replacements=[
                    (date_fragment, build_date_match_pattern("fecha")),
                    (amount_fragment, build_amount_match_pattern("monto")),
                    (number_fragment, build_digits_match_pattern("numero", min_length=5)),
                ],
                fallback_pattern=fallback,
            )
            rule = {
                "tipo": "cheque",
                "scope": PROFILE_SCOPE_LINE,
                "pattern": pattern,
                "groups": {
                    "fecha": "fecha",
                    "monto": "monto",
                    "numero": "numero",
                },
                "transforms": {
                    "fecha": "date",
                    "monto": "decimal",
                    "numero": "digits",
                },
                "defaults": {
                    "cuenta_nombre": movement.get("cuenta_nombre") or "Valores A Depositar",
                    "banco_nombre": movement.get("banco_nombre"),
                },
                "flags": ["ignorecase"],
            }
        elif movement_type == "retencion":
            fallback = (
                rf"(?i).*?(?:rg[.\s-]*|sicore\D{{0,5}})(?P<regimen>\d{{2,4}})"
                rf".*?{build_amount_match_pattern('monto')}"
            )
            pattern = source_based_pattern(
                source_text,
                replacements=[
                    (regimen_fragment, r"(?P<regimen>\d{2,4})"),
                    (amount_fragment, build_amount_match_pattern("monto")),
                    (number_fragment, build_digits_match_pattern("numero", min_length=1)),
                ],
                fallback_pattern=fallback,
            )
            groups: dict[str, str] = {"monto": "monto"}
            if "(?P<regimen>" in pattern:
                groups["regimen"] = "regimen"
            if "(?P<numero>" in pattern:
                groups["numero"] = "numero"
            rule = {
                "tipo": "retencion",
                "scope": PROFILE_SCOPE_LINE,
                "pattern": pattern,
                "groups": groups,
                "transforms": {
                    "monto": "decimal",
                    "numero": "digits",
                    "regimen": "digits",
                },
                "defaults": {
                    "cuenta_nombre": movement.get("cuenta_nombre"),
                    "fecha": PROFILE_DOCUMENT_DATE_TOKEN,
                },
                "flags": ["ignorecase"],
            }
        else:
            continue
        key = (movement_type, str(rule.get("pattern")))
        if key in seen:
            continue
        seen.add(key)
        rules.append(rule)
    return rules


def build_profile_factura_rules_from_draft(draft: dict[str, Any]) -> list[dict[str, Any]]:
    rules: list[dict[str, Any]] = []
    seen: set[str] = set()
    for factura in draft.get("asociacion", {}).get("facturas", []):
        if not isinstance(factura, dict):
            continue
        if factura.get("venta_id") and not factura.get("comprobante"):
            continue
        source_text = str(factura.get("source_text") or "")
        date_fragment = next(iter(date_candidates_in_line(source_text)), None)
        amount_fragment = next(iter(amount_candidates_in_line(source_text)[-1:]), None)
        comprobante_fragment = normalize_comprobante_reference(source_text) or factura.get("comprobante")
        fallback = (
            rf"(?i){build_date_match_pattern('fecha')}.*?"
            rf"(?P<comprobante>(?:[A-Z]{{1,3}}\s*[A-Z]?\s*\d{{1,5}}[-/]\d{{1,8}}|[A-Z]-\d{{4}}-\d{{8}})).*?"
            rf"{build_amount_match_pattern('total')}"
        )
        pattern = source_based_pattern(
            source_text,
            replacements=[
                (date_fragment, build_date_match_pattern("fecha")),
                (comprobante_fragment, r"(?P<comprobante>(?:[A-Z]{1,3}\s*[A-Z]?\s*\d{1,5}[-/]\d{1,8}|[A-Z]-\d{4}-\d{8}))"),
                (amount_fragment, build_amount_match_pattern("total")),
            ],
            fallback_pattern=fallback,
        )
        if pattern in seen:
            continue
        seen.add(pattern)
        rules.append(
            {
                "scope": PROFILE_SCOPE_LINE,
                "pattern": pattern,
                "groups": {
                    "fecha": "fecha",
                    "comprobante": "comprobante",
                    "total": "total",
                },
                "transforms": {
                    "fecha": "date",
                    "comprobante": "comprobante",
                    "total": "decimal",
                },
                "flags": ["ignorecase"],
            }
        )
    return rules


def build_local_profile_manifest(
    *,
    draft: dict[str, Any],
    name: str,
    profile_id: str,
    work_cuit: str,
    counterparty_cuit: str | None,
    document_kind: str,
    sources: list[Path],
) -> dict[str, Any]:
    client_name = str(draft.get("cliente", {}).get("nombre") or "").strip()
    comment_text = normalize_search_text(draft.get("recibo", {}).get("comentarios"))
    required_text: list[str] = []
    if client_name:
        required_text.append(client_name)
    if "orden de pago" in comment_text:
        required_text.append("orden de pago")
    return {
        "profile_id": profile_id,
        "name": name,
        "version": 1,
        "work_cuit": work_cuit,
        "counterparty_cuit": counterparty_cuit or "",
        "document_kind": document_kind,
        "priority": 100,
        "required_text": required_text,
        "filename_contains": [],
        "created_at": utc_now_iso(),
        "source_paths": [str(path) for path in sources],
        "created_from_draft_id": draft.get("draft_id"),
    }


def build_local_profile_rules(draft: dict[str, Any]) -> dict[str, Any]:
    return {
        "fields": build_profile_field_rules_from_draft(draft),
        "movimientos": build_profile_movement_rules_from_draft(draft),
        "facturas": build_profile_factura_rules_from_draft(draft),
    }


def resolve_profile_source_paths(args: argparse.Namespace, draft: dict[str, Any]) -> list[Path]:
    explicit_sources = [Path(path) for path in (getattr(args, "source", None) or [])]
    if explicit_sources:
        return explicit_sources
    sources: list[Path] = []
    for item in draft.get("contexto", {}).get("fuentes", []):
        if not isinstance(item, dict):
            continue
        raw_path = item.get("path")
        if raw_path:
            sources.append(Path(str(raw_path)))
    return sources


def command_cobro_profile_create(args: argparse.Namespace, client: SOSContadorClient) -> Any:
    draft = load_draft_payload(
        draft_id=getattr(args, "draft_id", None),
        draft_file=getattr(args, "draft_file", None),
    )
    work_cuit = digits_only(
        getattr(args, "cuit_trabajo", None)
        or draft.get("contexto", {}).get("cuit_trabajo", {}).get("cuit")
        or ""
    )
    if not work_cuit:
        raise CLIError("No se pudo determinar la CUIT de trabajo para crear el perfil local.")
    counterparty_cuit = digits_only(
        getattr(args, "counterparty_cuit", None)
        or draft.get("cliente", {}).get("cuit")
        or ""
    ) or None
    document_kind = str(getattr(args, "document_kind", None) or PROFILE_DOCUMENT_KIND_COBRO)
    default_name = str(draft.get("cliente", {}).get("nombre") or "Perfil local de cobranza").strip()
    profile_name = str(getattr(args, "name", None) or default_name)
    profile_id = str(getattr(args, "profile_id", None) or slugify(profile_name)).strip()
    if not profile_id:
        raise CLIError("Debe indicar un nombre o profile_id utilizable para el perfil local.")

    sources = resolve_profile_source_paths(args, draft)
    manifest = build_local_profile_manifest(
        draft=draft,
        name=profile_name,
        profile_id=profile_id,
        work_cuit=work_cuit,
        counterparty_cuit=counterparty_cuit,
        document_kind=document_kind,
        sources=sources,
    )
    rules = build_local_profile_rules(draft)
    destination = profile_directory(
        work_cuit=work_cuit,
        counterparty_cuit=counterparty_cuit,
        document_kind=document_kind,
        profile_id=profile_id,
    )
    manifest_path = destination / "manifest.json"
    rules_path = destination / "rules.json"
    save_json_file(manifest_path, manifest)
    save_json_file(rules_path, rules)
    payload = {
        "profile_id": profile_id,
        "path": str(destination),
        "manifest_file": str(manifest_path),
        "rules_file": str(rules_path),
        "work_cuit": work_cuit,
        "counterparty_cuit": counterparty_cuit,
        "document_kind": document_kind,
        "draft_id": draft.get("draft_id"),
        "movimiento_rules": len(rules.get("movimientos") or []),
        "factura_rules": len(rules.get("facturas") or []),
    }
    if getattr(args, "json_out", None):
        json_path = Path(args.json_out)
        save_json_file(json_path, payload)
        payload["json_written_to"] = str(json_path)
    return payload


def normalize_cobro_movimientos(raw_movimientos: Any) -> list[dict[str, str]]:
    if not isinstance(raw_movimientos, list) or not raw_movimientos:
        raise CLIError("Los movimientos deben ser una lista JSON no vacia.")

    normalized_items: list[dict[str, str]] = []
    valid_types = {"simple", "cheque", "retencion"}
    for index, item in enumerate(raw_movimientos, start=1):
        if not isinstance(item, dict):
            raise CLIError(f"El movimiento #{index} debe ser un objeto JSON.")
        tipo = str(item.get("tipo") or "").strip().lower()
        if tipo not in valid_types:
            raise CLIError(
                f"El movimiento #{index} debe indicar un 'tipo' valido ({', '.join(sorted(valid_types))})."
            )
        cuenta_id = str(item.get("cuenta_id") or "").strip()
        if not cuenta_id:
            raise CLIError(f"El movimiento #{index} debe indicar 'cuenta_id'.")

        movimiento = {
            "cuid": cuenta_id,
            "pf": "",
            "pn": "",
            "pi": "",
            "pa": "",
            "pm": "",
            "pr": "",
            "pb": "",
            "pc": "",
            "fv": parse_required_decimal_text(item.get("monto"), field_name=f"monto del movimiento #{index}"),
        }

        if tipo == "simple":
            normalized_items.append(movimiento)
            continue

        movimiento["pf"] = parse_web_date(item.get("fecha"), field_name=f"fecha del movimiento #{index}")

        if tipo == "cheque":
            banco_id = str(item.get("banco_id") or "").strip()
            numero = parse_numeric_text(item.get("numero"), field_name=f"numero del cheque #{index}")
            if not banco_id:
                raise CLIError(f"El movimiento #{index} de tipo cheque debe indicar 'banco_id'.")
            movimiento["pb"] = banco_id
            movimiento["pn"] = numero
            normalized_items.append(movimiento)
            continue

        regimen = str(item.get("regimen") or "").strip()
        if not regimen:
            raise CLIError(f"El movimiento #{index} de tipo retencion debe indicar 'regimen'.")
        movimiento["pr"] = regimen
        movimiento["pn"] = parse_numeric_text(
            item.get("numero"),
            field_name=f"numero de la retencion #{index}",
            required=False,
        )
        normalized_items.append(movimiento)

    return normalized_items


def validate_cobro_web_body(body: Any) -> dict[str, Any]:
    if not isinstance(body, dict):
        raise CLIError("El cuerpo del cobro detallado debe ser un objeto JSON.")
    if not body.get("idclipro"):
        raise CLIError("El cuerpo del cobro detallado debe incluir 'idclipro'.")
    if not body.get("fecha"):
        raise CLIError("El cuerpo del cobro detallado debe incluir 'fecha'.")
    movimientos = body.get("movimientos")
    if not isinstance(movimientos, list) or not movimientos:
        raise CLIError("El cuerpo del cobro detallado debe incluir 'movimientos' no vacios.")
    for index, item in enumerate(movimientos, start=1):
        if not isinstance(item, dict):
            raise CLIError(f"El movimiento web #{index} debe ser un objeto JSON.")
        if not item.get("cuid") or item.get("fv") in (None, ""):
            raise CLIError(f"El movimiento web #{index} debe incluir 'cuid' y 'fv'.")
    validated = dict(body)
    validated.setdefault("id", "0")
    validated.setdefault("idtipo", "12")
    validated["fecha"] = parse_web_date(validated.get("fecha"), field_name="fecha")
    validated["fechaiva"] = parse_web_date(validated.get("fechaiva") or validated["fecha"], field_name="fechaiva")
    validated["fechadesde"] = parse_web_date(
        validated.get("fechadesde") or validated["fecha"],
        field_name="fechadesde",
    )
    validated["fechahasta"] = parse_web_date(
        validated.get("fechahasta") or validated["fecha"],
        field_name="fechahasta",
    )
    validated.setdefault("numero", "")
    validated.setdefault("idcentrocosto", "")
    validated.setdefault("observaciones", "")
    validated.setdefault("moneda", "ARS")
    validated.setdefault("tipocambio", "1")
    validated.setdefault("nuevocentro", "")
    return validated


def args_use_draft_source(args: argparse.Namespace) -> bool:
    return bool(getattr(args, "draft_id", None) or getattr(args, "draft_file", None))


def draft_has_missing_fields(draft: dict[str, Any]) -> bool:
    missing = draft.get("validacion", {}).get("missing_fields", [])
    return isinstance(missing, list) and bool(missing)


def build_cobro_detail_body_from_draft(draft: dict[str, Any], client: SOSContadorClient) -> tuple[SOSContadorClient, dict[str, Any]]:
    if draft_has_missing_fields(draft):
        raise CLIError("El borrador todavia tiene campos faltantes. Complete los faltantes antes de cargar el recibo.")
    work_cuit = digits_only(str(draft.get("contexto", {}).get("cuit_trabajo", {}).get("cuit") or ""))
    if not work_cuit:
        raise CLIError("El borrador no incluye una CUIT de trabajo explicita.")
    bound_client = ensure_bound_client(client, cuit=work_cuit)
    web_client = SOSContadorWebClient(api_client=bound_client, explicit_cuit=work_cuit)
    cliente = draft.get("cliente", {})
    recibo = draft.get("recibo", {})
    fecha = str(recibo.get("fecha") or "").strip()
    if not fecha:
        raise CLIError("El borrador no incluye la fecha del recibo.")
    idclipro = str(cliente.get("idclipro") or "").strip()
    if not idclipro:
        idclipro = resolve_cliente_id(
            bound_client,
            explicit_id=None,
            cuit=cliente.get("cuit"),
            nombre=cliente.get("nombre"),
        )
    idcentrocosto = str(recibo.get("centro_costo_id") or "").strip()
    if not idcentrocosto:
        idcentrocosto = resolve_centrocosto_id(
            web_client,
            explicit_id=None,
            centrocosto=recibo.get("centro_costo"),
        )
    numero = str(recibo.get("numero") or "").strip()
    if not numero:
        numero = str(web_client.get_ultimo_comprobante_numero(idtipo_operacion=12) + 1)
    movimientos = []
    for index, item in enumerate(recibo.get("movimientos", []), start=1):
        if not isinstance(item, dict):
            raise CLIError(f"Movimiento invalido en borrador #{index}.")
        movement = {
            "tipo": item.get("tipo"),
            "cuenta_id": item.get("cuenta_id"),
            "monto": item.get("monto"),
        }
        if item.get("fecha"):
            movement["fecha"] = item.get("fecha")
        if item.get("numero"):
            movement["numero"] = item.get("numero")
        if item.get("banco_id"):
            movement["banco_id"] = item.get("banco_id")
        if item.get("regimen"):
            movement["regimen"] = item.get("regimen")
        movimientos.append(movement)
    body = validate_cobro_web_body(
        {
            "id": "0",
            "idtipo": "12",
            "idclipro": str(idclipro),
            "fecha": fecha,
            "fechaiva": fecha,
            "fechadesde": fecha,
            "fechahasta": fecha,
            "numero": parse_numeric_text(numero, field_name="numero del recibo"),
            "idcentrocosto": str(idcentrocosto),
            "observaciones": str(recibo.get("comentarios") or ""),
            "moneda": "ARS",
            "tipocambio": "1",
            "nuevocentro": "",
            "movimientos": normalize_cobro_movimientos(movimientos),
        }
    )
    return bound_client, body


def resolve_saved_cobro_identity(
    *,
    save_payload: dict[str, Any],
    body: dict[str, Any],
    client: SOSContadorClient,
) -> dict[str, Any]:
    cobro_id = str(find_value(save_payload, ("id", "idcomprobante")) or "").strip()
    if cobro_id:
        web_client = SOSContadorWebClient(api_client=client, explicit_cuit=client.get_session().cuit)
        payload = web_client.get_comprobante_historia(idcomprobante=cobro_id, idtipo_operacion=12)
        return {
            "id": cobro_id,
            "idclipro": str(payload.get("idclipro") or body.get("idclipro") or ""),
            "numero": str(payload.get("numero") or body.get("numero") or ""),
            "comprobante": format_comprobante_text(payload.get("letra"), payload.get("sucursal"), payload.get("numero")),
        }
    lookup_args = argparse.Namespace(
        id=None,
        recibo=str(body.get("numero") or ""),
        comprobante=None,
        fecha=str(body.get("fecha") or ""),
        desde=None,
        hasta=None,
        cliente_cuit=None,
        cliente_nombre=None,
        registros=200,
    )
    match = resolve_cobro_match_from_lookup(lookup_args, client)
    return {
        "id": str(match.get("id") or ""),
        "idclipro": str(match.get("idclipro") or body.get("idclipro") or ""),
        "numero": str(match.get("numero") or body.get("numero") or ""),
        "comprobante": str(match.get("comprobante") or ""),
    }


def build_cobro_detail_body(args: argparse.Namespace, client: SOSContadorClient) -> dict[str, Any]:
    explicit_body = load_json_source(getattr(args, "body_json", None), getattr(args, "body_file", None))
    if explicit_body is not None:
        if explicit_body.get("movimientos") is None:
            raise CLIError("Para el alta detallada de cobros, el cuerpo debe incluir 'movimientos'.")
        return validate_cobro_web_body(explicit_body)

    if not args.fecha:
        raise CLIError("Debe indicar --fecha para crear un cobro detallado.")
    raw_movimientos = load_json_source(getattr(args, "movimientos_json", None), getattr(args, "movimientos_file", None))
    movimientos = normalize_cobro_movimientos(raw_movimientos)
    web_client = SOSContadorWebClient(api_client=client)
    idclipro = resolve_cliente_id(
        client,
        explicit_id=args.idclipro,
        cuit=getattr(args, "cliente_cuit", None),
        nombre=getattr(args, "cliente_nombre", None),
    )
    idcentrocosto = resolve_centrocosto_id(
        web_client,
        explicit_id=getattr(args, "idcentrocosto", None),
        centrocosto=getattr(args, "centrocosto", None),
    )
    numero = str(getattr(args, "numero", "") or "").strip()
    if not numero:
        numero = str(web_client.get_ultimo_comprobante_numero(idtipo_operacion=12) + 1)

    fecha = parse_web_date(args.fecha, field_name="fecha")
    numero = parse_numeric_text(numero, field_name="numero del recibo")
    return {
        "id": "0",
        "idtipo": "12",
        "idclipro": str(idclipro),
        "fecha": fecha,
        "fechaiva": fecha,
        "fechadesde": fecha,
        "fechahasta": fecha,
        "numero": numero,
        "idcentrocosto": str(idcentrocosto),
        "observaciones": coalesce_cobro_observaciones(args),
        "moneda": "ARS",
        "tipocambio": "1",
        "nuevocentro": "",
        "movimientos": movimientos,
    }


def build_cobro_create_request(args: argparse.Namespace, client: SOSContadorClient) -> tuple[str, str, dict[str, Any]]:
    if args_use_draft_source(args):
        draft = load_draft_payload(
            draft_id=getattr(args, "draft_id", None),
            draft_file=getattr(args, "draft_file", None),
        )
        _, body = build_cobro_detail_body_from_draft(draft, client)
        return "web-session", "back/comprobante_altamodi.asp", body
    has_movimientos_source = any(
        (
            getattr(args, "movimientos_json", None),
            getattr(args, "movimientos_file", None),
        )
    )
    explicit_body = load_json_source(getattr(args, "body_json", None), getattr(args, "body_file", None))
    explicit_body_has_movimientos = isinstance(explicit_body, dict) and explicit_body.get("movimientos") is not None

    if has_movimientos_source or explicit_body_has_movimientos:
        if any((getattr(args, "imputaciones_json", None), getattr(args, "imputaciones_file", None))):
            raise CLIError("No combine --movimientos-json/--movimientos-file con --imputaciones-json/--imputaciones-file.")
        if getattr(args, "idcuenta", None):
            raise CLIError("No use --idcuenta cuando crea un cobro detallado con movimientos.")
        if getattr(args, "idprovinciaiibb", None) or getattr(args, "provincia", None):
            raise CLIError("No use --idprovinciaiibb/--provincia cuando crea un cobro detallado con movimientos.")
        body = build_cobro_detail_body(args, client)
        return "web-session", "back/comprobante_altamodi.asp", body

    body = build_cobro_or_pago_body(args, client)
    return "api", "cobro", body


def build_cobro_or_pago_body(args: argparse.Namespace, client: SOSContadorClient) -> dict[str, Any]:
    body = load_json_source(getattr(args, "body_json", None), getattr(args, "body_file", None))
    if body is not None:
        validate_imputaciones(body.get("imputaciones"))
        return body

    if not args.fecha:
        raise CLIError("Debe indicar --fecha o proveer el cuerpo completo con --body-json / --body-file.")
    imputaciones = load_json_source(getattr(args, "imputaciones_json", None), getattr(args, "imputaciones_file", None))
    validate_imputaciones(imputaciones)
    idclipro = resolve_cliente_id(
        client,
        explicit_id=args.idclipro,
        cuit=getattr(args, "cliente_cuit", None),
        nombre=getattr(args, "cliente_nombre", None),
    )
    idprovinciaiibb = resolve_provincia_id(
        client,
        explicit_id=args.idprovinciaiibb,
        provincia=getattr(args, "provincia", None),
    )
    if not args.idcuenta:
        raise CLIError("Debe indicar --idcuenta explicitamente.")
    if not args.idcentrocosto:
        raise CLIError("Debe indicar --idcentrocosto explicitamente.")

    return {
        "fecha": args.fecha,
        "idclipro": str(idclipro),
        "idcuenta": str(args.idcuenta),
        "idprovinciaiibb": str(idprovinciaiibb),
        "idcentrocosto": str(args.idcentrocosto),
        "memo": args.memo or "",
        "referencia": args.referencia or "",
        "imputaciones": imputaciones,
    }


def default_period_bounds_for_date(fecha_ddmmyyyy: str) -> tuple[str, str]:
    date_value = datetime.datetime.strptime(fecha_ddmmyyyy, "%d/%m/%Y")
    first_day = date_value.replace(day=1)
    next_month = (first_day + datetime.timedelta(days=32)).replace(day=1)
    last_day = next_month - datetime.timedelta(days=1)
    return first_day.strftime("%d/%m/%Y"), last_day.strftime("%d/%m/%Y")


def normalize_venta_productos(raw_productos: Any) -> list[dict[str, Any]]:
    if not isinstance(raw_productos, list) or not raw_productos:
        raise CLIError("Los productos de la venta deben ser una lista JSON no vacia.")
    normalized: list[dict[str, Any]] = []
    for index, item in enumerate(raw_productos, start=1):
        if not isinstance(item, dict):
            raise CLIError(f"El producto #{index} debe ser un objeto JSON.")
        required_fields = ("id", "u", "fc", "fu", "fa", "fi", "cuid")
        missing = [field for field in required_fields if item.get(field) in (None, "")]
        if missing:
            raise CLIError(f"El producto #{index} debe incluir {', '.join(missing)}.")
        normalized.append(dict(item))
    return normalized


def normalize_venta_imputa(raw_imputa: Any) -> dict[str, Any]:
    if raw_imputa is None:
        return {"imputa": []}
    if not isinstance(raw_imputa, dict):
        raise CLIError("La imputacion de la venta debe ser un objeto JSON.")
    imputa_items = raw_imputa.get("imputa")
    if imputa_items is None:
        return {"imputa": []}
    if not isinstance(imputa_items, list):
        raise CLIError("La clave 'imputa' de la venta debe ser una lista.")
    return dict(raw_imputa)


def validate_venta_web_body(body: Any) -> dict[str, Any]:
    if not isinstance(body, dict):
        raise CLIError("El cuerpo de la venta debe ser un objeto JSON.")
    required_fields = ("idclipro", "fecha", "fcncnd", "idtipo", "letra", "sucursal", "numero")
    missing = [field for field in required_fields if body.get(field) in (None, "")]
    if missing:
        raise CLIError(f"El cuerpo de la venta debe incluir {', '.join(missing)}.")
    productos = normalize_venta_productos(body.get("productos"))
    validated = dict(body)
    validated["id"] = str(validated.get("id") or "0")
    validated["idtipo"] = str(validated.get("idtipo") or "2")
    validated["fecha"] = parse_web_date(validated.get("fecha"), field_name="fecha")
    validated["fechaiva"] = parse_web_date(validated.get("fechaiva") or validated["fecha"], field_name="fechaiva")
    fecha_desde, fecha_hasta = default_period_bounds_for_date(validated["fecha"])
    validated["fechadesde"] = parse_web_date(validated.get("fechadesde") or fecha_desde, field_name="fechadesde")
    validated["fechahasta"] = parse_web_date(validated.get("fechahasta") or fecha_hasta, field_name="fechahasta")
    validated["numero"] = parse_numeric_text(validated.get("numero"), field_name="numero de la venta")
    validated["numero_hasta"] = parse_numeric_text(
        validated.get("numero_hasta") or validated["numero"],
        field_name="numero_hasta",
    )
    validated.setdefault("idcuenta", "")
    validated.setdefault("idprovinciaiibb", "")
    validated.setdefault("idcentrocosto", "")
    validated.setdefault("codactividad", "")
    validated.setdefault("observaciones", "")
    validated.setdefault("moneda", "ARS")
    validated.setdefault("tipocambio", "1")
    validated.setdefault("nuevocentro", "")
    validated.setdefault("cuentacobropago", "0")
    validated["obtienecae"] = str(validated.get("obtienecae", "false")).lower()
    validated["imputa"] = normalize_venta_imputa(validated.get("imputa"))
    validated["productos"] = productos
    return validated


def build_venta_create_body(args: argparse.Namespace, client: SOSContadorClient) -> dict[str, Any]:
    explicit_body = load_json_source(getattr(args, "body_json", None), getattr(args, "body_file", None))
    if explicit_body is not None:
        return validate_venta_web_body(explicit_body)

    if not args.fecha:
        raise CLIError("Debe indicar --fecha para crear una venta.")
    if not args.letra:
        raise CLIError("Debe indicar --letra para crear una venta.")
    if not args.sucursal:
        raise CLIError("Debe indicar --sucursal para crear una venta.")
    if not args.idcuenta:
        raise CLIError("Debe indicar --idcuenta para crear una venta.")

    raw_productos = load_json_source(getattr(args, "productos_json", None), getattr(args, "productos_file", None))
    productos = normalize_venta_productos(raw_productos)
    raw_imputa = load_json_source(getattr(args, "imputa_json", None), getattr(args, "imputa_file", None))
    imputa = normalize_venta_imputa(raw_imputa)

    web_client = SOSContadorWebClient(api_client=client)
    idclipro = resolve_cliente_id(
        client,
        explicit_id=args.idclipro,
        cuit=getattr(args, "cliente_cuit", None),
        nombre=getattr(args, "cliente_nombre", None),
    )
    idprovinciaiibb = resolve_provincia_id(
        client,
        explicit_id=getattr(args, "idprovinciaiibb", None),
        provincia=getattr(args, "provincia", None),
    )
    idcentrocosto = resolve_centrocosto_id(
        web_client,
        explicit_id=getattr(args, "idcentrocosto", None),
        centrocosto=getattr(args, "centrocosto", None),
    )
    fecha = parse_web_date(args.fecha, field_name="fecha")
    fechadesde, fechahasta = default_period_bounds_for_date(fecha)
    numero = str(getattr(args, "numero", "") or "").strip()
    if not numero:
        numero = str(web_client.get_ultimo_comprobante_numero(idtipo_operacion=2) + 1)

    return validate_venta_web_body(
        {
            "id": "0",
            "fcncnd": getattr(args, "fcncnd", None) or "F",
            "idtipo": "2",
            "fecha": fecha,
            "idclipro": str(idclipro),
            "letra": str(args.letra),
            "sucursal": str(args.sucursal),
            "numero": numero,
            "numero_hasta": str(getattr(args, "numero_hasta", None) or numero),
            "fechaiva": fecha,
            "fechadesde": fechadesde,
            "fechahasta": fechahasta,
            "codactividad": str(getattr(args, "codactividad", None) or ""),
            "idcuenta": str(args.idcuenta),
            "idprovinciaiibb": str(idprovinciaiibb),
            "idcentrocosto": str(idcentrocosto),
            "observaciones": str(getattr(args, "observaciones", None) or ""),
            "moneda": str(getattr(args, "moneda", None) or "ARS"),
            "tipocambio": str(getattr(args, "tipocambio", None) or "1"),
            "nuevocentro": str(getattr(args, "centrocosto", None) or "General"),
            "cuentacobropago": str(getattr(args, "cuentacobropago", None) or "0"),
            "obtienecae": "true" if getattr(args, "obtienecae", False) else "false",
            "imputa": imputa,
            "productos": productos,
        }
    )


def parse_human_date(value: str) -> str:
    for fmt in ("%d/%m/%Y", "%d/%m/%y"):
        try:
            return datetime.datetime.strptime(value, fmt).strftime("%d/%m/%Y")
        except ValueError:
            continue
    raise CLIError(f"Fecha invalida '{value}'. Use formato DD/MM/YYYY.")


def parse_isoish_date(value: str | None) -> str | None:
    if not value:
        return None
    text = str(value).strip()
    for fmt in ("%Y-%m-%dT%H:%M:%S.%f", "%Y-%m-%d", "%d/%m/%Y", "%d/%m/%y"):
        try:
            return datetime.datetime.strptime(text, fmt).strftime("%Y-%m-%d")
        except ValueError:
            continue
    return text


def decimal_or_zero(value: Any) -> Decimal:
    if value in (None, ""):
        return Decimal("0")
    try:
        return Decimal(str(value))
    except (InvalidOperation, ValueError):
        return Decimal("0")


def format_decimal_string(value: Decimal) -> str:
    return format(value.quantize(Decimal("0.01")), "f")


def format_comprobante_text(letra: Any, sucursal: Any, numero: Any) -> str | None:
    if not letra and sucursal in (None, "") and numero in (None, ""):
        return None
    letra_text = str(letra or "").strip() or "?"
    sucursal_num = digits_only(str(sucursal or "0")) or "0"
    numero_num = digits_only(str(numero or "0")) or "0"
    return f"{letra_text}-{int(sucursal_num):04d}-{int(numero_num):08d}"


def summarize_web_comprobante(item: dict[str, Any]) -> dict[str, Any]:
    return {
        "id": item.get("id"),
        "fecha": item.get("fechaiva") or item.get("fecha"),
        "comprobante": item.get("comprobante"),
        "numero": item.get("numero"),
        "cliente": item.get("clipro"),
        "cuit": item.get("cuitreceptor"),
        "idclipro": item.get("idclipro"),
        "total": item.get("total"),
        "cuenta": item.get("cuenta"),
        "centro_costo": item.get("rubro") or item.get("centrocosto"),
        "usuario": item.get("usuario"),
        "origen": item.get("origen"),
    }


def summarize_cobro_historia_item(item: dict[str, Any]) -> dict[str, Any]:
    return {
        "identificador": item.get("identificador"),
        "idcuenta": item.get("idcuenta"),
        "idrubro": item.get("idrubro"),
        "monto": item.get("montocarga") or item.get("montodebe") or item.get("fv"),
        "fecha": parse_isoish_date(item.get("pago_fecha")),
        "numero": item.get("pago_numero"),
        "banco": item.get("pago_banco"),
        "base_imponible": item.get("pago_baseimp"),
        "alicuota": item.get("pago_alicuota"),
        "impuesto": item.get("pago_impuesto"),
        "regimen": item.get("pago_regimen"),
        "memo": item.get("memo"),
        "rubroposicion": item.get("rubroposicion"),
        "cheque": item.get("cheque"),
    }


def summarize_cobro_historia(comprobante: dict[str, Any]) -> dict[str, Any]:
    raw_items = comprobante.get("item")
    if isinstance(raw_items, list):
        detail_items = [item for item in raw_items if isinstance(item, dict)]
    elif isinstance(raw_items, dict):
        detail_items = [raw_items]
    else:
        detail_items = []

    detalles = [summarize_cobro_historia_item(item) for item in detail_items]
    total = sum((decimal_or_zero(item.get("monto")) for item in detalles), Decimal("0"))
    return {
        "id": str(comprobante.get("id", "")),
        "fecha": parse_isoish_date(comprobante.get("fechaiva") or comprobante.get("fecha")),
        "comprobante": format_comprobante_text(
            comprobante.get("letra"),
            comprobante.get("sucursal"),
            comprobante.get("numero"),
        ),
        "numero": comprobante.get("numero"),
        "cliente": comprobante.get("clipro"),
        "cuit": comprobante.get("cuit"),
        "idclipro": comprobante.get("idclipro"),
        "idcentrocosto": comprobante.get("idcentrocosto"),
        "centro_costo": comprobante.get("centrocosto"),
        "idcuenta": comprobante.get("idcuenta"),
        "cuenta": comprobante.get("cuenta"),
        "idprovinciaiibb": comprobante.get("idprovinciaiibb"),
        "memo": comprobante.get("memo"),
        "moneda": comprobante.get("moneda"),
        "tipocambio": comprobante.get("tipocambio"),
        "total": format_decimal_string(total),
        "detalles": detalles,
    }


def normalize_optional_name(value: str | None) -> str | None:
    if not value:
        return None
    return normalize_name(value)


def search_cobros_range_payload(
    payload: dict[str, Any],
    *,
    recibo: str | None = None,
    comprobante: str | None = None,
    cliente_cuit: str | None = None,
    cliente_nombre: str | None = None,
) -> list[dict[str, Any]]:
    matches: list[dict[str, Any]] = []
    target_recibo = digits_only(recibo or "") if recibo else None
    target_comprobante = normalize_optional_name(comprobante)
    target_cliente_cuit = digits_only(cliente_cuit or "") if cliente_cuit else None
    target_cliente_nombre = normalize_optional_name(cliente_nombre)

    for item in payload.get("items", []):
        if not isinstance(item, dict):
            continue
        if target_recibo and digits_only(str(item.get("numero") or item.get("comprobante") or "")) != target_recibo:
            continue
        if target_comprobante and normalize_name(str(item.get("comprobante") or "")) != target_comprobante:
            continue
        if target_cliente_cuit and digits_only(str(item.get("cuitreceptor") or "")) != target_cliente_cuit:
            continue
        if target_cliente_nombre and normalize_name(str(item.get("clipro") or "")) != target_cliente_nombre:
            continue
        matches.append(item)
    return matches


def filtered_web_comprobantes_result(
    payload: dict[str, Any],
    *,
    desde: str,
    hasta: str,
    recibo: str | None = None,
    comprobante: str | None = None,
    cliente_cuit: str | None = None,
    cliente_nombre: str | None = None,
) -> dict[str, Any]:
    matches = search_cobros_range_payload(
        payload,
        recibo=recibo,
        comprobante=comprobante,
        cliente_cuit=cliente_cuit,
        cliente_nombre=cliente_nombre,
    )
    total = sum((decimal_or_zero(item.get("total")) for item in matches), Decimal("0"))
    result = {
        "transport": "web-session",
        "desde": desde,
        "hasta": hasta,
        "count": len(matches),
        "sumatotal": format_decimal_string(total),
        "items": [summarize_web_comprobante(item) for item in matches],
    }
    if cliente_cuit:
        result["cliente_cuit"] = cliente_cuit
    if cliente_nombre:
        result["cliente_nombre"] = cliente_nombre
    if recibo:
        result["recibo"] = recibo
    if comprobante:
        result["comprobante"] = comprobante
    return result


def resolve_cobro_match_from_lookup(args: argparse.Namespace, client: SOSContadorClient) -> dict[str, Any]:
    if getattr(args, "id", None):
        web_client = SOSContadorWebClient(api_client=client)
        payload = web_client.get_comprobante_historia(idcomprobante=str(args.id), idtipo_operacion=12)
        return {
            "id": str(args.id),
            "idclipro": payload.get("idclipro"),
            "numero": payload.get("numero"),
            "comprobante": format_comprobante_text(
                payload.get("letra"),
                payload.get("sucursal"),
                payload.get("numero"),
            ),
            "clipro": payload.get("clipro"),
            "cuitreceptor": payload.get("cuit"),
        }

    if not any([getattr(args, "recibo", None), getattr(args, "comprobante", None)]):
        raise CLIError("Debe indicar --id o bien --recibo / --comprobante para identificar el recibo.")

    if getattr(args, "fecha", None):
        desde = hasta = parse_human_date(args.fecha)
    else:
        if not (getattr(args, "desde", None) and getattr(args, "hasta", None)):
            raise CLIError("Para resolver un recibo sin ID use --fecha o bien --desde y --hasta.")
        desde = parse_human_date(args.desde)
        hasta = parse_human_date(args.hasta)

    web_client = SOSContadorWebClient(api_client=client)
    payload = web_client.list_comprobantes_range(idtipo_operacion=12, desde=desde, hasta=hasta, length=args.registros)
    matches = search_cobros_range_payload(
        payload,
        recibo=getattr(args, "recibo", None),
        comprobante=getattr(args, "comprobante", None),
        cliente_cuit=getattr(args, "cliente_cuit", None),
        cliente_nombre=getattr(args, "cliente_nombre", None),
    )
    if not matches:
        raise CLIError("No se encontro un recibo que coincida con los criterios indicados.")
    if len(matches) > 1:
        simplified = [summarize_web_comprobante(item) for item in matches[:10]]
        raise CLIError(
            "La busqueda del recibo devolvio multiples coincidencias. "
            f"Acote por fecha o cliente.\n{format_output(simplified)}"
        )
    return matches[0]


def resolve_cobro_id_from_lookup(args: argparse.Namespace, client: SOSContadorClient) -> str:
    match_id = resolve_cobro_match_from_lookup(args, client).get("id")
    if match_id in (None, ""):
        raise CLIError("El recibo encontrado no trae un ID utilizable.")
    return str(match_id)


def coerce_association_group_items(group: Any) -> list[dict[str, Any]]:
    if not isinstance(group, dict):
        return []
    raw_items = group.get("item")
    if isinstance(raw_items, list):
        return [item for item in raw_items if isinstance(item, dict)]
    if isinstance(raw_items, dict):
        return [raw_items]
    return []


def parse_association_groups(payload: dict[str, Any]) -> list[list[dict[str, Any]]]:
    raw_groups = payload.get("asociacion")
    if isinstance(raw_groups, list):
        groups = raw_groups
    elif isinstance(raw_groups, dict):
        groups = [raw_groups]
    else:
        return []
    return [items for items in (coerce_association_group_items(group) for group in groups) if items]


def summarize_association_document(item: dict[str, Any]) -> dict[str, Any]:
    return {
        "id": item.get("id"),
        "idtipo_operacion": item.get("idtipo_operacion"),
        "comprobante": item.get("comprobante")
        or format_comprobante_text(item.get("letra"), item.get("sucursal"), item.get("numero")),
        "fecha": parse_isoish_date(item.get("fechaiva") or item.get("fecha")),
        "total": item.get("total"),
    }


def summarize_association_groups(groups: list[list[dict[str, Any]]]) -> list[dict[str, Any]]:
    summary: list[dict[str, Any]] = []
    for index, group in enumerate(groups, start=1):
        total_cobros = Decimal("0")
        total_ventas = Decimal("0")
        for item in group:
            amount = decimal_or_zero(item.get("total"))
            if str(item.get("idtipo_operacion")) == "12":
                total_cobros += amount
            else:
                total_ventas += amount
        summary.append(
            {
                "grupo": index,
                "items": [summarize_association_document(item) for item in group],
                "total_cobros": format_decimal_string(total_cobros),
                "total_ventas": format_decimal_string(total_ventas),
                "diferencia": format_decimal_string(total_cobros - total_ventas),
            }
        )
    return summary


def item_matches_comprobante(item: dict[str, Any], lookup: str) -> bool:
    target_digits = digits_only(lookup)
    comprobante = str(
        item.get("comprobante")
        or format_comprobante_text(item.get("letra"), item.get("sucursal"), item.get("numero"))
        or ""
    )
    if target_digits and digits_only(comprobante) == target_digits:
        return True
    return normalize_name(comprobante) == normalize_name(lookup)


def find_association_document(
    groups: list[list[dict[str, Any]]],
    *,
    document_id: str | None = None,
    comprobante: str | None = None,
    idtipo_operacion: str | None = None,
    role_label: str,
) -> dict[str, Any]:
    matches: list[dict[str, Any]] = []
    for group in groups:
        for item in group:
            if idtipo_operacion is not None and str(item.get("idtipo_operacion")) != str(idtipo_operacion):
                continue
            if document_id and str(item.get("id")) == str(document_id):
                matches.append(item)
                continue
            if comprobante and item_matches_comprobante(item, comprobante):
                matches.append(item)
    if not matches:
        raise CLIError(f"No se encontro {role_label} en la ventana de asociacion del cliente.")
    unique_matches = {str(item.get('id')): item for item in matches if item.get("id") not in (None, "")}
    if len(unique_matches) > 1:
        raise CLIError(
            f"La busqueda de {role_label} devolvio multiples coincidencias. "
            f"Sea mas especifico.\n{format_output([summarize_association_document(item) for item in unique_matches.values()])}"
        )
    if unique_matches:
        return next(iter(unique_matches.values()))
    return matches[0]


def build_association_save_body(
    *,
    idclipro: str,
    groups: list[list[dict[str, Any]]],
    target_document_ids: list[str],
) -> dict[str, Any]:
    target_ids = {str(item) for item in target_document_ids}
    rebuilt_groups: list[list[str]] = []
    for group in groups:
        remaining = [str(item.get("id")) for item in group if str(item.get("id") or "") and str(item.get("id")) not in target_ids]
        if remaining:
            rebuilt_groups.append(remaining)
    rebuilt_groups.append(list(dict.fromkeys(str(item) for item in target_document_ids if str(item or "").strip())))

    asociados: list[dict[str, list[str]]] = []
    desasociados: list[str] = []
    seen_desasociados: set[str] = set()
    for group in rebuilt_groups:
        deduped_group = list(dict.fromkeys(group))
        if len(deduped_group) >= 2:
            asociados.append({"c": deduped_group})
            continue
        only_id = deduped_group[0]
        if only_id not in seen_desasociados:
            desasociados.append(only_id)
            seen_desasociados.add(only_id)

    return {
        "idclipro": str(idclipro),
        "a": asociados,
        "d": desasociados,
    }


def verify_association_pair(
    groups: list[list[dict[str, Any]]],
    *,
    expected_document_ids: list[str],
) -> dict[str, Any]:
    expected_ids = {str(item) for item in expected_document_ids}
    target_group: list[dict[str, Any]] | None = None
    for group in groups:
        group_ids = {str(item.get("id")) for item in group if item.get("id") not in (None, "")}
        if expected_ids.issubset(group_ids):
            target_group = group
            break
    if target_group is None:
        raise CLIError(
            "La asociacion no quedo grabada: no se encontro un grupo que contenga todos los comprobantes indicados."
        )
    target_ids = {str(item.get("id")) for item in target_group if item.get("id") not in (None, "")}
    if target_ids != expected_ids:
        raise CLIError(
            "La asociacion se grabo, pero el grupo final no coincide exactamente con el conjunto pedido.\n"
            f"{format_output([summarize_association_document(item) for item in target_group])}"
        )
    return {
        "grupo_objetivo": [summarize_association_document(item) for item in target_group],
        "grupos": summarize_association_groups(groups),
    }


def should_fallback_cobro_detail(payload: Any) -> bool:
    return isinstance(payload, dict) and bool(payload.get("error"))


def parse_csv_rows(text: str) -> list[dict[str, str]]:
    cleaned = text.lstrip("\ufeff")
    reader = csv.DictReader(io.StringIO(cleaned), delimiter=";")
    return [dict(row) for row in reader]


def command_auth_info(args: argparse.Namespace, client: SOSContadorClient) -> Any:
    session = client.get_login_session()
    available_entries = collect_cuit_entries(session.login_payload)
    available_cuits = [entry["cuit"] for entry in available_entries if entry["cuit"]]
    result = {
        "base_url": client.base_url,
        "cuit_id": None,
        "cuit": None,
        "available_cuits": available_cuits,
        "available_cuits_count": len(available_cuits),
        "runtime_home": str(RUNTIME_HOME),
        "local_env_loaded": LOCAL_ENV_LOADED,
        "local_env_path": str(LOADED_LOCAL_ENV_PATH or LOCAL_ENV_PATH),
        "jwt_obtenido": bool(session.jwt),
        "jwtc_obtenido": False,
        "deprecated_work_cuit_env_present": bool(deprecated_work_cuit_env_values()),
        "deprecated_work_cuit_env_ignored": bool(deprecated_work_cuit_env_values()),
        "deprecated_work_cuit_env_vars": sorted(deprecated_work_cuit_env_values().keys()),
    }
    if args.show_tokens:
        result["jwt"] = session.jwt
    return result


def command_auth_list_cuits(args: argparse.Namespace, client: SOSContadorClient) -> Any:
    items = refresh_cuit_catalog(client)
    return {
        "items": [format_cuit_catalog_item(item) for item in items],
        "cache_written_to": str(CUIT_CATALOG_CACHE_PATH),
    }


def command_auth_resolve_cuit(args: argparse.Namespace, client: SOSContadorClient) -> Any:
    match = resolve_cuit_name(client, name=args.name, refresh=args.refresh)
    match["aliases_loaded"] = bool(load_cuit_aliases())
    match["catalog_cache_path"] = str(CUIT_CATALOG_CACHE_PATH)
    match["aliases_path"] = str(CUIT_ALIASES_PATH)
    return match


def command_auth_web_info(args: argparse.Namespace, client: SOSContadorClient) -> Any:
    bound_client, target = bind_business_client(client, args, required=True)
    web_client = SOSContadorWebClient(
        api_client=bound_client,
        explicit_cuit=target.get("cuit") if target else None,
        explicit_cuit_id=target.get("cuit_id") if target else None,
    )
    session = web_client.get_session()
    return {
        "base_url": web_client.base_url,
        "transport": "web-session",
        "cuit_id": session.cuit_id,
        "cuit": session.cuit,
        "nombrecuit": session.nombrecuit,
        "runtime_home": str(RUNTIME_HOME),
        "local_env_loaded": LOCAL_ENV_LOADED,
        "local_env_path": str(LOADED_LOCAL_ENV_PATH or LOCAL_ENV_PATH),
        "login_message": session.login_message,
        "cookies_obtenidas": sum(1 for _ in web_client.cookie_jar) > 0,
    }


def command_api_catalog(args: argparse.Namespace, client: SOSContadorClient) -> Any:
    del client
    operations = public_api_operations()
    if getattr(args, "module", None):
        module = str(args.module).strip().lower()
        operations = [item for item in operations if str(item.get("module") or "").lower() == module]
    if getattr(args, "method", None):
        method = str(args.method).upper()
        operations = [item for item in operations if str(item.get("method") or "").upper() == method]
    if getattr(args, "status", None):
        status = str(args.status).lower()
        operations = [item for item in operations if str(item.get("status") or "").lower() == status]
    return {
        "source": load_public_api_catalog().get("source"),
        "reviewed_at": load_public_api_catalog().get("reviewed_at"),
        "count": len(operations),
        "items": [
            {
                "id": item["id"],
                "module": item.get("module"),
                "method": item["method"],
                "path": item["path"],
                "auth": item["auth"],
                "status": item.get("status", "documented"),
                "side_effect": bool(item.get("side_effect")),
                "invokable": item.get("invokable", True),
                "summary": item.get("summary", ""),
            }
            for item in operations
        ],
    }


def command_api_describe(args: argparse.Namespace, client: SOSContadorClient) -> Any:
    del client
    return public_api_operation(args.operation)


def command_api_invoke(args: argparse.Namespace, client: SOSContadorClient) -> Any:
    operation = public_api_operation(args.operation)
    if operation.get("invokable", True) is False:
        guidance = operation.get("guidance") or "Use el helper especializado indicado en el catalogo."
        raise CLIError(f"{operation['id']} no se invoca directamente. {guidance}")

    path_parameters = parse_unique_key_values(args.param or [], label="Parametro de path")
    path = render_operation_path(str(operation["path"]), path_parameters)
    query = [parse_key_value(item) for item in args.query or []]
    validate_operation_query(operation, query)
    body = load_json_source(args.body_json, args.body_file)
    validate_operation_body(operation, body)

    auth_mode = str(operation["auth"])
    bound_client = client
    work_target = None
    if auth_mode == "jwtc":
        bound_client, work_target = bind_business_client(client, args, required=True)

    if operation.get("side_effect", False):
        detail_rows = [
            {
                "Operacion": operation["id"],
                "Metodo": operation["method"],
                "Path": path,
                "Query": "&".join(f"{key}={value}" for key, value in query),
                "Body": format_output(redact_sensitive_payload(body)) if body is not None else "",
            }
        ]
        if auth_mode == "jwtc":
            preview, preview_markdown = build_generic_mutation_preview(
                str(operation["id"]).replace(".", "_"),
                bound_client=bound_client,
                work_target=work_target,
                detail_rows=detail_rows,
                detail_title="Operacion documentada",
            )
        else:
            preview = {
                "titulo": str(operation["id"]).replace(".", "_"),
                "contexto": [{"Ambito": "cuenta de usuario"}],
                "tablas": [{"titulo": "Operacion documentada", "rows": detail_rows}],
                "totales": {},
                "advertencias": [],
            }
            preview_markdown = render_business_preview_markdown(preview)
        set_mutation_preview(args, business_preview=preview, preview_markdown=preview_markdown)
        setattr(args, "auth_mode", auth_mode)
        ensure_mutation_allowed(str(operation["method"]), path, query, body, args)

    return bound_client.request(
        str(operation["method"]),
        path,
        query=query,
        body=body,
        auth_mode=auth_mode,
        out_path=args.out,
    )


def command_call(args: argparse.Namespace, client: SOSContadorClient) -> Any:
    body = load_json_source(args.body_json, args.body_file)
    query = [parse_key_value(item) for item in args.query or []]
    bound_client = client
    work_target = None
    if getattr(args, "auth_mode", "jwtc") == "jwtc":
        bound_client, work_target = bind_business_client(client, args, required=True)
    preview, preview_markdown = build_generic_mutation_preview(
        "call",
        bound_client=bound_client,
        work_target=work_target,
        detail_rows=[
            {
                "Método": str(args.method).upper(),
                "Path": args.path,
                "Query": "&".join(f"{key}={value}" for key, value in query),
                "Body": format_output(body) if body is not None else "",
            }
        ],
        detail_title="Operacion",
    )
    set_mutation_preview(args, business_preview=preview, preview_markdown=preview_markdown)
    ensure_mutation_allowed(args.method, args.path, query, body, args)
    return bound_client.request(args.method, args.path, query=query, body=body, auth_mode=args.auth_mode, out_path=args.out)


def command_cliente_list(args: argparse.Namespace, client: SOSContadorClient) -> Any:
    bound_client, _ = bind_business_client(client, args, required=True)
    query = [
        ("cliente", bool_to_str(args.cliente)),
        ("proveedor", bool_to_str(args.proveedor)),
        ("pagina", str(args.pagina)),
        ("registros", str(args.registros)),
    ]
    return bound_client.request("GET", "cliente/listado", query=query)


def command_cliente_get(args: argparse.Namespace, client: SOSContadorClient) -> Any:
    if not (args.id or args.cuit or args.nombre):
        raise CLIError("Debe indicar una via de busqueda: --id, --cuit o --nombre.")
    target_cuit = require_valid_argentina_cuit(args.cuit) if args.cuit else None
    bound_client, _ = bind_business_client(client, args, required=True)

    def predicate(item: dict[str, Any]) -> bool:
        if args.id:
            candidate = find_value(item, ("idclipro", "id", "idcliente"))
            if candidate is not None and str(candidate) == str(args.id):
                return True
        if target_cuit:
            candidate = find_value(item, ("cuit",))
            if candidate is not None and digits_only(str(candidate)) == target_cuit:
                return True
        if args.nombre:
            candidate = find_value(item, ("clipro", "cliente", "nombre"))
            if candidate is not None and normalize_name(str(candidate)) == normalize_name(args.nombre):
                return True
        return False

    match = search_paginated(
        bound_client,
        path="cliente/listado",
        query=[("cliente", "true"), ("proveedor", "true"), ("pagina", "1"), ("registros", "200")],
        predicate=predicate,
    )
    if not match:
        raise CLIError("No se encontro un cliente que coincida con el criterio indicado.")
    return match


def command_cliente_create(args: argparse.Namespace, client: SOSContadorClient) -> Any:
    body = load_json_source(args.body_json, args.body_file)
    if body is None:
        raise CLIError("cliente create requiere --body-json o --body-file.")
    if isinstance(body, dict) and body.get("cuit") not in (None, ""):
        require_valid_argentina_cuit(body["cuit"], label="CUIT del cliente/proveedor")
    bound_client, work_target = bind_business_client(client, args, required=True)
    preview, preview_markdown = build_generic_payload_preview(
        "cliente_create",
        bound_client=bound_client,
        work_target=work_target,
        body=body,
    )
    set_mutation_preview(args, business_preview=preview, preview_markdown=preview_markdown)
    ensure_mutation_allowed("POST", "cliente", None, body, args)
    return bound_client.request("POST", "cliente", body=body)


def command_cliente_update(args: argparse.Namespace, client: SOSContadorClient) -> Any:
    body = load_json_source(args.body_json, args.body_file)
    if body is None:
        raise CLIError("cliente update requiere --body-json o --body-file.")
    if isinstance(body, dict) and body.get("cuit") not in (None, ""):
        require_valid_argentina_cuit(body["cuit"], label="CUIT del cliente/proveedor")
    bound_client, work_target = bind_business_client(client, args, required=True)
    preview, preview_markdown = build_generic_payload_preview(
        "cliente_update",
        bound_client=bound_client,
        work_target=work_target,
        body=body,
    )
    set_mutation_preview(args, business_preview=preview, preview_markdown=preview_markdown)
    ensure_mutation_allowed("PUT", f"cliente/{args.id}", None, body, args)
    return bound_client.request("PUT", f"cliente/{args.id}", body=body)


def command_cliente_delete(args: argparse.Namespace, client: SOSContadorClient) -> Any:
    bound_client, work_target = bind_business_client(client, args, required=True)
    preview, preview_markdown = build_generic_mutation_preview(
        "cliente_delete",
        bound_client=bound_client,
        work_target=work_target,
        detail_rows=[{"ID cliente": args.id}],
        detail_title="Operacion",
    )
    set_mutation_preview(args, business_preview=preview, preview_markdown=preview_markdown)
    ensure_mutation_allowed("DELETE", f"cliente/{args.id}", None, None, args)
    return bound_client.request("DELETE", f"cliente/{args.id}")


def command_producto_list(args: argparse.Namespace, client: SOSContadorClient) -> Any:
    bound_client, _ = bind_business_client(client, args, required=True)
    query = [("pagina", str(args.pagina)), ("registros", str(args.registros))]
    return bound_client.request("GET", "producto/listado", query=query)


def command_producto_create(args: argparse.Namespace, client: SOSContadorClient) -> Any:
    bound_client, work_target = bind_business_client(client, args, required=True)
    body = load_json_source(args.body_json, args.body_file)
    if body is None:
        raise CLIError("producto create requiere --body-json o --body-file.")
    preview, preview_markdown = build_generic_payload_preview(
        "producto_create",
        bound_client=bound_client,
        work_target=work_target,
        body=body,
    )
    set_mutation_preview(args, business_preview=preview, preview_markdown=preview_markdown)
    ensure_mutation_allowed("POST", "producto", None, body, args)
    return bound_client.request("POST", "producto", body=body)


def command_producto_update(args: argparse.Namespace, client: SOSContadorClient) -> Any:
    bound_client, work_target = bind_business_client(client, args, required=True)
    body = load_json_source(args.body_json, args.body_file)
    if body is None:
        raise CLIError("producto update requiere --body-json o --body-file.")
    preview, preview_markdown = build_generic_payload_preview(
        "producto_update",
        bound_client=bound_client,
        work_target=work_target,
        body=body,
    )
    set_mutation_preview(args, business_preview=preview, preview_markdown=preview_markdown)
    ensure_mutation_allowed("PUT", f"producto/{args.id}", None, body, args)
    return bound_client.request("PUT", f"producto/{args.id}", body=body)


def command_periodic_list(args: argparse.Namespace, client: SOSContadorClient, entity: str) -> Any:
    bound_client, _ = bind_business_client(client, args, required=True)
    query = [("pagina", str(args.pagina)), ("registros", str(args.registros))]
    return bound_client.request("GET", f"{entity}/listado/{args.periodo}", query=query)


def command_detail(args: argparse.Namespace, client: SOSContadorClient, entity: str) -> Any:
    bound_client, _ = bind_business_client(client, args, required=True)
    return bound_client.request("GET", f"{entity}/detalle/{args.id}")


def extract_draft_association_targets(draft: dict[str, Any]) -> tuple[list[str], list[str]]:
    venta_ids: list[str] = []
    facturas: list[str] = []
    for item in draft.get("asociacion", {}).get("facturas", []):
        if not isinstance(item, dict):
            continue
        venta_id = str(item.get("venta_id") or item.get("id") or "").strip()
        comprobante = normalize_comprobante_reference(item.get("comprobante"))
        if venta_id:
            venta_ids.append(venta_id)
        elif comprobante:
            facturas.append(comprobante)
    return venta_ids, facturas


def associate_cobro_documents(
    *,
    client: SOSContadorClient,
    cobro_id: str,
    idclipro: str,
    venta_ids: list[str],
    facturas: list[str],
    args: argparse.Namespace,
) -> dict[str, Any]:
    if not cobro_id:
        raise CLIError("No se pudo determinar el ID del cobro a asociar.")
    if not idclipro:
        raise CLIError("No se pudo determinar el idclipro del cobro a asociar.")
    if not (venta_ids or facturas):
        return {}

    web_client = SOSContadorWebClient(api_client=client, explicit_cuit=client.get_session().cuit)
    before_payload = web_client.get_cobropago_vs_compraventa(idclipro=idclipro, idtipo_operacion=12, items=[cobro_id])
    before_groups = parse_association_groups(before_payload)
    cobro_doc = find_association_document(
        before_groups,
        document_id=cobro_id,
        idtipo_operacion="12",
        role_label="el cobro indicado",
    )
    venta_docs: list[dict[str, Any]] = []
    for venta_id in venta_ids:
        venta_docs.append(
            find_association_document(
                before_groups,
                document_id=venta_id,
                role_label=f"la venta {venta_id}",
            )
        )
    for factura in facturas:
        venta_docs.append(
            find_association_document(
                before_groups,
                comprobante=factura,
                role_label=f"la factura {factura}",
            )
        )
    venta_docs = list({str(item.get("id")): item for item in venta_docs}.values())
    target_document_ids = [str(cobro_doc.get("id"))] + [str(item.get("id")) for item in venta_docs]
    save_body = build_association_save_body(
        idclipro=idclipro,
        groups=before_groups,
        target_document_ids=target_document_ids,
    )
    args.auth_mode = "web-session"
    ensure_mutation_allowed("POST", "back/asociaciones_altamodi.asp", None, save_body, args)
    save_payload = web_client.save_asociaciones(save_body)

    after_payload = web_client.get_cobropago_vs_compraventa(idclipro=idclipro, idtipo_operacion=12, items=[cobro_id])
    after_groups = parse_association_groups(after_payload)
    verification = verify_association_pair(after_groups, expected_document_ids=target_document_ids)
    return {
        "transport": "web-session",
        "idclipro": idclipro,
        "cobro": summarize_association_document(cobro_doc),
        "ventas": [summarize_association_document(item) for item in venta_docs],
        "save_payload": save_payload,
        "before": summarize_association_groups(before_groups),
        "after": verification["grupos"],
        "grupo_objetivo": verification["grupo_objetivo"],
    }


def command_cobro_get(args: argparse.Namespace, client: SOSContadorClient) -> Any:
    bound_client, _ = bind_business_client(client, args, required=True)
    cobro_id = resolve_cobro_id_from_lookup(args, bound_client)
    api_error: str | None = None
    try:
        api_payload = bound_client.request("GET", f"cobro/detalle/{cobro_id}")
    except APIError as exc:
        api_payload = exc.payload
        api_error = f"HTTP {exc.status} {exc.reason}"

    if api_error is None and not should_fallback_cobro_detail(api_payload):
        return {
            "transport": "api",
            "id": cobro_id,
            "payload": api_payload,
        }

    web_client = SOSContadorWebClient(api_client=bound_client)
    web_payload = web_client.get_comprobante_historia(idcomprobante=cobro_id, idtipo_operacion=12)
    return {
        "transport": "web-session",
        "fallback_reason": "La API publica devolvio un error en cobro/detalle.",
        "api_error": api_error,
        "api_payload": api_payload,
        "payload": summarize_cobro_historia(web_payload),
    }


def command_cobro_draft(args: argparse.Namespace, client: SOSContadorClient) -> Any:
    draft = build_cobro_draft(args, client)
    preview = build_cobro_draft_preview(draft)
    preview_markdown = render_cobro_draft_markdown(draft)
    if getattr(args, "json_out", None):
        destination = Path(args.json_out)
        save_json_file(destination, draft)
        draft["json_written_to"] = str(destination)
    preview_format = str(getattr(args, "preview_format", "json") or "json").strip().lower()
    if preview_format == "markdown":
        return preview_markdown
    result = {
        "draft_id": draft.get("draft_id"),
        "draft_file": draft.get("draft_file"),
        "payload": draft,
        "preview": preview,
    }
    if preview_format == "both":
        result["preview_markdown"] = preview_markdown
    return result


def command_cobro_create(args: argparse.Namespace, client: SOSContadorClient) -> Any:
    if args_use_draft_source(args):
        draft = load_draft_payload(
            draft_id=getattr(args, "draft_id", None),
            draft_file=getattr(args, "draft_file", None),
        )
        bound_client, body = build_cobro_detail_body_from_draft(draft, client)
        args.auth_mode = "web-session"
        set_mutation_preview(
            args,
            business_preview=build_cobro_draft_preview(draft),
            preview_markdown=render_cobro_draft_markdown(draft),
        )
        ensure_mutation_allowed("POST", "back/comprobante_altamodi.asp", None, body, args)
        started_at = time.perf_counter()
        web_client = SOSContadorWebClient(api_client=bound_client, explicit_cuit=bound_client.get_session().cuit)
        payload = web_client.save_comprobante(body)
        cobro_identity = resolve_saved_cobro_identity(save_payload=payload, body=body, client=bound_client)
        venta_ids, facturas = extract_draft_association_targets(draft)
        association = associate_cobro_documents(
            client=bound_client,
            cobro_id=cobro_identity.get("id") or "",
            idclipro=cobro_identity.get("idclipro") or str(body.get("idclipro") or ""),
            venta_ids=venta_ids,
            facturas=facturas,
            args=args,
        )
        draft.setdefault("validacion", {}).setdefault("timings", {})["execution_seconds"] = round(
            time.perf_counter() - started_at,
            3,
        )
        save_draft_payload(draft)
        return {
            "transport": "web-session",
            "draft_id": draft.get("draft_id"),
            "numero": body.get("numero"),
            "fecha": body.get("fecha"),
            "idclipro": body.get("idclipro"),
            "id": cobro_identity.get("id"),
            "comprobante": cobro_identity.get("comprobante"),
            "payload": payload,
            "association": association or None,
            "timings": draft["validacion"]["timings"],
        }

    bound_client, work_target = bind_business_client(client, args, required=True)
    transport, path, body = build_cobro_create_request(args, bound_client)
    preview, preview_markdown = build_cobro_or_pago_business_preview(
        "cobro_create",
        bound_client=bound_client,
        work_target=work_target,
        body=body,
    )
    set_mutation_preview(args, business_preview=preview, preview_markdown=preview_markdown)
    if transport == "web-session":
        args.auth_mode = "web-session"
        ensure_mutation_allowed("POST", path, None, body, args)
        web_client = SOSContadorWebClient(api_client=bound_client)
        payload = web_client.save_comprobante(body)
        return {
            "transport": "web-session",
            "numero": body.get("numero"),
            "fecha": body.get("fecha"),
            "idclipro": body.get("idclipro"),
            "payload": payload,
        }
    ensure_mutation_allowed("PUT", path, None, body, args)
    return bound_client.request("PUT", path, body=body)


def command_cobro_asociar(args: argparse.Namespace, client: SOSContadorClient) -> Any:
    bound_client, _ = bind_business_client(client, args, required=True)
    venta_ids = list(getattr(args, "venta_id", None) or [])
    facturas = list(getattr(args, "factura", None) or [])
    if not (venta_ids or facturas):
        raise CLIError("cobro asociar requiere --venta-id o --factura para identificar al menos una venta.")

    cobro_match = resolve_cobro_match_from_lookup(args, bound_client)
    cobro_id = str(cobro_match.get("id") or "")
    if not cobro_id:
        raise CLIError("El recibo seleccionado no tiene un ID utilizable.")
    idclipro = str(cobro_match.get("idclipro") or "")
    if not idclipro:
        raise CLIError("No se pudo determinar el idclipro del recibo seleccionado.")
    return associate_cobro_documents(
        client=bound_client,
        cobro_id=cobro_id,
        idclipro=idclipro,
        venta_ids=venta_ids,
        facturas=facturas,
        args=args,
    )


def command_cobro_list_range(args: argparse.Namespace, client: SOSContadorClient) -> Any:
    bound_client, _ = bind_business_client(client, args, required=True)
    desde = parse_human_date(args.desde)
    hasta = parse_human_date(args.hasta)
    web_client = SOSContadorWebClient(api_client=bound_client)
    payload = web_client.list_comprobantes_range(idtipo_operacion=12, desde=desde, hasta=hasta, length=args.registros)
    result = filtered_web_comprobantes_result(
        payload,
        desde=desde,
        hasta=hasta,
        recibo=args.recibo,
        comprobante=args.comprobante,
        cliente_cuit=args.cliente_cuit,
        cliente_nombre=args.cliente_nombre,
    )
    if args.csv_out:
        csv_text = web_client.export_comprobantes_csv(idtipo_operacion=12, desde=desde, hasta=hasta)
        destination = Path(args.csv_out)
        destination.parent.mkdir(parents=True, exist_ok=True)
        destination.write_text(csv_text, encoding="utf-8")
        result["csv_written_to"] = str(destination)
        result["csv_rows"] = len(parse_csv_rows(csv_text))
    return result


def command_cobro_resolve_id(args: argparse.Namespace, client: SOSContadorClient) -> Any:
    bound_client, _ = bind_business_client(client, args, required=True)
    if not any([args.recibo, args.comprobante, args.cliente_cuit, args.cliente_nombre]):
        raise CLIError("cobro resolve-id requiere al menos --recibo, --comprobante, --cliente-cuit o --cliente-nombre.")
    if args.fecha:
        desde = hasta = parse_human_date(args.fecha)
    else:
        if not (args.desde and args.hasta):
            raise CLIError("Use --fecha o bien --desde y --hasta para acotar la busqueda.")
        desde = parse_human_date(args.desde)
        hasta = parse_human_date(args.hasta)

    web_client = SOSContadorWebClient(api_client=bound_client)
    payload = web_client.list_comprobantes_range(idtipo_operacion=12, desde=desde, hasta=hasta, length=args.registros)
    matches = [
        summarize_web_comprobante(item)
        for item in search_cobros_range_payload(
            payload,
            recibo=args.recibo,
            comprobante=args.comprobante,
            cliente_cuit=args.cliente_cuit,
            cliente_nombre=args.cliente_nombre,
        )
    ]

    return {
        "transport": "web-session",
        "desde": desde,
        "hasta": hasta,
        "matches": matches,
    }


def command_pago_create(args: argparse.Namespace, client: SOSContadorClient) -> Any:
    bound_client, work_target = bind_business_client(client, args, required=True)
    body = build_cobro_or_pago_body(args, bound_client)
    preview, preview_markdown = build_cobro_or_pago_business_preview(
        "pago_create",
        bound_client=bound_client,
        work_target=work_target,
        body=body,
    )
    set_mutation_preview(args, business_preview=preview, preview_markdown=preview_markdown)
    ensure_mutation_allowed("PUT", "pago", None, body, args)
    return bound_client.request("PUT", "pago", body=body)


def command_venta_list(args: argparse.Namespace, client: SOSContadorClient) -> Any:
    bound_client, _ = bind_business_client(client, args, required=True)
    query = [("pagina", str(args.pagina)), ("registros", str(args.registros))]
    if args.fecha_desde:
        query.append(("fecha_desde", args.fecha_desde))
    if args.fecha_hasta:
        query.append(("fecha_hasta", args.fecha_hasta))
    path = f"venta/listado/{args.modo}/{args.periodo}/{args.cae}"
    return bound_client.request("GET", path, query=query)


def command_venta_list_all(args: argparse.Namespace, client: SOSContadorClient) -> Any:
    args.modo = "todas"
    return command_venta_list(args, client)


def command_venta_create(args: argparse.Namespace, client: SOSContadorClient) -> Any:
    bound_client, work_target = bind_business_client(client, args, required=True)
    body = build_venta_create_body(args, bound_client)
    preview, preview_markdown = build_venta_business_preview(
        bound_client=bound_client,
        work_target=work_target,
        body=body,
    )
    set_mutation_preview(args, business_preview=preview, preview_markdown=preview_markdown)
    args.auth_mode = "web-session"
    ensure_mutation_allowed("POST", "back/comprobante_altamodi.asp", None, body, args)
    web_client = SOSContadorWebClient(api_client=bound_client)
    payload = web_client.save_comprobante(body)
    return {
        "transport": "web-session",
        "id": payload.get("id"),
        "idclipro": payload.get("idclipro") or body.get("idclipro"),
        "idcuenta": payload.get("idcuenta") or body.get("idcuenta"),
        "comprobante": format_comprobante_text(body.get("letra"), body.get("sucursal"), body.get("numero")),
        "fecha": body.get("fecha"),
        "payload": payload,
    }


def command_venta_get(args: argparse.Namespace, client: SOSContadorClient) -> Any:
    bound_client, _ = bind_business_client(client, args, required=True)
    return bound_client.request("GET", f"venta/detalle/{args.id}")


def summarize_pdf_transport_result(payload: Any) -> Any:
    if isinstance(payload, str):
        text = " ".join(payload.split())
        if len(text) > 200:
            return text[:197] + "..."
        return text
    return payload


def command_venta_pdf(args: argparse.Namespace, client: SOSContadorClient) -> Any:
    bound_client, _ = bind_business_client(client, args, required=True)
    api_result: Any = None
    api_content_type = ""
    try:
        raw, headers = bound_client.request_raw("GET", f"venta/pdf/{args.id}")
        api_content_type = headers.get("Content-Type", "")
        if response_looks_like_pdf(raw, headers):
            write_binary_output(args.out, raw)
            return {
                "written_to": args.out,
                "content_type": api_content_type,
                "transport": "api",
        }
        api_result = summarize_pdf_transport_result(parse_response_body(raw, headers))
    except APIError as exc:
        api_result = summarize_pdf_transport_result(exc.payload)

    api_detail = "La API no devolvio un PDF valido para la venta solicitada."
    if api_result is not None:
        api_detail = f"{api_detail} Resultado API: {api_result}"
    elif api_content_type:
        api_detail = f"{api_detail} Content-Type API: {api_content_type}"

    if not getattr(args, "allow_web_session_fallback", False):
        raise CLIError(
            f"{api_detail} El fallback web-session requiere permiso explicito; "
            "vuelva a ejecutar con --allow-web-session-fallback solo si ya se decidio usarlo."
        )

    web_client = SOSContadorWebClient(api_client=bound_client)
    try:
        result = web_client.download_venta_pdf(
            venta_id=str(args.id),
            out_path=args.out,
        )
        result["transport"] = "web-session"
        result["fallback_reason"] = api_detail
        return result
    except CLIError as exc:
        raise CLIError(
            "No se pudo descargar el PDF de la venta por API ni por web-session. "
            f"{api_detail} Resultado web-session: {exc}"
        ) from exc


def command_venta_search(args: argparse.Namespace, client: SOSContadorClient) -> Any:
    bound_client, _ = bind_business_client(client, args, required=True)
    body = load_json_source(args.body_json, args.body_file)
    if body is None:
        raise CLIError("venta search requiere --body-json o --body-file.")
    query = [("pagina", str(args.pagina)), ("registros", str(args.registros))]
    return bound_client.request("POST", "venta/consulta", query=query, body=body)


def quantize_money(value: Decimal | Any) -> Decimal:
    if not isinstance(value, Decimal):
        value = decimal_or_zero(value)
    return value.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def make_json_safe(value: Any) -> Any:
    if isinstance(value, Decimal):
        return format_decimal_string(value)
    if isinstance(value, dict):
        return {str(key): make_json_safe(item) for key, item in value.items()}
    if isinstance(value, list):
        return [make_json_safe(item) for item in value]
    return value


def afip_type_rule(label: str) -> dict[str, Any] | None:
    normalized = normalize_search_text(label)
    mapping = {
        normalize_search_text("1 - Factura A"): {"fcncnd": "F", "letra": "A", "tipocomprobante": 1, "sign": Decimal("1")},
        normalize_search_text("81 - Tique Factura A"): {"fcncnd": "F", "letra": "A", "tipocomprobante": 1, "sign": Decimal("1")},
        normalize_search_text("11 - Factura C"): {"fcncnd": "F", "letra": "C", "tipocomprobante": 11, "sign": Decimal("1")},
        normalize_search_text("3 - Nota de Credito A"): {"fcncnd": "C", "letra": "A", "tipocomprobante": 3, "sign": Decimal("-1")},
    }
    return mapping.get(normalized)


def read_afip_mis_comprobantes_workbook(path: Path) -> dict[str, Any]:
    if load_workbook is None:
        raise CLIError("Falta openpyxl para leer archivos XLSX.")
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    row1 = " ".join(str(cell).strip() for cell in next(sheet.iter_rows(min_row=1, max_row=1, values_only=True)) if cell not in (None, ""))
    if not row1:
        raise CLIError("El archivo AFIP no tiene una fila 1 utilizable.")
    row1_normalized = normalize_search_text(row1)
    if "comprobantes recibidos" in row1_normalized:
        operation_kind = "compra"
    elif "comprobantes emitidos" in row1_normalized:
        operation_kind = "venta"
    else:
        raise CLIError("No se pudo identificar si el archivo AFIP es de comprobantes recibidos o emitidos.")
    work_cuit = digits_only(row1)
    if len(work_cuit) != 11:
        raise CLIError("No se pudo extraer el CUIT de trabajo desde la fila 1 del archivo AFIP.")

    header_row = next(sheet.iter_rows(min_row=2, max_row=2, values_only=True))
    header_map = {normalize_search_text(str(value or "")): index for index, value in enumerate(header_row)}

    def get_cell(row_values: tuple[Any, ...], *aliases: str) -> Any:
        for alias in aliases:
            key = normalize_search_text(alias)
            if key in header_map:
                index = header_map[key]
                if index < len(row_values):
                    return row_values[index]
        raise CLIError(f"No se encontro la columna esperada en el archivo AFIP: {aliases[0]}")

    rows: list[dict[str, Any]] = []
    for source_row, row_values in enumerate(sheet.iter_rows(min_row=3, values_only=True), start=3):
        if not any(value not in (None, "") for value in row_values):
            continue
        tipo_raw = str(get_cell(row_values, "Tipo") or "").strip()
        type_rule = afip_type_rule(tipo_raw)
        numero_desde = get_cell(row_values, "Número Desde", "Numero Desde")
        numero_hasta = get_cell(row_values, "Número Hasta", "Numero Hasta")
        fecha_iso = guess_date_from_text(str(get_cell(row_values, "Fecha") or "").strip())
        imp_total_raw = decimal_or_zero(get_cell(row_values, "Imp. Total"))
        validation_errors: list[str] = []
        if numero_desde in (None, "") or numero_hasta in (None, ""):
            validation_errors.append("Número Desde y Número Hasta son obligatorios.")
        elif str(numero_desde) != str(numero_hasta):
            validation_errors.append("Número Desde y Número Hasta difieren.")
        if not fecha_iso:
            validation_errors.append("Fecha invalida.")
        if type_rule is None:
            validation_errors.append(f"Tipo AFIP no soportado: {tipo_raw}.")
        if imp_total_raw == Decimal("0") and get_cell(row_values, "Imp. Total") in (None, ""):
            validation_errors.append("Imp. Total vacio.")
        rows.append(
            {
                "source_row": source_row,
                "fecha": fecha_iso,
                "tipo_raw": tipo_raw,
                "type_rule": type_rule,
                "puntoventa": int(decimal_or_zero(get_cell(row_values, "Punto de Venta"))) if get_cell(row_values, "Punto de Venta") not in (None, "") else 0,
                "numero": int(decimal_or_zero(numero_desde)) if numero_desde not in (None, "") else 0,
                "numerohasta": int(decimal_or_zero(numero_hasta)) if numero_hasta not in (None, "") else 0,
                "cae": digits_only(str(get_cell(row_values, "Cód. Autorización", "Cod. Autorizacion") or "")),
                "counterparty_cuit": digits_only(str(get_cell(row_values, "Nro. Doc. Emisor") or "")),
                "counterparty_name": str(get_cell(row_values, "Denominación Emisor", "Denominacion Emisor") or "").strip(),
                "amounts_raw": {
                    "neto_0": decimal_or_zero(get_cell(row_values, "Neto Grav. IVA 0%")),
                    "neto_2_5": decimal_or_zero(get_cell(row_values, "Neto Grav. IVA 2,5%")),
                    "iva_2_5": decimal_or_zero(get_cell(row_values, "IVA 2,5%")),
                    "neto_5": decimal_or_zero(get_cell(row_values, "Neto Grav. IVA 5%")),
                    "iva_5": decimal_or_zero(get_cell(row_values, "IVA 5%")),
                    "neto_10_5": decimal_or_zero(get_cell(row_values, "Neto Grav. IVA 10,5%")),
                    "iva_10_5": decimal_or_zero(get_cell(row_values, "IVA 10,5%")),
                    "neto_21": decimal_or_zero(get_cell(row_values, "Neto Grav. IVA 21%")),
                    "iva_21": decimal_or_zero(get_cell(row_values, "IVA 21%")),
                    "neto_27": decimal_or_zero(get_cell(row_values, "Neto Grav. IVA 27%")),
                    "iva_27": decimal_or_zero(get_cell(row_values, "IVA 27%")),
                    "nogravado": decimal_or_zero(get_cell(row_values, "Neto No Gravado")),
                    "exento": decimal_or_zero(get_cell(row_values, "Op. Exentas")),
                    "otros": decimal_or_zero(get_cell(row_values, "Otros Tributos")),
                    "imp_total": imp_total_raw,
                },
                "validation_errors": validation_errors,
            }
        )
    return {
        "row1": row1,
        "operation_kind": operation_kind,
        "work_cuit": work_cuit,
        "rows": rows,
    }


def resolve_afip_work_target(
    client: SOSContadorClient,
    args: argparse.Namespace,
    *,
    workbook_work_cuit: str,
) -> tuple[SOSContadorClient, dict[str, Any]]:
    explicit = any(
        getattr(args, key, None)
        for key in ("cuit_trabajo", "cuit_trabajo_id", "cuit_trabajo_nombre")
    )
    if explicit:
        bound_client, target = bind_business_client(client, args, required=True)
        if target is None:
            raise CLIError("No se pudo resolver la CUIT de trabajo indicada.")
        target_cuit = digits_only(str(target.get("cuit") or ""))
        if target_cuit and target_cuit != workbook_work_cuit:
            raise CLIError(
                "La CUIT de trabajo indicada no coincide con la fila 1 del archivo AFIP. "
                f"Archivo: {workbook_work_cuit}. Indicada: {target_cuit}."
            )
        return bound_client, target
    target = resolve_work_cuit_target(client, explicit_cuit=workbook_work_cuit, required=True)
    bound_client = ensure_bound_client(client, cuit=workbook_work_cuit, cuit_id=str(target.get("cuit_id") or ""))
    return bound_client, target


def fetch_all_clientes_catalog(bound_client: SOSContadorClient) -> dict[str, dict[str, Any]]:
    results: dict[str, dict[str, Any]] = {}
    page = 1
    while True:
        payload = bound_client.request(
            "GET",
            "cliente/listado",
            query=[("cliente", "true"), ("proveedor", "true"), ("pagina", str(page)), ("registros", "200")],
        )
        items = payload.get("items", [])
        for item in items if isinstance(items, list) else []:
            if not isinstance(item, dict):
                continue
            cuit_value = digits_only(str(item.get("cuit") or ""))
            if cuit_value:
                results[cuit_value] = item
        paginas = int(payload.get("paginas") or page)
        if page >= paginas:
            break
        page += 1
    return results


def fetch_compra_consulta_items(bound_client: SOSContadorClient, *, desde: str, hasta: str) -> list[dict[str, Any]]:
    payload = bound_client.request(
        "POST",
        "compra/consulta",
        query=[("pagina", "1"), ("registros", "500")],
        body={"fecha_desde": desde, "fecha_hasta": hasta},
    )
    items = payload.get("items", [])
    if not isinstance(items, list):
        return []
    return [item for item in items if isinstance(item, dict)]


def list_web_comprobantes_range_all(
    web_client: SOSContadorWebClient,
    *,
    idtipo_operacion: int,
    desde: str,
    hasta: str,
    batch_size: int = 200,
    max_pages: int = 20,
) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    seen_ids: set[str] = set()
    start = 0
    for _ in range(max_pages):
        payload = web_client.list_comprobantes_range(
            idtipo_operacion=idtipo_operacion,
            desde=desde,
            hasta=hasta,
            start=start,
            length=batch_size,
        )
        items = payload.get("items", [])
        if not isinstance(items, list) or not items:
            break
        page_rows = [item for item in items if isinstance(item, dict)]
        for item in page_rows:
            item_id = str(item.get("id") or "")
            dedupe_key = item_id or json.dumps(make_json_safe(item), ensure_ascii=False, sort_keys=True)
            if dedupe_key in seen_ids:
                continue
            seen_ids.add(dedupe_key)
            rows.append(item)
        if len(page_rows) < batch_size:
            break
        start += batch_size
    return rows


def is_web_comprobante_annulled(item: dict[str, Any]) -> bool:
    cancelado = str(find_value(item, ("cancelado", "anulado")) or "").strip().lower()
    if cancelado in {"1", "true", "si", "sí", "yes", "y"}:
        return True
    if str(find_value(item, ("fechabaja", "fecha_baja")) or "").strip():
        return True
    clipro = str(find_value(item, ("clipro", "razonsocial", "cliente")) or "").strip()
    if normalize_search_text(clipro).startswith("anulado "):
        return True
    return False


def fetch_web_comprobante_status_index(
    bound_client: SOSContadorClient,
    *,
    idtipo_operacion: int,
    desde_iso: str,
    hasta_iso: str,
) -> dict[str, dict[str, Any]]:
    web_client = SOSContadorWebClient(
        api_client=bound_client,
        explicit_cuit=bound_client.get_session().cuit,
        explicit_cuit_id=bound_client.get_session().cuit_id,
    )
    desde = parse_web_date(desde_iso, field_name="fecha desde")
    hasta = parse_web_date(hasta_iso, field_name="fecha hasta")
    index: dict[str, dict[str, Any]] = {}
    for item in list_web_comprobantes_range_all(
        web_client,
        idtipo_operacion=idtipo_operacion,
        desde=desde,
        hasta=hasta,
    ):
        item_id = str(item.get("id") or "").strip()
        if not item_id:
            continue
        index[item_id] = {
            "id": item_id,
            "comprobante": str(item.get("comprobante") or ""),
            "clipro": str(item.get("clipro") or ""),
            "cancelado": str(item.get("cancelado") or ""),
            "fechabaja": str(item.get("fechabaja") or ""),
            "annulled": is_web_comprobante_annulled(item),
        }
    return index


def expected_afip_purchase_amounts(row: dict[str, Any]) -> dict[str, Decimal]:
    type_rule = row.get("type_rule") or {}
    sign = type_rule.get("sign") or Decimal("1")
    amounts_raw = row.get("amounts_raw", {})
    raw_total = decimal_or_zero(amounts_raw.get("imp_total"))
    component_sum = sum((decimal_or_zero(value) for key, value in amounts_raw.items() if key != "imp_total"), Decimal("0"))
    residual = quantize_money(raw_total - component_sum)
    expected_nogravado = decimal_or_zero(amounts_raw.get("nogravado"))
    expected_otros = decimal_or_zero(amounts_raw.get("otros"))
    if type_rule.get("letra") == "C" and not any(
        decimal_or_zero(amounts_raw.get(key))
        for key in ("neto_0", "neto_10_5", "neto_21", "neto_27", "nogravado", "exento", "otros")
    ):
        expected_nogravado = residual
        expected_otros = Decimal("0")
    else:
        expected_otros = quantize_money(expected_otros + residual)
    return {
        "neto_10_5": quantize_money(decimal_or_zero(amounts_raw.get("neto_10_5")) * sign),
        "neto_21": quantize_money(decimal_or_zero(amounts_raw.get("neto_21")) * sign),
        "nogravado": quantize_money(expected_nogravado * sign),
        "exento": quantize_money(decimal_or_zero(amounts_raw.get("exento")) * sign),
        "otros": quantize_money(expected_otros * sign),
        "total": quantize_money(raw_total * sign),
    }


def summarize_compra_detail_for_afip(detail: dict[str, Any]) -> dict[str, Any]:
    cabecera = detail.get("cabecera", {})
    fcncnd = str(cabecera.get("fcncnd") or "")
    sign = Decimal("-1") if fcncnd == "C" else Decimal("1")
    summary = {
        "id": cabecera.get("id"),
        "cuit": digits_only(str(cabecera.get("cuit") or "")),
        "clipro": str(cabecera.get("clipro") or ""),
        "fecha": str(cabecera.get("fecha") or "")[:10],
        "fcncnd": fcncnd,
        "letra": str(cabecera.get("letra") or ""),
        "puntoventa": int(cabecera.get("puntoventa") or 0),
        "numero": int(cabecera.get("numero") or 0),
        "cae": digits_only(str(cabecera.get("cae") or "")),
        "memo": str(cabecera.get("memo") or ""),
        "neto_0": Decimal("0.00"),
        "neto_10_5": Decimal("0.00"),
        "neto_21": Decimal("0.00"),
        "neto_27": Decimal("0.00"),
        "iva_0": Decimal("0.00"),
        "iva_10_5": Decimal("0.00"),
        "iva_21": Decimal("0.00"),
        "iva_27": Decimal("0.00"),
        "nogravado": Decimal("0.00"),
        "exento": Decimal("0.00"),
        "percepcion_iibb": Decimal("0.00"),
        "otros": Decimal("0.00"),
        "total": Decimal("0.00"),
    }
    for imputacion in detail.get("imputaciones", []):
        if not isinstance(imputacion, dict):
            continue
        identifier = str(imputacion.get("identificador") or "")
        alicuota = decimal_or_zero(imputacion.get("alicuota"))
        monto = decimal_or_zero(imputacion.get("montodebe"))
        iva = decimal_or_zero(imputacion.get("iva_debe"))
        if identifier == "neto":
            rate_key = {
                Decimal("0"): "0",
                Decimal("10.5"): "10_5",
                Decimal("21"): "21",
                Decimal("27"): "27",
            }.get(alicuota)
            if rate_key:
                summary[f"neto_{rate_key}"] += monto * sign
                summary[f"iva_{rate_key}"] += iva * sign
            summary["total"] += (monto + iva) * sign
        elif identifier == "nogravado":
            summary["nogravado"] += monto * sign
            summary["total"] += monto * sign
        elif identifier == "exento":
            summary["exento"] += monto * sign
            summary["total"] += monto * sign
        elif identifier == "percepcioniibb":
            summary["percepcion_iibb"] += monto * sign
            summary["total"] += monto * sign
        elif identifier == "percepcionotra":
            summary["otros"] += monto * sign
            summary["total"] += monto * sign
    for key in (
        "neto_0",
        "neto_10_5",
        "neto_21",
        "neto_27",
        "iva_0",
        "iva_10_5",
        "iva_21",
        "iva_27",
        "nogravado",
        "exento",
        "percepcion_iibb",
        "otros",
        "total",
    ):
        summary[key] = quantize_money(summary[key])
    return summary


def build_afip_identity_only_match(existing_rows: list[dict[str, Any]], row: dict[str, Any]) -> dict[str, Any] | None:
    type_rule = row.get("type_rule") or {}
    for candidate in existing_rows:
        summary = candidate.get("summary", {})
        if summary.get("cuit") != row.get("counterparty_cuit"):
            continue
        if summary.get("fecha") != row.get("fecha"):
            continue
        if summary.get("fcncnd") != type_rule.get("fcncnd"):
            continue
        if summary.get("letra") != type_rule.get("letra"):
            continue
        if int(summary.get("puntoventa") or 0) != int(row.get("puntoventa") or 0):
            continue
        if int(summary.get("numero") or 0) != int(row.get("numero") or 0):
            continue
        return candidate
    return None


def afip_exact_match(existing_rows: list[dict[str, Any]], row: dict[str, Any]) -> dict[str, Any] | None:
    expected = expected_afip_purchase_amounts(row)
    identity_match = build_afip_identity_only_match(existing_rows, row)
    if identity_match is None:
        return None
    summary = identity_match.get("summary", {})
    if row.get("cae"):
        if summary.get("cae") == row.get("cae") or row.get("cae") in str(summary.get("memo") or ""):
            return identity_match
    if (
        summary.get("clipro") == row.get("counterparty_name")
        and summary.get("neto_10_5") == expected.get("neto_10_5")
        and summary.get("neto_21") == expected.get("neto_21")
        and summary.get("nogravado") == expected.get("nogravado")
        and summary.get("exento") == expected.get("exento")
        and summary.get("otros") == expected.get("otros")
        and summary.get("total") == expected.get("total")
    ):
        return identity_match
    return None


def build_afip_compra_body(row: dict[str, Any], provider_id: str | int) -> dict[str, Any]:
    type_rule = row.get("type_rule") or {}
    sign = type_rule.get("sign") or Decimal("1")
    amounts_raw = row.get("amounts_raw", {})
    expected = expected_afip_purchase_amounts(row)
    imputa: list[dict[str, Any]] = []
    if decimal_or_zero(amounts_raw.get("neto_21")):
        imputa.append({"i": "neto", "a": 21.0, "v": float(quantize_money(decimal_or_zero(amounts_raw.get("neto_21")) * sign))})
    if decimal_or_zero(amounts_raw.get("neto_10_5")):
        imputa.append({"i": "neto", "a": 10.5, "v": float(quantize_money(decimal_or_zero(amounts_raw.get("neto_10_5")) * sign))})
    if decimal_or_zero(amounts_raw.get("neto_27")):
        imputa.append({"i": "neto", "a": 27.0, "v": float(quantize_money(decimal_or_zero(amounts_raw.get("neto_27")) * sign))})
    if decimal_or_zero(amounts_raw.get("neto_0")):
        imputa.append({"i": "neto", "a": 0.0, "v": float(quantize_money(decimal_or_zero(amounts_raw.get("neto_0")) * sign))})
    if expected.get("nogravado"):
        imputa.append({"i": "nogravado", "a": 0.0, "v": float(expected["nogravado"])})
    if expected.get("exento"):
        imputa.append({"i": "exento", "a": 0.0, "v": float(expected["exento"])})
    if expected.get("otros"):
        imputa.append({"i": "percepcionotra", "a": 0.0, "v": float(expected["otros"])})
    return {
        "fecha": row.get("fecha"),
        "idclipro": int(provider_id),
        "cuitclipro": row.get("counterparty_cuit"),
        "fcncnd": type_rule.get("fcncnd"),
        "letra": type_rule.get("letra"),
        "puntoventa": int(row.get("puntoventa") or 0),
        "numero": int(row.get("numero") or 0),
        "numerohasta": int(row.get("numerohasta") or 0),
        "obtienecae": False,
        "fechaiva": row.get("fecha"),
        "idprovinciaiibb": AFIP_PURCHASE_PROVIDER_PROVINCE_ID,
        "memo": f"CAE: {row.get('cae')} - " if row.get("cae") else "",
        "referencia": "",
        "descuento": 0,
        "uniqueid": str(uuid.uuid4()),
        "controlainconsistencia": 0,
        "imputaciones": [{"imputa": imputa}],
        "productos": [],
    }


def build_afip_provider_create_body(row: dict[str, Any]) -> dict[str, Any]:
    type_rule = row.get("type_rule") or {}
    idtipocondicioniva = AFIP_PURCHASE_PROVIDER_COND_IVA_C if type_rule.get("letra") == "C" else AFIP_PURCHASE_PROVIDER_COND_IVA_A
    return {
        "cuit": row.get("counterparty_cuit"),
        "clipro": row.get("counterparty_name"),
        "idprovincia": AFIP_PURCHASE_PROVIDER_PROVINCE_ID,
        "idtipocondicioniva": idtipocondicioniva,
        "email": "",
    }


def build_afip_mis_comprobantes_draft(args: argparse.Namespace, client: SOSContadorClient) -> dict[str, Any]:
    source_path = Path(str(getattr(args, "source", "") or "")).expanduser()
    if not source_path.exists():
        raise CLIError(f"No existe el archivo AFIP indicado: {source_path}")
    workbook_payload = read_afip_mis_comprobantes_workbook(source_path)
    if workbook_payload.get("operation_kind") != "compra":
        raise CLIError("Por ahora el flujo AFIP validado en esta skill solo soporta 'Mis Comprobantes Recibidos'.")
    bound_client, work_target = resolve_afip_work_target(client, args, workbook_work_cuit=workbook_payload["work_cuit"])
    provider_catalog = fetch_all_clientes_catalog(bound_client)
    fechas = [row.get("fecha") for row in workbook_payload["rows"] if row.get("fecha")]
    desde = min(fechas) if fechas else None
    hasta = max(fechas) if fechas else None
    if not desde or not hasta:
        raise CLIError("El archivo AFIP no contiene fechas validas para analizar.")
    web_status_index = fetch_web_comprobante_status_index(
        bound_client,
        idtipo_operacion=4,
        desde_iso=desde,
        hasta_iso=hasta,
    )
    existing_rows: list[dict[str, Any]] = []
    for item in fetch_compra_consulta_items(bound_client, desde=desde, hasta=hasta):
        item_id = str(item.get("id") or "")
        status_info = web_status_index.get(item_id)
        if status_info and status_info.get("annulled"):
            continue
        detail = bound_client.request("GET", f"compra/detalle/{item.get('id')}")
        existing_rows.append({"id": item.get("id"), "summary": summarize_compra_detail_for_afip(detail)})

    draft_rows: list[dict[str, Any]] = []
    stats = {"ya_cargado": 0, "pendiente": 0, "verificar": 0}
    warnings: list[str] = []
    for row in workbook_payload["rows"]:
        provider = provider_catalog.get(row.get("counterparty_cuit") or "")
        expected = expected_afip_purchase_amounts(row)
        document_text = format_comprobante_text(
            (row.get("type_rule") or {}).get("letra"),
            row.get("puntoventa"),
            row.get("numero"),
        )
        status = "pendiente"
        reason = ""
        match = None
        identity_only = None
        if row.get("validation_errors"):
            status = "verificar"
            reason = "; ".join(row["validation_errors"])
        else:
            match = afip_exact_match(existing_rows, row)
            if match is not None:
                status = "ya_cargado"
                reason = "Coincidencia exacta en SOS."
            else:
                identity_only = build_afip_identity_only_match(existing_rows, row)
                if identity_only is not None:
                    status = "verificar"
                    reason = "Ya existe un comprobante con la misma identidad, pero los importes o imputaciones no coinciden."
        if status == "pendiente" and not provider:
            reason = "Proveedor faltante: se creara primero por API."
        draft_row = {
            "source_row": row.get("source_row"),
            "status": status,
            "reason": reason,
            "fecha": row.get("fecha"),
            "documento": document_text,
            "counterparty_name": row.get("counterparty_name"),
            "counterparty_cuit": row.get("counterparty_cuit"),
            "cae": row.get("cae"),
            "expected_amounts": {key: format_decimal_string(value) for key, value in expected.items()},
            "provider_exists": bool(provider),
            "provider_id": str(find_value(provider or {}, ("id", "idclipro")) or ""),
            "existing_match_id": str((match or identity_only or {}).get("id") or ""),
            "row_payload": row,
        }
        if status == "verificar":
            warnings.append(f"Fila {row.get('source_row')}: {reason}")
        stats[status] += 1
        draft_rows.append(draft_row)

    draft = {
        "draft_id": uuid.uuid4().hex[:12],
        "draft_kind": AFIP_DRAFT_KIND,
        "source": str(source_path),
        "contexto": {
            "operation_kind": workbook_payload["operation_kind"],
            "row1": workbook_payload["row1"],
            "cuit_trabajo": work_target,
            "fecha_desde": desde,
            "fecha_hasta": hasta,
        },
        "rows": draft_rows,
        "stats": stats,
        "validacion": {"warnings": warnings},
    }
    draft = make_json_safe(draft)
    saved_path = save_draft_payload(draft)
    draft["draft_file"] = str(saved_path)
    return draft


def build_afip_draft_preview(draft: dict[str, Any]) -> dict[str, Any]:
    contexto = draft.get("contexto", {})
    stats = draft.get("stats", {})
    rows = draft.get("rows", [])
    pending_rows = [
        {
            "Fila": row.get("source_row"),
            "Fecha": row.get("fecha"),
            "Comprobante": row.get("documento"),
            "Proveedor": row.get("counterparty_name"),
            "CUIT": row.get("counterparty_cuit"),
            "Total": row.get("expected_amounts", {}).get("total"),
            "Proveedor en SOS": "si" if row.get("provider_exists") else "crear",
        }
        for row in rows
        if row.get("status") == "pendiente"
    ]
    verify_rows = [
        {
            "Fila": row.get("source_row"),
            "Fecha": row.get("fecha"),
            "Comprobante": row.get("documento"),
            "Proveedor": row.get("counterparty_name"),
            "Motivo": row.get("reason"),
        }
        for row in rows
        if row.get("status") == "verificar"
    ]
    loaded_rows = [
        {
            "Fila": row.get("source_row"),
            "Comprobante": row.get("documento"),
            "Proveedor": row.get("counterparty_name"),
            "ID SOS": row.get("existing_match_id"),
        }
        for row in rows
        if row.get("status") == "ya_cargado"
    ]
    preview = {
        "titulo": "afip_draft",
        "contexto": [
            {
                "CUIT de trabajo": contexto.get("cuit_trabajo", {}).get("cuit"),
                "Contribuyente": contexto.get("cuit_trabajo", {}).get("nombre"),
                "Operacion": contexto.get("operation_kind"),
                "Archivo": draft.get("source"),
                "Periodo": f"{contexto.get('fecha_desde')} a {contexto.get('fecha_hasta')}",
            }
        ],
        "tablas": [],
        "totales": {
            "Pendientes": stats.get("pendiente", 0),
            "Ya cargados": stats.get("ya_cargado", 0),
            "Verificar": stats.get("verificar", 0),
        },
        "advertencias": list(draft.get("validacion", {}).get("warnings", [])),
    }
    if pending_rows:
        preview["tablas"].append({"titulo": "Pendientes", "rows": pending_rows})
    if loaded_rows:
        preview["tablas"].append({"titulo": "Ya cargados", "rows": loaded_rows})
    if verify_rows:
        preview["tablas"].append({"titulo": "Verificar", "rows": verify_rows})
    return preview


def render_afip_draft_markdown(draft: dict[str, Any]) -> str:
    return render_business_preview_markdown(build_afip_draft_preview(draft))


def command_afip_draft(args: argparse.Namespace, client: SOSContadorClient) -> Any:
    draft = build_afip_mis_comprobantes_draft(args, client)
    preview = build_afip_draft_preview(draft)
    preview_markdown = render_afip_draft_markdown(draft)
    if getattr(args, "json_out", None):
        destination = Path(args.json_out)
        save_json_file(destination, draft)
        draft["json_written_to"] = str(destination)
    return {
        "draft_id": draft.get("draft_id"),
        "draft_file": draft.get("draft_file"),
        "preview_markdown": preview_markdown,
        "payload": draft,
        "business_preview": preview,
    }


def command_afip_import(args: argparse.Namespace, client: SOSContadorClient) -> Any:
    draft = build_afip_mis_comprobantes_draft(args, client)
    preview = build_afip_draft_preview(draft)
    preview_markdown = render_afip_draft_markdown(draft)
    context = draft.get("contexto", {})
    work_target = context.get("cuit_trabajo", {})
    set_mutation_preview(args, business_preview=preview, preview_markdown=preview_markdown)
    ensure_mutation_allowed("PUT", "afip/import", None, {"source": draft.get("source")}, args)

    bound_client = ensure_bound_client(
        client,
        cuit=str(work_target.get("cuit") or ""),
        cuit_id=str(work_target.get("cuit_id") or ""),
    )
    provider_catalog = fetch_all_clientes_catalog(bound_client)
    created_providers: list[dict[str, Any]] = []
    created_compras: list[dict[str, Any]] = []

    for draft_row in draft.get("rows", []):
        if draft_row.get("status") != "pendiente":
            continue
        row_payload = draft_row.get("row_payload", {})
        counterparty_cuit = row_payload.get("counterparty_cuit") or ""
        provider = provider_catalog.get(counterparty_cuit)
        if provider is None:
            provider_body = build_afip_provider_create_body(row_payload)
            provider_result = bound_client.request("POST", "cliente", body=provider_body)
            provider = dict(provider_body)
            provider["id"] = provider_result.get("id")
            provider_catalog[counterparty_cuit] = provider
            created_providers.append(
                {
                    "cuit": counterparty_cuit,
                    "clipro": provider_body.get("clipro"),
                    "id": provider_result.get("id"),
                }
            )
        provider_id = find_value(provider, ("id", "idclipro"))
        if not provider_id:
            raise CLIError(f"No se pudo resolver el proveedor para la fila {draft_row.get('source_row')}.")
        compra_body = build_afip_compra_body(row_payload, provider_id)
        compra_result = bound_client.request("PUT", "compra/0", body=compra_body)
        created_compras.append(
            {
                "source_row": draft_row.get("source_row"),
                "id": compra_result.get("id"),
                "documento": draft_row.get("documento"),
                "proveedor": draft_row.get("counterparty_name"),
            }
        )

    created_index = fetch_web_comprobante_status_index(
        bound_client,
        idtipo_operacion=4,
        desde_iso=str(context.get("fecha_desde") or ""),
        hasta_iso=str(context.get("fecha_hasta") or ""),
    )
    post_create_checks: list[dict[str, Any]] = []
    invalid_annulled: list[dict[str, Any]] = []
    for item in created_compras:
        status_info = created_index.get(str(item.get("id") or ""))
        check = {
            "id": item.get("id"),
            "documento": item.get("documento"),
            "proveedor": item.get("proveedor"),
            "annulled": bool((status_info or {}).get("annulled")),
            "cancelado": (status_info or {}).get("cancelado", ""),
            "fechabaja": (status_info or {}).get("fechabaja", ""),
        }
        post_create_checks.append(check)
        if check["annulled"]:
            invalid_annulled.append(check)

    if invalid_annulled:
        details = ", ".join(
            f"{item.get('documento') or item.get('id')} (cancelado={item.get('cancelado')}, fechabaja={item.get('fechabaja')})"
            for item in invalid_annulled
        )
        raise CLIError(
            "La verificacion posterior detecto compras anuladas inmediatamente despues de la carga: "
            f"{details}. Revise el flujo antes de considerar exitosa la importacion."
        )

    return {
        "draft_id": draft.get("draft_id"),
        "draft_file": draft.get("draft_file"),
        "created_providers": created_providers,
        "created_compras": created_compras,
        "post_create_checks": post_create_checks,
        "stats": draft.get("stats"),
        "preview_markdown": preview_markdown,
    }


COMPRA_RATE_KEYS = {
    Decimal("0"): "0",
    Decimal("10.5"): "10_5",
    Decimal("21"): "21",
    Decimal("27"): "27",
}
COMPRA_AMOUNT_KEYS = (
    "neto_0",
    "neto_10_5",
    "neto_21",
    "neto_27",
    "iva_0",
    "iva_10_5",
    "iva_21",
    "iva_27",
    "nogravado",
    "exento",
    "percepcion_iibb",
    "otros",
    "total",
)
COMPRA_DISCOUNT_KEYS = ("descuento_0", "descuento_10_5", "descuento_21", "descuento_27", "descuento_global")


def extract_compra_line_amount(line: str) -> Decimal | None:
    matches = re.findall(r"-?\d[\d., ]*[,.]\d{2,4}", str(line or ""))
    if not matches:
        return None
    try:
        return quantize_money(smart_decimal_from_text(matches[-1]))
    except CLIError:
        return None


def extract_compra_rate(line: str, *, prefix: str) -> Decimal | None:
    pattern = re.compile(
        rf"(?i){prefix}[^0-9]*(10\s*[,.]\s*5(?:0)?|21(?:\s*[,.]\s*0+)?|27(?:\s*[,.]\s*0+)?|0(?:\s*[,.]\s*0+)?)"
    )
    match = pattern.search(str(line or ""))
    if not match:
        return None
    rate_text = re.sub(r"\s+", "", match.group(1)).replace(",", ".")
    try:
        rate = Decimal(rate_text)
    except InvalidOperation:
        return None
    return rate if rate in COMPRA_RATE_KEYS else None


def extract_compra_document_date(lines: list[str], explicit_date: Any = None) -> str:
    if explicit_date not in (None, ""):
        parsed = guess_date_from_text(str(explicit_date)) or parse_isoish_date(str(explicit_date))
        if parsed and re.fullmatch(r"\d{4}-\d{2}-\d{2}", parsed):
            return parsed
        raise CLIError(f"fecha inválida '{explicit_date}'. Use DD/MM/YYYY o YYYY-MM-DD.")
    pattern = re.compile(r"\b(?:\d{1,2}[/-]\d{1,2}[/-]\d{2,4}|\d{4}[/-]\d{1,2}[/-]\d{1,2})\b")
    candidates: list[str] = []
    for line in lines:
        normalized = normalize_search_text(line)
        if "fecha" not in normalized or "inicio de actividades" in normalized:
            continue
        match = pattern.search(line)
        if match:
            candidates.append(match.group(0))
    if len(set(candidates)) == 1:
        return guess_date_from_text(candidates[0]) or ""
    return ""


def extract_compra_document_identity(lines: list[str]) -> dict[str, Any]:
    result = {"fcncnd": "", "letra": "", "puntoventa": "", "numero": ""}
    for line in lines:
        normalized = normalize_search_text(line)
        if not any(token in normalized for token in ("factura", "nota de credito", "nota credito", "nota de debito", "nota debito")):
            continue
        if "nota" in normalized and "credito" in normalized:
            result["fcncnd"] = "C"
        elif "nota" in normalized and "debito" in normalized:
            result["fcncnd"] = "D"
        else:
            result["fcncnd"] = "F"
        letter_match = re.search(r"(?i)(?:factura|nota(?:\s+de)?\s+(?:credito|debito))\s*[\"']?([A-Z])\b", line)
        if letter_match:
            result["letra"] = letter_match.group(1).upper()
        number_match = re.search(r"\b(\d{1,5})\s*[-–]\s*(\d{1,8})\b", line)
        if number_match:
            result["puntoventa"] = str(int(number_match.group(1)))
            result["numero"] = str(int(number_match.group(2)))
        if result["letra"] and result["puntoventa"] and result["numero"]:
            break
    return result


def extract_compra_document_amounts(lines: list[str]) -> dict[str, Decimal]:
    amounts: dict[str, Decimal] = {}
    neto_gravado_total: Decimal | None = None
    otros_candidate: Decimal | None = None
    for line in lines:
        normalized = normalize_search_text(line)
        amount = extract_compra_line_amount(line)
        if amount is None:
            continue
        if "descuento" in normalized:
            rate = extract_compra_rate(line, prefix=r"(?:descuento|dto)")
            if rate is not None:
                amounts[f"descuento_{COMPRA_RATE_KEYS[rate]}"] = amount
            else:
                amounts["descuento_global"] = amount
            continue
        if "neto" in normalized and "gravado" in normalized:
            rate = extract_compra_rate(line, prefix=r"(?:neto(?:\s+gravado)?|gravado)")
            if rate is None:
                neto_gravado_total = amount
            else:
                amounts[f"neto_{COMPRA_RATE_KEYS[rate]}"] = amount
            continue
        if "alicuota" in normalized or re.search(r"\biva\b", normalized):
            rate = extract_compra_rate(line, prefix=r"(?:al[ií]cuota|iva)")
            if rate is not None:
                amounts[f"iva_{COMPRA_RATE_KEYS[rate]}"] = amount
            continue
        if any(token in normalized for token in ("iibb", "ingresos brutos", "arba")):
            amounts["percepcion_iibb"] = amount
            continue
        if "no gravado" in normalized or "nogravado" in normalized:
            amounts["nogravado"] = amount
            continue
        if "exento" in normalized:
            amounts["exento"] = amount
            continue
        if "otros tributos" in normalized:
            otros_candidate = amount
            continue
        if normalized == "total" or normalized.startswith("total "):
            amounts["total"] = amount
    if neto_gravado_total is not None:
        amounts["neto_gravado_total"] = neto_gravado_total
    if otros_candidate is not None and not amounts.get("percepcion_iibb"):
        amounts["otros"] = otros_candidate
    return amounts


def normalize_compra_amounts(raw_fields: dict[str, Any]) -> tuple[dict[str, Decimal], list[str]]:
    warnings: list[str] = []
    amounts: dict[str, Decimal] = {}
    for key in (*COMPRA_AMOUNT_KEYS, *COMPRA_DISCOUNT_KEYS, "neto_gravado_total"):
        value = raw_fields.get(key)
        if value in (None, ""):
            continue
        try:
            amounts[key] = quantize_money(smart_decimal_from_text(value))
        except CLIError as exc:
            warnings.append(f"{key}: {exc}")

    present_rates = [
        rate
        for rate, key in COMPRA_RATE_KEYS.items()
        if amounts.get(f"iva_{key}") is not None or amounts.get(f"neto_{key}") is not None
    ]
    neto_total = amounts.get("neto_gravado_total")
    if neto_total is not None and len(present_rates) == 1:
        key = COMPRA_RATE_KEYS[present_rates[0]]
        amounts.setdefault(f"neto_{key}", neto_total)
    elif neto_total is not None and len(present_rates) > 1:
        for rate in present_rates:
            key = COMPRA_RATE_KEYS[rate]
            if amounts.get(f"neto_{key}") is None and amounts.get(f"iva_{key}") is not None and rate:
                amounts[f"neto_{key}"] = quantize_money(amounts[f"iva_{key}"] * Decimal("100") / rate)

    for rate, key in COMPRA_RATE_KEYS.items():
        neto_key = f"neto_{key}"
        iva_key = f"iva_{key}"
        if amounts.get(neto_key) is not None and amounts.get(iva_key) is None:
            amounts[iva_key] = quantize_money(amounts[neto_key] * rate / Decimal("100"))
        elif rate and amounts.get(neto_key) is not None and amounts.get(iva_key) is not None:
            expected_iva = quantize_money(amounts[neto_key] * rate / Decimal("100"))
            if abs(expected_iva - amounts[iva_key]) > Decimal("0.02"):
                warnings.append(
                    f"El IVA de la alícuota {format_decimal_string(rate)} no coincide con su neto gravado."
                )

    if amounts.get("descuento_global") is not None and len(present_rates) > 1:
        warnings.append(
            "El descuento global no puede asignarse con seguridad entre varias alícuotas. Indique los netos finales por alícuota."
        )

    neto_sum = sum((amounts.get(f"neto_{key}", Decimal("0")) for key in COMPRA_RATE_KEYS.values()), Decimal("0"))
    if neto_total is not None and abs(quantize_money(neto_sum - neto_total)) > Decimal("0.01"):
        warnings.append(
            f"Los netos por alícuota ({format_decimal_string(neto_sum)}) no coinciden con el neto gravado total ({format_decimal_string(neto_total)})."
        )

    calculated_total = sum((amounts.get(f"neto_{key}", Decimal("0")) for key in COMPRA_RATE_KEYS.values()), Decimal("0"))
    calculated_total += sum((amounts.get(f"iva_{key}", Decimal("0")) for key in COMPRA_RATE_KEYS.values()), Decimal("0"))
    calculated_total += sum((amounts.get(key, Decimal("0")) for key in ("nogravado", "exento", "percepcion_iibb", "otros")), Decimal("0"))
    if amounts.get("total") is not None and abs(quantize_money(calculated_total - amounts["total"])) > Decimal("0.01"):
        warnings.append(
            f"La suma de netos, IVA y tributos ({format_decimal_string(calculated_total)}) no coincide con el total ({format_decimal_string(amounts['total'])})."
        )
    return amounts, warnings


def load_compra_document_overrides(args: argparse.Namespace, source_count: int) -> list[dict[str, Any]]:
    if (getattr(args, "document_json", None) or []) and (getattr(args, "document_file", None) or []):
        raise CLIError("Use --document-json o --document-file, no ambas opciones.")
    raw_items: list[Any] = []
    for raw in getattr(args, "document_json", None) or []:
        try:
            raw_items.append(json.loads(raw))
        except json.JSONDecodeError as exc:
            raise CLIError(f"--document-json inválido: {exc}") from exc
    for file_name in getattr(args, "document_file", None) or []:
        raw_items.append(load_json_file(Path(file_name), None))
    if not raw_items:
        return [{} for _ in range(source_count)]
    if len(raw_items) != source_count:
        raise CLIError("Debe indicar un --document-json/--document-file por cada --source.")
    if not all(isinstance(item, dict) for item in raw_items):
        raise CLIError("Cada corrección documental debe ser un objeto JSON.")
    return [dict(item) for item in raw_items]


def build_compra_source_fields(
    extracted_source: dict[str, Any],
    *,
    work_cuit: str,
    overrides: dict[str, Any],
) -> tuple[dict[str, Any], list[str]]:
    lines = extracted_source.get("lines", [])
    issues: list[str] = []
    identity = extract_compra_document_identity(lines)
    extracted_amounts = extract_compra_document_amounts(lines)
    raw_fields: dict[str, Any] = {
        **identity,
        "fecha": extract_compra_document_date(lines),
        **extracted_amounts,
    }
    counterparty_cuits = collect_counterparty_cuits(lines, work_cuit=work_cuit, explicit_cuit=overrides.get("proveedor_cuit"))
    if len(counterparty_cuits) == 1:
        raw_fields["proveedor_cuit"] = counterparty_cuits[0]
    raw_fields.update({key: value for key, value in overrides.items() if value not in (None, "")})
    if decimal_or_zero(extracted_amounts.get("otros")):
        tax_override_keys = ("nogravado", "exento", "percepcion_iibb", "otros")
        has_explicit_tax_classification = any(
            key in overrides and overrides.get(key) not in (None, "")
            for key in tax_override_keys
        )
        if has_explicit_tax_classification:
            if "otros" not in overrides:
                raw_fields.pop("otros", None)
        else:
            issues.append(
                "El comprobante informa otros tributos sin clasificación. Indique explícitamente si corresponden a no gravado, exento, percepción de IIBB u otra percepción."
            )
    if raw_fields.get("fecha"):
        raw_fields["fecha"] = extract_compra_document_date(lines, raw_fields.get("fecha"))
    if raw_fields.get("proveedor_cuit"):
        try:
            raw_fields["proveedor_cuit"] = require_valid_argentina_cuit(
                raw_fields["proveedor_cuit"],
                label="CUIT del proveedor",
            )
        except CLIError as exc:
            issues.append(str(exc))
            raw_fields["proveedor_cuit"] = ""
    for key in ("puntoventa", "numero", "numerohasta"):
        if raw_fields.get(key) not in (None, ""):
            raw_fields[key] = int(digits_only(str(raw_fields[key])) or "0")
    if raw_fields.get("numero") and not raw_fields.get("numerohasta"):
        raw_fields["numerohasta"] = raw_fields["numero"]
    raw_fields["letra"] = str(raw_fields.get("letra") or "").strip().upper()
    raw_fields["fcncnd"] = str(raw_fields.get("fcncnd") or "F").strip().upper()
    amounts, amount_issues = normalize_compra_amounts(raw_fields)
    raw_fields["amounts"] = amounts
    return raw_fields, [*issues, *amount_issues]


def fetch_compra_items_complete(
    bound_client: SOSContadorClient,
    *,
    desde: str,
    hasta: str,
) -> tuple[list[dict[str, Any]], bool]:
    items = fetch_compra_consulta_items(bound_client, desde=desde, hasta=hasta)
    if len(items) < 50:
        return items, False
    start = datetime.date.fromisoformat(desde)
    end = datetime.date.fromisoformat(hasta)
    if start >= end:
        return items, True
    midpoint = start + datetime.timedelta(days=(end - start).days // 2)
    left, left_truncated = fetch_compra_items_complete(
        bound_client,
        desde=start.isoformat(),
        hasta=midpoint.isoformat(),
    )
    right, right_truncated = fetch_compra_items_complete(
        bound_client,
        desde=(midpoint + datetime.timedelta(days=1)).isoformat(),
        hasta=end.isoformat(),
    )
    merged: list[dict[str, Any]] = []
    seen: set[str] = set()
    for item in [*left, *right]:
        item_id = str(item.get("id") or "")
        key = item_id or json.dumps(make_json_safe(item), ensure_ascii=False, sort_keys=True)
        if key in seen:
            continue
        seen.add(key)
        merged.append(item)
    return merged, left_truncated or right_truncated


def compra_row_provider_cuit(row: dict[str, Any]) -> str:
    clipro = row.get("clipro")
    if isinstance(clipro, dict):
        return digits_only(str(clipro.get("cuit") or ""))
    return digits_only(str(row.get("cuit") or ""))


def compra_document_code(*, fcncnd: Any, letra: Any, puntoventa: Any, numero: Any) -> str:
    prefix = {"F": "F", "C": "C", "D": "D"}.get(str(fcncnd or "").upper(), str(fcncnd or "").upper())
    return f"{prefix}{str(letra or '').upper()}-{int(puntoventa or 0):04d}-{int(numero or 0):08d}"


def compra_row_is_active(row: dict[str, Any]) -> bool:
    true_values = {"1", "true", "si", "sí", "yes", "y"}
    for key in ("eliminado", "archivado", "cancelado", "anulado"):
        value = row.get(key)
        if isinstance(value, bool) and value:
            return False
        if isinstance(value, (int, float, Decimal)) and value != 0:
            return False
        if isinstance(value, str) and value.strip().lower() in true_values:
            return False
    return not bool(str(row.get("fechabaja") or row.get("fecha_baja") or "").strip())


def compra_summary_matches_expected(
    summary: dict[str, Any],
    expected: dict[str, Any],
    *,
    fcncnd: Any = None,
) -> bool:
    credit_note = str(fcncnd or "").strip().upper() == "C"
    for key in COMPRA_AMOUNT_KEYS:
        if key not in expected:
            continue
        actual_value = quantize_money(summary.get(key))
        if credit_note:
            actual_value = abs(actual_value)
        if actual_value != quantize_money(expected.get(key)):
            return False
    return True


def find_compra_duplicate(
    bound_client: SOSContadorClient,
    *,
    fields: dict[str, Any],
) -> dict[str, Any] | None:
    fecha = str(fields.get("fecha") or "")
    proveedor_cuit = digits_only(str(fields.get("proveedor_cuit") or ""))
    if not fecha or not proveedor_cuit:
        return None
    rows, truncated = fetch_compra_items_complete(bound_client, desde=fecha, hasta=fecha)
    expected_code = compra_document_code(
        fcncnd=fields.get("fcncnd"),
        letra=fields.get("letra"),
        puntoventa=fields.get("puntoventa"),
        numero=fields.get("numero"),
    )
    for row in rows:
        if compra_row_provider_cuit(row) != proveedor_cuit:
            continue
        actual_code = normalize_comprobante_reference(str(row.get("factura") or ""))
        normalized_expected = normalize_comprobante_reference(expected_code)
        if actual_code != normalized_expected:
            continue
        row_id = row.get("id")
        detail = bound_client.request("GET", f"compra/detalle/{row_id}")
        summary = summarize_compra_detail_for_afip(detail)
        active = compra_row_is_active(row)
        return {
            "kind": "exact"
            if active
            and compra_summary_matches_expected(
                summary,
                fields.get("amounts", {}),
                fcncnd=fields.get("fcncnd"),
            )
            else "identity_conflict",
            "id": row_id,
            "active": active,
            "summary": summary,
        }
    if truncated:
        return {"kind": "truncated", "reason": "La consulta del día alcanzó el límite de 50 compras."}
    return None


def previous_month_start(value: datetime.date) -> datetime.date:
    if value.month == 1:
        return datetime.date(value.year - 1, 12, 1)
    return datetime.date(value.year, value.month - 1, 1)


def month_end(value: datetime.date) -> datetime.date:
    if value.month == 12:
        return datetime.date(value.year + 1, 1, 1) - datetime.timedelta(days=1)
    return datetime.date(value.year, value.month + 1, 1) - datetime.timedelta(days=1)


def fetch_recent_supplier_purchase_details(
    bound_client: SOSContadorClient,
    *,
    proveedor_cuit: str,
    fecha: str,
    limit: int = 3,
    max_months: int = 24,
) -> list[dict[str, Any]]:
    target_date = datetime.date.fromisoformat(fecha)
    cursor = target_date.replace(day=1)
    details: list[dict[str, Any]] = []
    seen_ids: set[str] = set()
    for month_index in range(max_months):
        end = target_date if month_index == 0 else month_end(cursor)
        rows, _ = fetch_compra_items_complete(bound_client, desde=cursor.isoformat(), hasta=end.isoformat())
        matches = [
            row
            for row in rows
            if compra_row_provider_cuit(row) == proveedor_cuit
            and compra_row_is_active(row)
        ]
        matches.sort(key=lambda row: str(row.get("fecha") or ""), reverse=True)
        for row in matches:
            row_id = str(row.get("id") or "")
            if not row_id or row_id in seen_ids:
                continue
            seen_ids.add(row_id)
            details.append(bound_client.request("GET", f"compra/detalle/{row_id}"))
            if len(details) >= limit:
                return details
        cursor = previous_month_start(cursor)
    return details


def infer_compra_historical_defaults(details: list[dict[str, Any]]) -> tuple[dict[str, Any], list[str]]:
    warnings: list[str] = []
    defaults: dict[str, Any] = {}
    definitions = {
        "idcuenta": ("idcuenta",),
        "idcentrocosto": ("idcentrocosto",),
        "idprovinciaiibb": ("idprovinciaiibb", "idprovincia"),
    }
    for target_key, source_keys in definitions.items():
        values: set[str] = set()
        for detail in details:
            cabecera = detail.get("cabecera", {}) if isinstance(detail, dict) else {}
            for source_key in source_keys:
                value = cabecera.get(source_key)
                if value not in (None, ""):
                    values.add(str(value))
                    break
        if len(values) == 1:
            defaults[target_key] = next(iter(values))
        elif len(values) > 1:
            warnings.append(f"Los antecedentes activos no coinciden en {target_key}.")
    if details:
        cabecera = details[0].get("cabecera", {})
        defaults["cuenta_nombre"] = cabecera.get("cuenta")
        defaults["centrocosto_nombre"] = cabecera.get("centrocosto")
        defaults["antecedente_id"] = cabecera.get("id")
    return defaults, warnings


def compra_payload_hash(body: dict[str, Any]) -> str:
    canonical = json.dumps(body, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return hashlib.sha256(canonical.encode("utf-8")).hexdigest()


def build_compra_body_from_fields(fields: dict[str, Any]) -> dict[str, Any]:
    amounts = fields.get("amounts", {})
    imputa: list[dict[str, Any]] = []
    for rate in (Decimal("21"), Decimal("10.5"), Decimal("27"), Decimal("0")):
        key = COMPRA_RATE_KEYS[rate]
        value = decimal_or_zero(amounts.get(f"neto_{key}"))
        if value:
            imputa.append({"i": "neto", "a": float(rate), "v": float(quantize_money(value))})
    for identifier, key in (
        ("nogravado", "nogravado"),
        ("exento", "exento"),
        ("percepcioniibb", "percepcion_iibb"),
        ("percepcionotra", "otros"),
    ):
        value = decimal_or_zero(amounts.get(key))
        if value:
            imputa.append({"i": identifier, "a": 0.0, "v": float(quantize_money(value))})
    account_id = str(fields.get("idcuenta") or "")
    if not account_id:
        raise CLIError("No se puede construir la compra sin idcuenta.")
    return {
        "fecha": fields.get("fecha"),
        "fechaiva": fields.get("fecha"),
        "idclipro": int(str(fields.get("idclipro"))),
        "cuitclipro": fields.get("proveedor_cuit"),
        "fcncnd": fields.get("fcncnd"),
        "letra": fields.get("letra"),
        "puntoventa": int(fields.get("puntoventa") or 0),
        "numero": int(fields.get("numero") or 0),
        "numerohasta": int(fields.get("numerohasta") or fields.get("numero") or 0),
        "obtienecae": False,
        "idprovinciaiibb": int(str(fields.get("idprovinciaiibb"))),
        "idcentrocosto": int(str(fields.get("idcentrocosto"))),
        "memo": str(fields.get("memo") or ""),
        "referencia": str(fields.get("referencia") or ""),
        "descuento": 0,
        "uniqueid": str(uuid.uuid4()),
        "controlainconsistencia": 0,
        "imputaciones": [{"imputa": imputa, "cuid": int(account_id)}],
        "productos": [],
    }


def build_compra_document_draft(args: argparse.Namespace, client: SOSContadorClient) -> dict[str, Any]:
    sources = [Path(path).expanduser().resolve() for path in (getattr(args, "source", None) or [])]
    if not sources:
        raise CLIError("compra draft requiere al menos un --source.")
    extracted_sources = [extract_source_document(path) for path in sources]
    overrides = load_compra_document_overrides(args, len(sources))
    explicit_work_target = resolve_work_cuit_target_from_args(client, args, required=False)
    if explicit_work_target is not None:
        work_target = explicit_work_target
    else:
        inferred_targets = [extract_document_work_cuit(source.get("lines", []), client) for source in extracted_sources]
        inferred_cuits = {digits_only(str(item.get("cuit") or "")) for item in inferred_targets if item.get("cuit")}
        if len(inferred_cuits) != 1:
            raise CLIError("No se pudo inferir una única CUIT de trabajo. Indíquela explícitamente.")
        inferred_cuit = next(iter(inferred_cuits))
        work_target = resolve_work_cuit_target(client, explicit_cuit=inferred_cuit, required=True)
    bound_client = ensure_bound_client(
        client,
        cuit=str(work_target.get("cuit") or ""),
        cuit_id=str(work_target.get("cuit_id") or ""),
    )

    rows: list[dict[str, Any]] = []
    provider_cache: dict[str, dict[str, Any] | None] = {}
    history_cache: dict[tuple[str, str], list[dict[str, Any]]] = {}
    stats = {"pendiente": 0, "ya_cargado": 0, "verificar": 0}
    draft_warnings: list[str] = []
    for index, (source_path, extracted, document_overrides) in enumerate(zip(sources, extracted_sources, overrides), start=1):
        fields, issues = build_compra_source_fields(
            extracted,
            work_cuit=str(work_target.get("cuit") or ""),
            overrides=document_overrides,
        )
        warnings = [*extracted.get("warnings", []), *extracted.get("capabilities", [])]
        proveedor_cuit = digits_only(str(fields.get("proveedor_cuit") or ""))
        provider: dict[str, Any] | None = None
        if proveedor_cuit:
            if proveedor_cuit not in provider_cache:
                try:
                    provider_cache[proveedor_cuit] = resolve_cliente_match(
                        bound_client,
                        explicit_id=None,
                        cuit=proveedor_cuit,
                        nombre=str(fields.get("proveedor_nombre") or "") or None,
                    )
                except CLIError as exc:
                    provider_cache[proveedor_cuit] = None
                    issues.append(str(exc))
            provider = provider_cache[proveedor_cuit]
        if provider is not None:
            fields["idclipro"] = str(find_value(provider, ("idclipro", "id", "idcliente")) or "")
            fields["proveedor_nombre"] = str(find_value(provider, ("clipro", "cliente", "nombre")) or fields.get("proveedor_nombre") or "")

        history: list[dict[str, Any]] = []
        if proveedor_cuit and fields.get("fecha"):
            cache_key = (proveedor_cuit, str(fields.get("fecha")))
            if cache_key not in history_cache:
                history_cache[cache_key] = fetch_recent_supplier_purchase_details(
                    bound_client,
                    proveedor_cuit=proveedor_cuit,
                    fecha=str(fields.get("fecha")),
                )
            history = history_cache[cache_key]
        defaults, history_warnings = infer_compra_historical_defaults(history)
        warnings.extend(history_warnings)
        for key in ("idcuenta", "idcentrocosto", "idprovinciaiibb"):
            explicit_value = document_overrides.get(key)
            if explicit_value not in (None, ""):
                fields[key] = str(explicit_value)
            elif defaults.get(key) not in (None, ""):
                fields[key] = str(defaults[key])
        if not fields.get("idprovinciaiibb") and provider is not None:
            provider_province = find_value(provider, ("idprovincia", "idprovinciaiibb"))
            if provider_province not in (None, ""):
                fields["idprovinciaiibb"] = str(provider_province)
        fields["cuenta_nombre"] = defaults.get("cuenta_nombre")
        fields["centrocosto_nombre"] = defaults.get("centrocosto_nombre")
        fields["antecedente_id"] = defaults.get("antecedente_id")
        fields["referencia"] = str(document_overrides.get("referencia") or source_path.name)
        fields["memo"] = str(document_overrides.get("memo") or "")

        required = {
            "fecha": fields.get("fecha"),
            "fcncnd": fields.get("fcncnd"),
            "letra": fields.get("letra"),
            "puntoventa": fields.get("puntoventa"),
            "numero": fields.get("numero"),
            "proveedor_cuit": proveedor_cuit,
            "idclipro": fields.get("idclipro"),
            "idcuenta": fields.get("idcuenta"),
            "idcentrocosto": fields.get("idcentrocosto"),
            "idprovinciaiibb": fields.get("idprovinciaiibb"),
            "total": fields.get("amounts", {}).get("total"),
        }
        missing = [key for key, value in required.items() if value in (None, "", 0)]
        if not any(decimal_or_zero(fields.get("amounts", {}).get(key)) for key in ("neto_0", "neto_10_5", "neto_21", "neto_27", "nogravado", "exento")):
            missing.append("imputaciones")

        duplicate = find_compra_duplicate(bound_client, fields=fields) if not missing else None
        status = "pendiente"
        reason = ""
        body: dict[str, Any] | None = None
        if duplicate and duplicate.get("kind") == "exact":
            status = "ya_cargado"
            reason = "Coincidencia exacta activa en SOS."
        elif duplicate:
            status = "verificar"
            reason = "La identidad ya existe, pero el estado o los importes no coinciden."
        elif missing or issues:
            status = "verificar"
            reason_parts = [f"Faltan: {', '.join(missing)}"] if missing else []
            reason = "; ".join([*reason_parts, *issues])
        else:
            body = build_compra_body_from_fields(fields)
        stats[status] += 1
        if status == "verificar":
            draft_warnings.append(f"Archivo {source_path.name}: {reason}")
        rows.append(
            {
                "source_index": index,
                "source": str(source_path),
                "status": status,
                "reason": reason,
                "documento": compra_document_code(
                    fcncnd=fields.get("fcncnd"),
                    letra=fields.get("letra"),
                    puntoventa=fields.get("puntoventa"),
                    numero=fields.get("numero"),
                ),
                "fields": fields,
                "expected_amounts": fields.get("amounts", {}),
                "missing_fields": missing,
                "warnings": warnings,
                "existing_match_id": (duplicate or {}).get("id"),
                "body": body,
                "body_sha256": compra_payload_hash(body) if body is not None else "",
            }
        )

    draft = {
        "draft_id": uuid.uuid4().hex[:12],
        "draft_kind": COMPRA_DOCUMENT_DRAFT_KIND,
        "contexto": {
            "cuit_trabajo": work_target,
            "fuentes": [{"path": source.get("path"), "kind": source.get("kind")} for source in extracted_sources],
        },
        "rows": rows,
        "stats": stats,
        "validacion": {"warnings": draft_warnings},
    }
    draft = make_json_safe(draft)
    saved_path = save_draft_payload(draft)
    draft["draft_file"] = str(saved_path)
    return draft


def build_compra_draft_preview(draft: dict[str, Any]) -> dict[str, Any]:
    context = draft.get("contexto", {}).get("cuit_trabajo", {})
    rows: list[dict[str, Any]] = []
    for row in draft.get("rows", []):
        fields = row.get("fields", {})
        amounts = row.get("expected_amounts", {})
        rows.append(
            {
                "Estado": row.get("status"),
                "Fecha": fields.get("fecha"),
                "Comprobante": row.get("documento"),
                "Proveedor": fields.get("proveedor_nombre"),
                "Neto 0 %": amounts.get("neto_0"),
                "Neto 10,5 %": amounts.get("neto_10_5"),
                "Neto 21 %": amounts.get("neto_21"),
                "Neto 27 %": amounts.get("neto_27"),
                "IVA 10,5 %": amounts.get("iva_10_5"),
                "IVA 21 %": amounts.get("iva_21"),
                "IVA 27 %": amounts.get("iva_27"),
                "No gravado": amounts.get("nogravado"),
                "Exento": amounts.get("exento"),
                "Percepción IIBB": amounts.get("percepcion_iibb"),
                "Otros tributos": amounts.get("otros"),
                "Total": amounts.get("total"),
                "Cuenta": fields.get("cuenta_nombre") or fields.get("idcuenta"),
                "Centro de costo": fields.get("centrocosto_nombre") or fields.get("idcentrocosto"),
                "Motivo": row.get("reason"),
            }
        )
    return {
        "titulo": "compra_draft",
        "contexto": [{"CUIT de trabajo": context.get("cuit"), "Contribuyente": context.get("nombre")}],
        "tablas": [{"titulo": "Compras", "rows": rows}],
        "totales": {
            "Pendientes": draft.get("stats", {}).get("pendiente", 0),
            "Ya cargadas": draft.get("stats", {}).get("ya_cargado", 0),
            "Verificar": draft.get("stats", {}).get("verificar", 0),
        },
        "advertencias": draft.get("validacion", {}).get("warnings", []),
    }


def render_compra_draft_markdown(draft: dict[str, Any]) -> str:
    return render_business_preview_markdown(build_compra_draft_preview(draft))


def command_compra_draft(args: argparse.Namespace, client: SOSContadorClient) -> Any:
    draft = build_compra_document_draft(args, client)
    preview = build_compra_draft_preview(draft)
    preview_markdown = render_compra_draft_markdown(draft)
    if getattr(args, "json_out", None):
        destination = Path(args.json_out)
        save_json_file(destination, draft)
        draft["json_written_to"] = str(destination)
    preview_format = str(getattr(args, "preview_format", "both") or "both").strip().lower()
    if preview_format == "markdown":
        return preview_markdown
    result = {
        "draft_id": draft.get("draft_id"),
        "draft_file": draft.get("draft_file"),
        "payload": draft,
        "business_preview": preview,
    }
    if preview_format == "both":
        result["preview_markdown"] = preview_markdown
    return result


def bind_compra_draft_client(
    draft: dict[str, Any],
    args: argparse.Namespace,
    client: SOSContadorClient,
) -> SOSContadorClient:
    context = draft.get("contexto", {}).get("cuit_trabajo", {})
    draft_cuit = digits_only(str(context.get("cuit") or ""))
    if not draft_cuit:
        raise CLIError("El borrador no contiene una CUIT de trabajo válida.")
    if any(
        getattr(args, key, None)
        for key in ("cuit_trabajo", "cuit_trabajo_id", "cuit_trabajo_nombre")
    ):
        explicit_target = resolve_work_cuit_target_from_args(client, args, required=True)
        if digits_only(str(explicit_target.get("cuit") or "")) != draft_cuit:
            raise CLIError("La CUIT de trabajo indicada no coincide con la del borrador.")
    return ensure_bound_client(
        client,
        cuit=draft_cuit,
        cuit_id=str(context.get("cuit_id") or ""),
    )


def verify_created_compra(
    bound_client: SOSContadorClient,
    *,
    compra_id: Any,
    row: dict[str, Any],
) -> dict[str, Any]:
    fields = row.get("fields", {})
    detail = bound_client.request("GET", f"compra/detalle/{compra_id}")
    cabecera = detail.get("cabecera", {}) if isinstance(detail, dict) else {}
    summary = summarize_compra_detail_for_afip(detail if isinstance(detail, dict) else {})
    issues: list[str] = []
    expected_identity = {
        "fecha": str(fields.get("fecha") or ""),
        "cuit": digits_only(str(fields.get("proveedor_cuit") or "")),
        "fcncnd": str(fields.get("fcncnd") or ""),
        "letra": str(fields.get("letra") or ""),
        "puntoventa": int(fields.get("puntoventa") or 0),
        "numero": int(fields.get("numero") or 0),
    }
    actual_identity = {
        "fecha": str(cabecera.get("fecha") or "")[:10],
        "cuit": digits_only(str(cabecera.get("cuit") or "")),
        "fcncnd": str(cabecera.get("fcncnd") or ""),
        "letra": str(cabecera.get("letra") or ""),
        "puntoventa": int(cabecera.get("puntoventa") or 0),
        "numero": int(cabecera.get("numero") or 0),
    }
    for key, expected_value in expected_identity.items():
        if actual_identity.get(key) != expected_value:
            issues.append(f"{key}: esperado {expected_value}, obtenido {actual_identity.get(key)}")
    for key in ("idcuenta", "idcentrocosto", "idprovinciaiibb"):
        if str(cabecera.get(key) or "") != str(fields.get(key) or ""):
            issues.append(f"{key}: esperado {fields.get(key)}, obtenido {cabecera.get(key)}")
    if not compra_summary_matches_expected(
        summary,
        row.get("expected_amounts", {}),
        fcncnd=fields.get("fcncnd"),
    ):
        issues.append("Los importes o tratamientos impositivos persistidos no coinciden con el borrador.")

    rows, truncated = fetch_compra_items_complete(
        bound_client,
        desde=str(fields.get("fecha") or ""),
        hasta=str(fields.get("fecha") or ""),
    )
    listed = next((item for item in rows if str(item.get("id") or "") == str(compra_id)), None)
    if listed is None:
        issues.append("La compra no aparece en la consulta del período esperado.")
    elif not compra_row_is_active(listed):
        issues.append("La compra aparece eliminada, archivada o anulada.")
    if truncated and listed is None:
        issues.append("La consulta del día quedó truncada en 50 registros.")
    if issues:
        raise CLIError(
            f"La compra {compra_id} fue creada, pero la verificación posterior falló: " + "; ".join(issues)
        )
    return {
        "id": compra_id,
        "documento": row.get("documento"),
        "fecha": fields.get("fecha"),
        "proveedor": fields.get("proveedor_nombre"),
        "total": row.get("expected_amounts", {}).get("total"),
        "public_active": True,
    }


def command_compra_create(args: argparse.Namespace, client: SOSContadorClient) -> Any:
    draft = load_draft_payload(
        draft_id=getattr(args, "draft_id", None),
        draft_file=getattr(args, "draft_file", None),
    )
    if draft.get("draft_kind") != COMPRA_DOCUMENT_DRAFT_KIND:
        raise CLIError("El borrador indicado no corresponde a compras desde documentos.")
    if int(draft.get("stats", {}).get("verificar") or 0):
        raise CLIError("El borrador contiene compras que requieren verificación. Genere un nuevo borrador corregido.")
    preview = build_compra_draft_preview(draft)
    preview_markdown = render_compra_draft_markdown(draft)
    pending_rows = [row for row in draft.get("rows", []) if row.get("status") == "pendiente"]
    mutation_payload = {
        "draft_id": draft.get("draft_id"),
        "bodies": [row.get("body") for row in pending_rows],
    }
    set_mutation_preview(args, business_preview=preview, preview_markdown=preview_markdown)
    ensure_mutation_allowed("PUT", "compra/create", None, mutation_payload, args)
    bound_client = bind_compra_draft_client(draft, args, client)

    ready_rows: list[dict[str, Any]] = []
    skipped_existing: list[dict[str, Any]] = []
    for row in pending_rows:
        body = row.get("body")
        if not isinstance(body, dict) or compra_payload_hash(body) != row.get("body_sha256"):
            raise CLIError("El payload del borrador cambió. Genere un nuevo borrador antes de confirmar.")
        fields = row.get("fields", {})
        provider = resolve_cliente_match(
            bound_client,
            explicit_id=None,
            cuit=str(fields.get("proveedor_cuit") or ""),
            nombre=str(fields.get("proveedor_nombre") or "") or None,
        )
        current_provider_id = str(find_value(provider, ("idclipro", "id", "idcliente")) or "")
        if current_provider_id != str(body.get("idclipro") or ""):
            raise CLIError("El proveedor del borrador cambió. Genere un nuevo borrador antes de confirmar.")
        duplicate = find_compra_duplicate(bound_client, fields=fields)
        if duplicate and duplicate.get("kind") == "exact":
            skipped_existing.append(
                {"documento": row.get("documento"), "id": duplicate.get("id"), "motivo": "ya_cargado"}
            )
            continue
        if duplicate:
            raise CLIError(
                f"El comprobante {row.get('documento')} apareció antes de la escritura con datos incompatibles. No se registró el lote."
            )
        ready_rows.append(row)

    created: list[dict[str, Any]] = []
    for row in ready_rows:
        result = bound_client.request("PUT", "compra/0", body=row.get("body"))
        compra_id = result.get("id") if isinstance(result, dict) else None
        if not compra_id:
            raise CLIError("SOS no devolvió el ID de la compra creada.")
        created.append(verify_created_compra(bound_client, compra_id=compra_id, row=row))

    internal_status_warning = ""
    if created:
        fechas = [str(item.get("fecha") or "") for item in created]
        try:
            status_index = fetch_web_comprobante_status_index(
                bound_client,
                idtipo_operacion=4,
                desde_iso=min(fechas),
                hasta_iso=max(fechas),
            )
            annulled = [
                item
                for item in created
                if bool((status_index.get(str(item.get("id") or "")) or {}).get("annulled"))
            ]
            if annulled:
                ids = ", ".join(str(item.get("id")) for item in annulled)
                raise CLIError(f"La verificación interna detectó compras anuladas: {ids}.")
            not_listed_ids: list[str] = []
            for item in created:
                item["internal_active"] = str(item.get("id") or "") in status_index
                if not item["internal_active"]:
                    not_listed_ids.append(str(item.get("id") or ""))
            if not_listed_ids:
                internal_status_warning = (
                    "La verificación pública fue satisfactoria, pero la consulta interna no mostró las compras: "
                    + ", ".join(not_listed_ids)
                    + "."
                )
        except (APIError, CLIError) as exc:
            if isinstance(exc, CLIError) and "compras anuladas" in str(exc):
                raise
            internal_status_warning = f"No se pudo completar la verificación interna: {exc}"

    return {
        "draft_id": draft.get("draft_id"),
        "created_compras": created,
        "skipped_existing": skipped_existing,
        "warnings": [internal_status_warning] if internal_status_warning else [],
        "preview_markdown": preview_markdown,
    }


def command_puntoventa_list(args: argparse.Namespace, client: SOSContadorClient) -> Any:
    bound_client, _ = bind_business_client(client, args, required=True)
    return bound_client.request("GET", "puntoventa/listado")


def add_body_source_arguments(parser: argparse.ArgumentParser) -> None:
    parser.add_argument("--body-json")
    parser.add_argument("--body-file")


def add_mutation_flags(parser: argparse.ArgumentParser) -> None:
    parser.add_argument("--dry-run", action="store_true")
    parser.add_argument("--confirm", action="store_true")


def add_work_cuit_arguments(parser: argparse.ArgumentParser) -> None:
    parser.add_argument("--cuit-trabajo")
    parser.add_argument("--cuit-trabajo-id")
    parser.add_argument("--cuit-trabajo-nombre")


def add_cobro_pago_creation_arguments(parser: argparse.ArgumentParser) -> None:
    add_body_source_arguments(parser)
    add_mutation_flags(parser)
    add_work_cuit_arguments(parser)
    parser.add_argument("--fecha")
    parser.add_argument("--idclipro")
    parser.add_argument("--cliente-cuit")
    parser.add_argument("--cliente-nombre")
    parser.add_argument("--numero")
    parser.add_argument("--idcuenta")
    parser.add_argument("--idprovinciaiibb")
    parser.add_argument("--provincia")
    parser.add_argument("--idcentrocosto")
    parser.add_argument("--centrocosto")
    parser.add_argument("--comentarios")
    parser.add_argument("--memo")
    parser.add_argument("--referencia")
    parser.add_argument("--imputaciones-json")
    parser.add_argument("--imputaciones-file")
    parser.add_argument("--movimientos-json")
    parser.add_argument("--movimientos-file")


def add_cobro_draft_arguments(parser: argparse.ArgumentParser) -> None:
    parser.add_argument("--source", action="append", required=True)
    add_work_cuit_arguments(parser)
    parser.add_argument("--cliente-cuit")
    parser.add_argument("--cliente-nombre")
    parser.add_argument("--fecha")
    parser.add_argument("--comentarios")
    parser.add_argument("--centrocosto")
    parser.add_argument("--factura", action="append")
    parser.add_argument("--venta-id", action="append")
    parser.add_argument("--json-out")
    parser.add_argument("--preview-format", choices=("json", "markdown", "both"), default="json")


def add_cobro_profile_create_arguments(parser: argparse.ArgumentParser) -> None:
    parser.add_argument("--draft-id")
    parser.add_argument("--draft-file")
    parser.add_argument("--source", action="append")
    parser.add_argument("--cuit-trabajo")
    parser.add_argument("--counterparty-cuit")
    parser.add_argument("--document-kind", default=PROFILE_DOCUMENT_KIND_COBRO)
    parser.add_argument("--name")
    parser.add_argument("--profile-id")
    parser.add_argument("--json-out")


def add_compra_draft_arguments(parser: argparse.ArgumentParser) -> None:
    parser.add_argument("--source", action="append", required=True)
    parser.add_argument("--document-json", action="append")
    parser.add_argument("--document-file", action="append")
    add_work_cuit_arguments(parser)
    parser.add_argument("--json-out")
    parser.add_argument("--preview-format", choices=("json", "markdown", "both"), default="both")


def add_afip_source_arguments(parser: argparse.ArgumentParser) -> None:
    parser.add_argument("--source", required=True)
    add_work_cuit_arguments(parser)
    parser.add_argument("--json-out")


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="CLI helper para la API de SOS Contador")
    subparsers = parser.add_subparsers(dest="command", required=True)

    auth = subparsers.add_parser("auth")
    auth_sub = auth.add_subparsers(dest="auth_command", required=True)
    auth_info = auth_sub.add_parser("info")
    auth_info.add_argument("--show-tokens", action="store_true")
    auth_info.set_defaults(handler=command_auth_info)
    auth_list_cuits = auth_sub.add_parser("list-cuits")
    auth_list_cuits.set_defaults(handler=command_auth_list_cuits)
    auth_resolve_cuit = auth_sub.add_parser("resolve-cuit")
    auth_resolve_cuit.add_argument("--name", required=True)
    auth_resolve_cuit.add_argument("--refresh", action="store_true")
    auth_resolve_cuit.set_defaults(handler=command_auth_resolve_cuit)
    auth_web_info = auth_sub.add_parser("web-info")
    add_work_cuit_arguments(auth_web_info)
    auth_web_info.set_defaults(handler=command_auth_web_info)

    api = subparsers.add_parser("api")
    api_sub = api.add_subparsers(dest="api_command", required=True)

    api_catalog = api_sub.add_parser("catalog")
    api_catalog.add_argument("--module")
    api_catalog.add_argument("--method", choices=("GET", "POST", "PUT", "DELETE", "PATCH"))
    api_catalog.add_argument("--status")
    api_catalog.set_defaults(handler=command_api_catalog)

    api_describe = api_sub.add_parser("describe")
    api_describe.add_argument("--operation", required=True)
    api_describe.set_defaults(handler=command_api_describe)

    api_invoke = api_sub.add_parser("invoke")
    api_invoke.add_argument("--operation", required=True)
    add_work_cuit_arguments(api_invoke)
    api_invoke.add_argument("--param", action="append")
    api_invoke.add_argument("--query", action="append")
    api_invoke.add_argument("--out")
    add_body_source_arguments(api_invoke)
    add_mutation_flags(api_invoke)
    api_invoke.set_defaults(handler=command_api_invoke)

    call = subparsers.add_parser("call")
    add_work_cuit_arguments(call)
    call.add_argument("--method", required=True)
    call.add_argument("--path", required=True)
    call.add_argument("--query", action="append")
    call.add_argument("--auth-mode", choices=("none", "jwt", "jwtc"), default="jwtc")
    call.add_argument("--out")
    add_body_source_arguments(call)
    add_mutation_flags(call)
    call.set_defaults(handler=command_call)

    afip = subparsers.add_parser("afip")
    afip_sub = afip.add_subparsers(dest="afip_command", required=True)

    afip_draft = afip_sub.add_parser("draft")
    add_afip_source_arguments(afip_draft)
    afip_draft.set_defaults(handler=command_afip_draft)

    afip_import = afip_sub.add_parser("import")
    add_afip_source_arguments(afip_import)
    add_mutation_flags(afip_import)
    afip_import.set_defaults(handler=command_afip_import)

    cliente = subparsers.add_parser("cliente")
    cliente_sub = cliente.add_subparsers(dest="cliente_command", required=True)

    cliente_list = cliente_sub.add_parser("list")
    add_work_cuit_arguments(cliente_list)
    cliente_list.add_argument("--cliente", action=argparse.BooleanOptionalAction, default=True)
    cliente_list.add_argument("--proveedor", action=argparse.BooleanOptionalAction, default=True)
    cliente_list.add_argument("--pagina", type=int, default=1)
    cliente_list.add_argument("--registros", type=int, default=50)
    cliente_list.set_defaults(handler=command_cliente_list)

    cliente_get = cliente_sub.add_parser("get")
    add_work_cuit_arguments(cliente_get)
    cliente_get.add_argument("--id")
    cliente_get.add_argument("--cuit")
    cliente_get.add_argument("--nombre")
    cliente_get.set_defaults(handler=command_cliente_get)

    cliente_create = cliente_sub.add_parser("create")
    add_work_cuit_arguments(cliente_create)
    add_body_source_arguments(cliente_create)
    add_mutation_flags(cliente_create)
    cliente_create.set_defaults(handler=command_cliente_create)

    cliente_update = cliente_sub.add_parser("update")
    add_work_cuit_arguments(cliente_update)
    cliente_update.add_argument("--id", required=True)
    add_body_source_arguments(cliente_update)
    add_mutation_flags(cliente_update)
    cliente_update.set_defaults(handler=command_cliente_update)

    cliente_delete = cliente_sub.add_parser("delete")
    add_work_cuit_arguments(cliente_delete)
    cliente_delete.add_argument("--id", required=True)
    add_mutation_flags(cliente_delete)
    cliente_delete.set_defaults(handler=command_cliente_delete)

    compra = subparsers.add_parser("compra")
    compra_sub = compra.add_subparsers(dest="compra_command", required=True)

    compra_draft = compra_sub.add_parser("draft")
    add_compra_draft_arguments(compra_draft)
    compra_draft.set_defaults(handler=command_compra_draft)

    compra_create = compra_sub.add_parser("create")
    compra_create.add_argument("--draft-id")
    compra_create.add_argument("--draft-file")
    add_work_cuit_arguments(compra_create)
    add_mutation_flags(compra_create)
    compra_create.set_defaults(handler=command_compra_create)

    producto = subparsers.add_parser("producto")
    producto_sub = producto.add_subparsers(dest="producto_command", required=True)

    producto_list = producto_sub.add_parser("list")
    add_work_cuit_arguments(producto_list)
    producto_list.add_argument("--pagina", type=int, default=1)
    producto_list.add_argument("--registros", type=int, default=50)
    producto_list.set_defaults(handler=command_producto_list)

    producto_create = producto_sub.add_parser("create")
    add_work_cuit_arguments(producto_create)
    add_body_source_arguments(producto_create)
    add_mutation_flags(producto_create)
    producto_create.set_defaults(handler=command_producto_create)

    producto_update = producto_sub.add_parser("update")
    add_work_cuit_arguments(producto_update)
    producto_update.add_argument("--id", required=True)
    add_body_source_arguments(producto_update)
    add_mutation_flags(producto_update)
    producto_update.set_defaults(handler=command_producto_update)

    cobro = subparsers.add_parser("cobro")
    cobro_sub = cobro.add_subparsers(dest="cobro_command", required=True)

    cobro_list = cobro_sub.add_parser("list")
    add_work_cuit_arguments(cobro_list)
    cobro_list.add_argument("--periodo", default="mes")
    cobro_list.add_argument("--pagina", type=int, default=1)
    cobro_list.add_argument("--registros", type=int, default=50)
    cobro_list.set_defaults(handler=lambda args, client: command_periodic_list(args, client, "cobro"))

    cobro_get = cobro_sub.add_parser("get")
    add_work_cuit_arguments(cobro_get)
    cobro_get.add_argument("--id")
    cobro_get.add_argument("--recibo")
    cobro_get.add_argument("--comprobante")
    cobro_get.add_argument("--fecha")
    cobro_get.add_argument("--desde")
    cobro_get.add_argument("--hasta")
    cobro_get.add_argument("--cliente-cuit")
    cobro_get.add_argument("--cliente-nombre")
    cobro_get.add_argument("--registros", type=int, default=500)
    cobro_get.set_defaults(handler=command_cobro_get)

    cobro_list_range = cobro_sub.add_parser("list-range")
    add_work_cuit_arguments(cobro_list_range)
    cobro_list_range.add_argument("--desde", required=True)
    cobro_list_range.add_argument("--hasta", required=True)
    cobro_list_range.add_argument("--recibo")
    cobro_list_range.add_argument("--comprobante")
    cobro_list_range.add_argument("--cliente-cuit")
    cobro_list_range.add_argument("--cliente-nombre")
    cobro_list_range.add_argument("--registros", type=int, default=500)
    cobro_list_range.add_argument("--csv-out")
    cobro_list_range.set_defaults(handler=command_cobro_list_range)

    cobro_resolve_id = cobro_sub.add_parser("resolve-id")
    add_work_cuit_arguments(cobro_resolve_id)
    cobro_resolve_id.add_argument("--fecha")
    cobro_resolve_id.add_argument("--desde")
    cobro_resolve_id.add_argument("--hasta")
    cobro_resolve_id.add_argument("--recibo")
    cobro_resolve_id.add_argument("--comprobante")
    cobro_resolve_id.add_argument("--cliente-cuit")
    cobro_resolve_id.add_argument("--cliente-nombre")
    cobro_resolve_id.add_argument("--registros", type=int, default=500)
    cobro_resolve_id.set_defaults(handler=command_cobro_resolve_id)

    cobro_draft = cobro_sub.add_parser("draft")
    add_cobro_draft_arguments(cobro_draft)
    cobro_draft.set_defaults(handler=command_cobro_draft)

    cobro_profile = cobro_sub.add_parser("profile")
    cobro_profile_sub = cobro_profile.add_subparsers(dest="cobro_profile_command", required=True)
    cobro_profile_create = cobro_profile_sub.add_parser("create")
    add_cobro_profile_create_arguments(cobro_profile_create)
    cobro_profile_create.set_defaults(handler=command_cobro_profile_create)

    cobro_create = cobro_sub.add_parser("create")
    add_cobro_pago_creation_arguments(cobro_create)
    cobro_create.add_argument("--draft-id")
    cobro_create.add_argument("--draft-file")
    cobro_create.set_defaults(handler=command_cobro_create)

    cobro_asociar = cobro_sub.add_parser("asociar")
    add_mutation_flags(cobro_asociar)
    add_work_cuit_arguments(cobro_asociar)
    cobro_asociar.add_argument("--id")
    cobro_asociar.add_argument("--recibo")
    cobro_asociar.add_argument("--comprobante")
    cobro_asociar.add_argument("--fecha")
    cobro_asociar.add_argument("--desde")
    cobro_asociar.add_argument("--hasta")
    cobro_asociar.add_argument("--cliente-cuit")
    cobro_asociar.add_argument("--cliente-nombre")
    cobro_asociar.add_argument("--registros", type=int, default=500)
    cobro_asociar.add_argument("--venta-id", action="append")
    cobro_asociar.add_argument("--factura", action="append")
    cobro_asociar.set_defaults(handler=command_cobro_asociar)

    pago = subparsers.add_parser("pago")
    pago_sub = pago.add_subparsers(dest="pago_command", required=True)

    pago_list = pago_sub.add_parser("list")
    add_work_cuit_arguments(pago_list)
    pago_list.add_argument("--periodo", default="mes")
    pago_list.add_argument("--pagina", type=int, default=1)
    pago_list.add_argument("--registros", type=int, default=50)
    pago_list.set_defaults(handler=lambda args, client: command_periodic_list(args, client, "pago"))

    pago_get = pago_sub.add_parser("get")
    add_work_cuit_arguments(pago_get)
    pago_get.add_argument("--id", required=True)
    pago_get.set_defaults(handler=lambda args, client: command_detail(args, client, "pago"))

    pago_create = pago_sub.add_parser("create")
    add_cobro_pago_creation_arguments(pago_create)
    pago_create.set_defaults(handler=command_pago_create)

    puntoventa = subparsers.add_parser("puntoventa")
    puntoventa_sub = puntoventa.add_subparsers(dest="puntoventa_command", required=True)
    puntoventa_list = puntoventa_sub.add_parser("list")
    add_work_cuit_arguments(puntoventa_list)
    puntoventa_list.set_defaults(handler=command_puntoventa_list)

    venta = subparsers.add_parser("venta")
    venta_sub = venta.add_subparsers(dest="venta_command", required=True)

    venta_list = venta_sub.add_parser("list")
    add_work_cuit_arguments(venta_list)
    venta_list.add_argument("--modo", default="facturas")
    venta_list.add_argument("--periodo", default="mes")
    venta_list.add_argument("--cae", default="todas")
    venta_list.add_argument("--pagina", type=int, default=1)
    venta_list.add_argument("--registros", type=int, default=50)
    venta_list.add_argument("--fecha-desde", dest="fecha_desde")
    venta_list.add_argument("--fecha-hasta", dest="fecha_hasta")
    venta_list.set_defaults(handler=command_venta_list)

    venta_list_all = venta_sub.add_parser("list-all")
    add_work_cuit_arguments(venta_list_all)
    venta_list_all.add_argument("--periodo", default="mes")
    venta_list_all.add_argument("--cae", default="todas")
    venta_list_all.add_argument("--pagina", type=int, default=1)
    venta_list_all.add_argument("--registros", type=int, default=50)
    venta_list_all.add_argument("--fecha-desde", dest="fecha_desde")
    venta_list_all.add_argument("--fecha-hasta", dest="fecha_hasta")
    venta_list_all.set_defaults(handler=command_venta_list_all)

    venta_create = venta_sub.add_parser("create")
    add_work_cuit_arguments(venta_create)
    add_body_source_arguments(venta_create)
    add_mutation_flags(venta_create)
    venta_create.add_argument("--fecha")
    venta_create.add_argument("--idclipro")
    venta_create.add_argument("--cliente-cuit")
    venta_create.add_argument("--cliente-nombre")
    venta_create.add_argument("--fcncnd", default="F")
    venta_create.add_argument("--letra")
    venta_create.add_argument("--sucursal")
    venta_create.add_argument("--numero")
    venta_create.add_argument("--numero-hasta", dest="numero_hasta")
    venta_create.add_argument("--idcuenta")
    venta_create.add_argument("--idprovinciaiibb")
    venta_create.add_argument("--provincia")
    venta_create.add_argument("--idcentrocosto")
    venta_create.add_argument("--centrocosto")
    venta_create.add_argument("--codactividad")
    venta_create.add_argument("--observaciones")
    venta_create.add_argument("--moneda", default="ARS")
    venta_create.add_argument("--tipocambio", default="1")
    venta_create.add_argument("--cuentacobropago", default="0")
    venta_create.add_argument("--obtienecae", action="store_true")
    venta_create.add_argument("--productos-json")
    venta_create.add_argument("--productos-file")
    venta_create.add_argument("--imputa-json")
    venta_create.add_argument("--imputa-file")
    venta_create.set_defaults(handler=command_venta_create)

    venta_get = venta_sub.add_parser("get")
    add_work_cuit_arguments(venta_get)
    venta_get.add_argument("--id", required=True)
    venta_get.set_defaults(handler=command_venta_get)

    venta_pdf = venta_sub.add_parser("pdf")
    add_work_cuit_arguments(venta_pdf)
    venta_pdf.add_argument("--id", required=True)
    venta_pdf.add_argument("--out", required=True)
    venta_pdf.add_argument("--allow-web-session-fallback", action="store_true")
    venta_pdf.set_defaults(handler=command_venta_pdf)

    venta_search = venta_sub.add_parser("search")
    add_work_cuit_arguments(venta_search)
    venta_search.add_argument("--pagina", type=int, default=1)
    venta_search.add_argument("--registros", type=int, default=50)
    add_body_source_arguments(venta_search)
    venta_search.set_defaults(handler=command_venta_search)

    return parser


def main(argv: list[str] | None = None) -> int:
    parser = build_parser()
    args = parser.parse_args(argv)
    client = SOSContadorClient()
    try:
        result = args.handler(args, client)
    except APIError as exc:
        print(str(exc), file=sys.stderr)
        return 1
    except CLIError as exc:
        print(str(exc), file=sys.stderr)
        return 1

    if result is not None:
        print(format_output(result))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
